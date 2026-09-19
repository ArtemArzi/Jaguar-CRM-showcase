"""Bounded XLSX reader. Source values stay private and never enter diagnostics."""

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from apps.common.exceptions import BusinessLogicError

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_MEMBERS = 100
MAX_CELLS = 100_000
MAX_ENTITLEMENTS = 2000
ENTITLEMENTS_SHEET = "Абонементы"
SETTLEMENTS_SHEET = "Расчёты с тренерами"
METADATA_SHEET = "_Источник"


def invalid(message, code="import_invalid_workbook"):
    raise BusinessLogicError(message, code=code)


@dataclass(frozen=True)
class WorkbookRow:
    sheet: str
    number: int
    values: dict


@dataclass(frozen=True)
class ParsedWorkbook:
    namespace: str | None
    rows: tuple[WorkbookRow, ...]


def read_workbook(*, path: Path) -> ParsedWorkbook:
    path = Path(path)
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > MAX_FILE_BYTES:
        invalid("Нужен файл .xlsx размером не более 10 МБ.")
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            names = [member.filename for member in members]
            if (
                len(members) > MAX_MEMBERS
                or len(set(names)) != len(names)
                or sum(member.file_size for member in members) > MAX_EXPANDED_BYTES
                or any(member.flag_bits & 1 for member in members)
                or "[Content_Types].xml" not in names
                or "xl/workbook.xml" not in names
            ):
                invalid("Контейнер Excel повреждён или превышает допустимый размер.")
            if any("vbaproject" in name.lower() or "externallinks/" in name.lower() for name in names):
                invalid("Удалите макросы и внешние ссылки из книги.")
            content_types = archive.read("[Content_Types].xml").lower()
            if b"macroenabled" in content_types or b"vba" in content_types:
                invalid("Книга с макросами не поддерживается.")
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if ENTITLEMENTS_SHEET not in book.sheetnames:
                invalid("В книге отсутствует лист «Абонементы».")
            rows, namespace, nonempty_cells, entitlement_count = [], None, 0, 0
            # Inspect every sheet: hidden sheets cannot bypass size/formula guards.
            for sheet in book:
                sheet.reset_dimensions()
                headers = None
                for number, cells in enumerate(sheet.iter_rows(), start=1):
                    if number > 10_000 or len(cells) > 256:
                        invalid("Лист превышает предел 10 000 строк или 256 столбцов.")
                    values = []
                    for cell in cells:
                        value = cell.value
                        if value is not None:
                            nonempty_cells += 1
                            if nonempty_cells > MAX_CELLS:
                                invalid("В книге больше 100 000 заполненных ячеек.")
                            if cell.data_type == "f":
                                invalid("Замените формулы подтверждёнными значениями.")
                        values.append(value)
                    if not any(value is not None for value in values):
                        continue
                    if sheet.title == METADATA_SHEET:
                        if values[0] == "source_namespace" and len(values) > 1:
                            if namespace is not None or not isinstance(values[1], str):
                                invalid("Проверьте технический источник книги.")
                            namespace = values[1].strip()
                            if not namespace or len(namespace) > 120:
                                invalid("Проверьте технический источник книги.")
                        continue
                    if sheet.title not in {ENTITLEMENTS_SHEET, SETTLEMENTS_SHEET}:
                        continue
                    if headers is None:
                        headers = [str(value).strip() if value is not None else "" for value in values]
                        populated = [header for header in headers if header]
                        if len(set(populated)) != len(populated):
                            invalid("В заголовке повторяются названия столбцов.")
                        continue
                    if any(value is not None for value in values[len(headers):]):
                        invalid("У заполненного столбца отсутствует заголовок.")
                    if sheet.title == ENTITLEMENTS_SHEET:
                        entitlement_count += 1
                        if entitlement_count > MAX_ENTITLEMENTS:
                            invalid("Разделите книгу на партии не более 2 000 абонементов.")
                    rows.append(WorkbookRow(sheet.title, number, {
                        header: values[index] if index < len(values) else None
                        for index, header in enumerate(headers) if header
                    }))
            return ParsedWorkbook(namespace=namespace, rows=tuple(rows))
        finally:
            book.close()
    except (BadZipFile, KeyError, ValueError, OSError, ParseError) as exc:
        raise BusinessLogicError("Не удалось прочитать контейнер Excel.", code="import_invalid_workbook") from exc


def parse_money(value) -> Decimal:
    if value is None or isinstance(value, bool):
        invalid("Укажите сумму явно; пустая ячейка не означает ноль.", "import_invalid_value")
    try:
        amount = Decimal(str(value).strip().replace(",", "."))
        if not amount.is_finite() or abs(amount) > Decimal("99999999.99"):
            raise InvalidOperation
        if amount != amount.quantize(Decimal("0.01")):
            raise InvalidOperation
        return amount.quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        invalid("Сумма должна быть числом с точностью до копеек.", "import_invalid_value")


def parse_date(value) -> date:
    if isinstance(value, datetime):
        if value.time().isoformat() != "00:00:00":
            invalid("Для даты укажите день без времени.", "import_invalid_value")
        return value.date()
    if isinstance(value, date):
        return value
    try:
        parsed = date.fromisoformat(value)
        if parsed.isoformat() != value:
            raise ValueError
        return parsed
    except (TypeError, ValueError):
        invalid("Дата должна быть датой Excel или иметь вид ГГГГ-ММ-ДД.", "import_invalid_value")


def parse_phone(value) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        invalid("Телефон должен быть текстом: проверьте сохранность всех цифр.", "import_invalid_value")
    return value.strip()


def write_text(cell, value):
    """Export source text literally, including leading =, +, - and @."""
    cell.value = str(value)
    cell.data_type = "s"
