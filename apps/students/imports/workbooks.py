"""Prepared books preserve source identities independently of row order or filename."""

import io
import uuid
from datetime import date, datetime

from django.urls import reverse
from openpyxl import Workbook

from apps.billing.models import Tariff
from apps.billing.service_modules.opening_subscriptions import _authorize
from apps.students.imports.parsing import (
    ENTITLEMENTS_SHEET,
    METADATA_SHEET,
    SETTLEMENTS_SHEET,
    write_text,
)
from apps.students.imports.schemas import COLUMNS, EXTRA_COLUMNS
from apps.trainers.models import Trainer

SETTLEMENT_COLUMNS = (
    "source_key",
    "Вид записи",
    "ID тренера",
    "Дата",
    "Сумма",
    "Причина",
    "Подтверждено",
    "Способ выплаты",
    "Аванс подтверждён",
    "Ключи абонементов",
)


def _append(sheet, values):
    row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, start=1):
        if value is None:
            continue
        cell = sheet.cell(row, column)
        if isinstance(value, (date, datetime)):
            write_text(cell, value.isoformat())
        elif isinstance(value, str):
            write_text(cell, value)
        else:
            cell.value = value


def workbook_bytes(*, club_id: int, actor_user_id: int, namespace: str, rows=(), results=()):
    from apps.attendance.models import TrainingGroup

    _authorize(club_id=club_id, actor_user_id=actor_user_id)
    book = Workbook()
    book.active.title = ENTITLEMENTS_SHEET
    columns = list(COLUMNS.values()) + list(EXTRA_COLUMNS)
    _append(book.active, columns)
    settlement = book.create_sheet(SETTLEMENTS_SHEET)
    _append(settlement, SETTLEMENT_COLUMNS)
    for row in rows:
        sheet = book[row.sheet]
        headers = columns if row.sheet == ENTITLEMENTS_SHEET else SETTLEMENT_COLUMNS
        _append(sheet, [row.values.get(header) for header in headers])
    reference = book.create_sheet("Справочники")
    _append(reference, ["Вид", "ID", "Название"])
    for trainer in Trainer.objects.for_club(club_id).order_by("id"):
        _append(reference, ["Тренер", trainer.id, f"{trainer.first_name} {trainer.last_name}".strip()])
    for model, label in ((TrainingGroup, "Группа"), (Tariff, "Тариф")):
        for record in model.objects.for_club(club_id).order_by("id"):
            _append(reference, [label, record.id, record.name])
    check = book.create_sheet("Проверка")
    _append(check, ["ID записи", "Статус", "Пояснение", "Ученик", "Карточка", "Тренер", "Расчётная запись"])
    for result in results:
        student_id = result.get("student_id")
        card = reverse("student-card", kwargs={"student_id": student_id}) if student_id else ""
        if result.get("trainer_id"):
            card = reverse("trainer-detail", kwargs={"trainer_id": result["trainer_id"]})
        _append(
            check,
            [
                result.get("id"),
                result.get("status"),
                result.get("message"),
                student_id,
                card,
                result.get("trainer_id"),
                result.get("settlement_entry_id"),
            ],
        )
    instructions = book.create_sheet("Как заполнить")
    for values in (
        ["Поле / шаг", "Правило"],
        [
            "Одна строка",
            "Один полностью оплаченный конечный абонемент. Не добавляйте исторические долги как посещения.",
        ],
        ["Сначала проверка", "Пустая ячейка не означает ноль. Спорные строки остаются на сверке."],
        ["Ключи", "Сохраняйте подготовленную копию и технические ключи. Перестановка строк безопасна."],
        [
            "Тип пакета",
            "Групповой / Персональный. Смешанные, безлимитные, недельные, "
            "замороженные и частично оплаченные требуют сверки.",
        ],
        [
            "Правило начислений",
            "С оплаты / За занятие / Без начисления (также принимаются on_payment / on_checkin / none).",
        ],
        ["Способ оплаты", "Наличные / Перевод / Неизвестен. Исторический внешний канал укажите в комментарии."],
        [
            "Ребёнок и подтверждения",
            "Да / Нет. Ответственность, владелец пакета и получатель комиссии — отдельные поля.",
        ],
        ["Суммы", "Рубли; точность до копеек. Исходная цена должна равняться подтверждённой оплате."],
        ["Занятия", "Всего = использовано + осталось. Исходный лимит сохраняет денежную базу занятия."],
        ["Даты", "ГГГГ-ММ-ДД. Сверено по и Начало работы в CRM — точные моменты в часовом поясе клуба."],
        [
            "Граница переноса",
            "Посещения до включительной границы Сверено по уже входят в использовано. "
            "Первое занятие после границы выбирается точно.",
        ],
        ["Идентификаторы", "ID и названия возьмите из Справочников. Если переданы оба, они должны совпадать."],
        [
            "Начальный расчёт",
            "Долг на начало дня со знаком: положительный долг, отрицательный аванс. Ноль вводится явно.",
        ],
        ["Выплата", "Уже выданные деньги: дата, сумма, способ и причина. Выплата не меняет заработанное."],
        ["Безопасность", "Не храните в книге пароли. Замените формулы значениями; макросы и внешние связи запрещены."],
    ):
        _append(instructions, values)
    instructions.column_dimensions["A"].width = 28
    instructions.column_dimensions["B"].width = 100
    metadata = book.create_sheet(METADATA_SHEET)
    _append(metadata, ["source_namespace", namespace or uuid.uuid4().hex])
    metadata.sheet_state = "hidden"
    for sheet in book:
        sheet.freeze_panes = "A2"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
