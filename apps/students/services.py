from __future__ import annotations

import logging
from dataclasses import dataclass, field
from datetime import date

from django.db import transaction

from apps.common.exceptions import BusinessLogicError
from apps.common.phone import is_valid_normalized_phone, normalize_phone
from apps.students.duplicates import (
    DuplicateStudentError,
    find_duplicate_for_contact,
)
from apps.students.identity_services import (
    arbitrate_person_identity,
    lock_club_person_identity_arbitration,
    normalize_person_identity,
    resolve_staff_intake_person_identity,
)
from apps.students.models import Student, StudentNote

logger = logging.getLogger(__name__)

ALLOWED_TRANSITIONS: dict[str, list[str]] = {
    Student.Status.LEAD: [Student.Status.TRIAL],
    Student.Status.TRIAL: [Student.Status.ACTIVE, Student.Status.LEAD],
    Student.Status.ACTIVE: [Student.Status.AT_RISK],
    Student.Status.AT_RISK: [Student.Status.ACTIVE, Student.Status.CHURNED],
    Student.Status.CHURNED: [Student.Status.ACTIVE, Student.Status.LOST],
    Student.Status.LOST: [],
}


@dataclass(frozen=True)
class ImportResult:
    created: int
    skipped: int
    errors: list[str] = field(default_factory=list)


def _normalize_phone(phone: str) -> str:
    """Strip formatting, convert 8→+7 prefix for Russian numbers."""
    return normalize_phone(phone)


def _normalize_optional_phone(phone: str | None) -> str:
    return normalize_phone(phone or "") if str(phone or "").strip() else ""


def _require_valid_phone(phone: str, *, code: str = "invalid_phone") -> None:
    if not phone or not is_valid_normalized_phone(phone):
        raise BusinessLogicError("Введите корректный телефон", code=code)


def _prepare_student_contact(*, is_child: bool, phone: str, guardian_phone: str = "") -> tuple[str, str]:
    normalized_phone = _normalize_optional_phone(phone)
    normalized_guardian_phone = _normalize_optional_phone(guardian_phone)

    if is_child:
        if not normalized_guardian_phone:
            normalized_guardian_phone = normalized_phone
            normalized_phone = ""
        _require_valid_phone(normalized_guardian_phone, code="guardian_phone_required")
        if normalized_phone:
            _require_valid_phone(normalized_phone)
        return normalized_phone, normalized_guardian_phone

    _require_valid_phone(normalized_phone)
    return normalized_phone, ""


def _get_duplicate_by_phone(*, club_id: int, phone: str) -> Student | None:
    if not phone:
        return None
    return (
        Student.objects.for_club(club_id)
        .filter(phone=phone, deleted_at__isnull=True)
        .order_by("id")
        .first()
    )


def _get_duplicate_child_by_guardian(
    *,
    club_id: int,
    first_name: str,
    last_name: str,
    guardian_phone: str,
    exclude_student_id: int | None = None,
) -> Student | None:
    if not guardian_phone:
        return None
    qs = (
        Student.objects.for_club(club_id)
        .filter(
            is_child=True,
            guardian_phone=guardian_phone,
            first_name__iexact=first_name.strip(),
            last_name__iexact=last_name.strip(),
            deleted_at__isnull=True,
        )
        .order_by("id")
    )
    if exclude_student_id is not None:
        qs = qs.exclude(id=exclude_student_id)
    return qs.first()


def create_student(
    *,
    club_id: int,
    first_name: str,
    last_name: str,
    phone: str,
    email: str = "",
    date_of_birth: date | None = None,
    is_child: bool = False,
    guardian_phone: str = "",
    source: str = "other",
    assigned_trainer_id: int | None = None,
) -> Student:
    from apps.trainers.models import Trainer

    identity = normalize_person_identity(
        first_name=first_name,
        last_name=last_name,
        date_of_birth=date_of_birth,
        is_child=is_child,
        phone=phone,
        guardian_phone=guardian_phone,
    )

    with transaction.atomic():
        lock_club_person_identity_arbitration(club_id=club_id)
        resolution = resolve_staff_intake_person_identity(
            club_id=club_id,
            identity=identity,
            confirm_distinct_child=False,
        )
        duplicate = resolution.duplicate or resolution.soft_deleted or resolution.ambiguous
        if duplicate is not None:
            raise DuplicateStudentError(duplicate, message="Ученик с таким телефоном уже существует")
        if assigned_trainer_id is not None and not Trainer.objects.for_club(club_id).filter(
            id=assigned_trainer_id,
        ).exists():
            raise BusinessLogicError("Тренер не найден в клубе", code="trainer_not_found")
        student = Student.objects.create(
            club_id=club_id,
            first_name=first_name,
            last_name=last_name,
            phone=identity.phone,
            guardian_phone=identity.guardian_phone,
            email=email,
            date_of_birth=date_of_birth,
            is_child=is_child,
            source=source,
            status=Student.Status.LEAD,
            lead_status=Student.LeadStatus.NEW,
            assigned_trainer_id=assigned_trainer_id,
        )
        if assigned_trainer_id is not None:
            from apps.leads.models import LeadLifecycleEvent

            LeadLifecycleEvent.objects.create(
                club_id=club_id,
                student=student,
                event_type=LeadLifecycleEvent.EventType.LEAD_ASSIGNED,
                old_lead_status="",
                new_lead_status=student.lead_status,
                old_trainer_id=None,
                new_trainer_id=assigned_trainer_id,
            )
    logger.info("student_created", extra={"student_id": student.id, "club_id": club_id})

    return student


_UPDATE_STUDENT_FIELDS = frozenset({
    "first_name",
    "last_name",
    "phone",
    "guardian_phone",
    "email",
    "date_of_birth",
    "is_child",
    "source",
    "contraindications",
})


def update_student(*, student_id: int, club_id: int, **fields) -> Student:
    bad = set(fields) - _UPDATE_STUDENT_FIELDS
    if bad:
        raise BusinessLogicError(f"Fields not allowed: {bad}", code="invalid_fields")

    student = Student.objects.for_club(club_id).get(id=student_id, deleted_at__isnull=True)

    if "phone" in fields:
        fields["phone"] = _normalize_optional_phone(fields["phone"])
    if "guardian_phone" in fields:
        fields["guardian_phone"] = _normalize_optional_phone(fields["guardian_phone"])
        if fields["guardian_phone"]:
            _require_valid_phone(fields["guardian_phone"], code="invalid_guardian_phone")

    next_is_child = fields.get("is_child", student.is_child)
    next_phone = fields.get("phone", student.phone)
    next_guardian_phone = fields.get("guardian_phone", student.guardian_phone)
    next_first_name = fields.get("first_name", student.first_name)
    next_last_name = fields.get("last_name", student.last_name)

    if next_is_child:
        if next_phone:
            _require_valid_phone(next_phone)
        if next_guardian_phone:
            _require_valid_phone(next_guardian_phone, code="invalid_guardian_phone")
        if not next_phone and not next_guardian_phone:
            raise BusinessLogicError(
                "Для ребёнка нужен телефон родителя или личный телефон",
                code="child_contact_required",
            )
        duplicate_child = find_duplicate_for_contact(
            club_id=club_id,
            first_name=next_first_name,
            last_name=next_last_name,
            is_child=True,
            phone=next_phone,
            guardian_phone=next_guardian_phone,
            exclude_student_id=student_id,
        )
        if duplicate_child is not None:
            raise DuplicateStudentError(
                duplicate_child,
                message="Ребёнок с таким именем уже есть у этого родителя",
            )
    else:
        _require_valid_phone(next_phone)
        fields["guardian_phone"] = ""
        duplicate = find_duplicate_for_contact(
            club_id=club_id,
            first_name=next_first_name,
            last_name=next_last_name,
            is_child=False,
            phone=next_phone,
            guardian_phone="",
            exclude_student_id=student_id,
        )
        if duplicate is not None:
            raise DuplicateStudentError(
                duplicate,
                message="Ученик с таким телефоном уже существует",
            )

    for attr, value in fields.items():
        setattr(student, attr, value)
    student.save(update_fields=[*fields.keys(), "updated_at"])
    return student


def delete_student(*, student_id: int, club_id: int) -> None:
    student = Student.objects.for_club(club_id).get(id=student_id, deleted_at__isnull=True)
    student.soft_delete()
    logger.info("student_deleted", extra={"student_id": student_id, "club_id": club_id})


def reactivate_student(*, student_id: int, club_id: int) -> Student | None:
    """T3: Auto-reactivate AT_RISK/CHURNED/LOST student on checkin.

    Bypasses the manual ALLOWED_TRANSITIONS FSM intentionally — this is a
    system event triggered by attendance, not a manual operator action.
    Sets ACTIVE if any active subscription exists, otherwise TRIAL.
    Returns None if student wasn't in a reactivatable state.
    """
    from django.db import transaction

    from apps.billing.models import Subscription

    inactive = (
        Student.Status.AT_RISK,
        Student.Status.CHURNED,
        Student.Status.LOST,
    )
    with transaction.atomic():
        student = (
            Student.objects.for_club(club_id)
            .select_for_update()
            .get(id=student_id)
        )
        if student.status not in inactive:
            return None

        old_status = student.status  # snapshot BEFORE mutation for logging
        has_active_sub = Subscription.objects.for_club(club_id).filter(
            student=student,
            status=Subscription.Status.ACTIVE,
            deleted_at__isnull=True,
        ).exists()
        new_status = (
            Student.Status.ACTIVE if has_active_sub else Student.Status.TRIAL
        )
        student.status = new_status
        student.save(update_fields=["status", "updated_at"])

    logger.info(
        "student_reactivated",
        extra={
            "student_id": student_id,
            "from": old_status,
            "to": new_status,
            "club_id": club_id,
        },
    )
    return student


def transition_status(
    *,
    student_id: int,
    club_id: int,
    new_status: str,
    actor_user_id: int | None = None,
    source: str = "student_status_transition",
) -> Student:
    should_finalize_lead_conversion = False
    with transaction.atomic():
        student = (
            Student.objects.for_club(club_id)
            .select_for_update()
            .get(id=student_id, deleted_at__isnull=True)
        )

        old_status = student.status
        old_lead_status = student.lead_status
        old_trainer_id = student.assigned_trainer_id
        allowed = ALLOWED_TRANSITIONS.get(student.status, [])
        if new_status not in allowed:
            current_label = dict(Student.Status.choices).get(student.status, student.status)
            new_label = dict(Student.Status.choices).get(new_status, new_status)
            raise BusinessLogicError(
                f"Нельзя перевести из «{current_label}» в «{new_label}»",
                code="invalid_transition",
            )

        student.status = new_status
        update_fields = ["status", "updated_at"]
        if new_status == Student.Status.LEAD:
            student.lead_status = Student.LeadStatus.NEW
        elif student.lead_status is not None:
            student.lead_status = None
        if student.lead_status != old_lead_status:
            update_fields.append("lead_status")
        student.save(update_fields=update_fields)

        if student.lead_status != old_lead_status:
            from apps.leads.models import LeadLifecycleEvent

            LeadLifecycleEvent.objects.create(
                club_id=club_id,
                student=student,
                actor_id=actor_user_id,
                event_type=LeadLifecycleEvent.EventType.STATUS_CHANGED,
                old_lead_status=old_lead_status or "",
                new_lead_status=student.lead_status or "",
                old_trainer_id=old_trainer_id,
                new_trainer_id=student.assigned_trainer_id,
                metadata={
                    "student_status_from": old_status,
                    "student_status_to": new_status,
                    "source": source,
                },
            )
        if old_status == Student.Status.TRIAL and new_status == Student.Status.ACTIVE:
            from apps.leads.models import LeadLifecycleEvent

            LeadLifecycleEvent.objects.create(
                club_id=club_id,
                student=student,
                actor_id=actor_user_id,
                event_type=LeadLifecycleEvent.EventType.LEAD_CONVERTED,
                old_lead_status=old_lead_status or "",
                new_lead_status="",
                old_trainer_id=old_trainer_id,
                new_trainer_id=student.assigned_trainer_id,
                metadata={
                    "student_status_from": old_status,
                    "student_status_to": new_status,
                    "source": source,
                },
            )
            should_finalize_lead_conversion = True
        if should_finalize_lead_conversion:
            from apps.leads.services import _finalize_lead_conversion_side_effects

            transaction.on_commit(
                lambda: _finalize_lead_conversion_side_effects(
                    club_id=club_id,
                    student_id=student_id,
                )
            )
    logger.info(
        "student_status_changed",
        extra={
            "student_id": student.id,
            "club_id": club_id,
            "new_status": new_status,
        },
    )
    return student


MAX_IMPORT_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_IMPORT_ROWS = 5000


def _parse_import_row(row: tuple) -> tuple[str, str, str]:
    """Return (name, normalized_phone, error) for a row. Error is '' on success."""
    name = str(row[0]).strip() if row[0] else ""
    phone = str(row[1]).strip() if len(row) > 1 and row[1] else ""
    if phone:
        phone = _normalize_phone(phone)
    if not name:
        return "", "", "empty name"
    if not phone:
        return "", "", "empty phone"
    return name, phone, ""


def import_students_from_excel(*, club_id: int, file) -> ImportResult:
    from openpyxl import load_workbook

    if hasattr(file, "size") and file.size > MAX_IMPORT_FILE_SIZE:
        raise BusinessLogicError("File too large (max 10MB)", code="file_too_large")

    wb = load_workbook(file, read_only=True)
    ws = wb.active
    skipped = 0
    errors: list[str] = []

    parsed_rows: list[tuple[int, str, str, str]] = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row_num - 1 > MAX_IMPORT_ROWS:
            wb.close()
            raise BusinessLogicError(f"Too many rows (max {MAX_IMPORT_ROWS})", code="too_many_rows")
        if not row or all(cell is None for cell in row):
            continue

        name, phone, err = _parse_import_row(row)
        if err:
            errors.append(f"Row {row_num}: {err}")
            continue

        parts = name.split(maxsplit=1)
        first_name = parts[0]
        last_name = parts[1] if len(parts) > 1 else ""
        parsed_rows.append((row_num, first_name, last_name, phone))

    wb.close()

    created = 0
    with transaction.atomic():
        for row_num, first_name, last_name, phone in parsed_rows:
            try:
                identity = normalize_person_identity(
                    first_name=first_name,
                    last_name=last_name,
                    is_child=False,
                    phone=phone,
                )
            except BusinessLogicError:
                errors.append(f"Row {row_num}: invalid phone")
                continue
            duplicate = arbitrate_person_identity(
                club_id=club_id,
                identity=identity,
                purpose="staff_create",
            )
            if duplicate is not None:
                skipped += 1
                continue
            Student.objects.create(
                club_id=club_id,
                first_name=identity.first_name,
                last_name=identity.last_name,
                phone=identity.phone,
                status=Student.Status.LEAD,
                lead_status=Student.LeadStatus.NEW,
            )
            created += 1
    logger.info(
        "excel_import_completed",
        extra={"club_id": club_id, "created_count": created, "skipped": skipped},
    )
    return ImportResult(created=created, skipped=skipped, errors=errors)


def add_student_note(*, club_id: int, student_id: int, author_id: int, text: str) -> StudentNote:
    Student.objects.for_club(club_id).get(id=student_id, deleted_at__isnull=True)
    return StudentNote.objects.create(
        club_id=club_id,
        student_id=student_id,
        author_id=author_id,
        text=text,
    )
