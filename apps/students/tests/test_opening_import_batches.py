from dataclasses import replace
from io import StringIO
from unittest.mock import patch

import pytest
from django.core.exceptions import ValidationError
from django.core.management import call_command

from apps.billing.models import OpeningEntitlementSnapshot, Payment
from apps.billing.tests.test_opening_issuer import command as opening_command
from apps.clubs.models import ClubMembership
from apps.common.exceptions import BusinessLogicError
from apps.students.imports.parsing import ENTITLEMENTS_SHEET, WorkbookRow
from apps.students.imports.runner import accept_batch, run_batch_chunk
from apps.students.imports.schemas import COLUMNS
from apps.students.imports.services import export_result, prepare_batch, update_draft_item, validate_batch
from apps.students.imports.storage import private_path
from apps.students.imports.workbooks import workbook_bytes
from apps.students.models import OpeningImportItemReceipt, Student
from apps.students.tests.factories import StudentFactory

pytestmark = pytest.mark.django_db
command = opening_command


def prepare(club, actor, terms, tmp_path, settings, extra_terms=()):
    settings.STUDENT_IMPORT_PRIVATE_ROOT = tmp_path / "private"
    rows = []
    for index, value in enumerate((terms, *extra_terms), start=2):
        payload = value.as_payload()
        values = {column: payload[field] for field, column in COLUMNS.items()}
        values.update({"Исходная цена": payload["paid_amount"], "Валюта": "RUB",
                       "Исторический долг": 0, "Заморожен": False, "Тип пакета": "Персональный"})
        rows.append(WorkbookRow(ENTITLEMENTS_SHEET, index, values))
    path = tmp_path / "input.xlsx"
    path.write_bytes(workbook_bytes(
        club_id=club.id, actor_user_id=actor.id, namespace=terms.source_namespace, rows=rows,
    ))
    return prepare_batch(club_id=club.id, actor_user_id=actor.id, path=path)


def accept(club, actor, batch):
    preview = validate_batch(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id)
    assert preview.selection, list(batch.items.values("errors"))
    accept_batch(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id, revision=preview.revision,
                 selection_token=str(preview.selection_token), channel="assistant_cli")
    return preview


def test_prepare_preview_are_private_and_have_no_domain_effects(club, command, tmp_path, settings):
    actor, terms = command
    batch = prepare(club, actor, terms, tmp_path, settings)
    preview = validate_batch(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id)
    assert len(preview.selection) == 1
    assert not Student.objects.for_club(club).exists()
    assert not Payment.objects.for_club(club).exists()
    assert private_path(name=batch.prepared_file).stat().st_mode & 0o777 == 0o600
    with pytest.raises(ValidationError):
        preview.save()


def test_apply_and_new_batch_replay_do_not_duplicate_financial_facts(club, command, tmp_path, settings):
    actor, terms = command
    batch = prepare(club, actor, terms, tmp_path, settings)
    accept(club, actor, batch)
    result = run_batch_chunk(club_id=club.id, batch_id=batch.id)
    assert result["counts"] == {"applied": 1}
    assert OpeningImportItemReceipt.objects.for_club(club).get().channel == "assistant_cli"
    settings.STUDENT_OPENING_IMPORT_ENABLED = False
    second = prepare(club, actor, terms, tmp_path, settings)
    accept(club, actor, second)
    assert run_batch_chunk(club_id=club.id, batch_id=second.id)["counts"] == {"replayed": 1}
    assert Payment.objects.for_club(club).count() == OpeningEntitlementSnapshot.objects.for_club(club).count() == 1


def test_preview_draft_edit_invalidates_old_revision(club, command, tmp_path, settings):
    actor, terms = command
    batch = prepare(club, actor, terms, tmp_path, settings)
    preview = validate_batch(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id)
    item = batch.items.get()
    values = {**item.source_data, "Комментарий": "Сверено повторно"}
    update_draft_item(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id, item_id=item.id, values=values)
    with pytest.raises(BusinessLogicError) as error:
        accept_batch(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id, revision=preview.revision,
                     selection_token=str(preview.selection_token), channel="assistant_cli")
    assert error.value.code == "import_preview_stale"


def test_two_packages_share_expected_own_student_creation(club, command, tmp_path, settings):
    actor, terms = command
    terms = replace(terms, distinct_payment_reference="source-one")
    second = replace(terms, entitlement_source_key="package-two", payment_source_key="payment-two",
                     distinct_payment_reference="source-two")
    batch = prepare(club, actor, terms, tmp_path, settings, extra_terms=(second,))
    accept(club, actor, batch)
    result = run_batch_chunk(club_id=club.id, batch_id=batch.id)
    assert result["counts"] == {"applied": 2}, list(batch.items.values("errors"))
    assert Student.objects.for_club(club).count() == 1
    assert Payment.objects.for_club(club).count() == 2


def test_receipt_failure_rolls_back_all_domain_effects(club, command, tmp_path, settings):
    actor, terms = command
    batch = prepare(club, actor, terms, tmp_path, settings)
    accept(club, actor, batch)
    with patch.object(OpeningImportItemReceipt.objects, "create", side_effect=RuntimeError("synthetic")):
        with pytest.raises(RuntimeError):
            run_batch_chunk(club_id=club.id, batch_id=batch.id)
    assert not Payment.objects.for_club(club).exists()
    assert not Student.objects.for_club(club).exists()
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["status"] == "completed"


def test_revoked_actor_cannot_run_or_read_replay(club, command, tmp_path, settings):
    actor, terms = command
    batch = prepare(club, actor, terms, tmp_path, settings)
    accept(club, actor, batch)
    ClubMembership.objects.filter(club=club, user=actor).update(is_active=False)
    with pytest.raises(BusinessLogicError) as error:
        run_batch_chunk(club_id=club.id, batch_id=batch.id)
    assert error.value.code == "actor_not_authorized"
    assert not Payment.objects.for_club(club).exists()


def test_external_student_change_after_preview_is_not_absorbed(club, command, tmp_path, settings):
    actor, terms = command
    student = StudentFactory(club=club, first_name=terms.first_name, last_name=terms.last_name,
                             phone=terms.phone, date_of_birth=None, status="active")
    terms = replace(terms, student_id=student.id)
    batch = prepare(club, actor, terms, tmp_path, settings)
    accept(club, actor, batch)
    # Simulates an existing writer which does not update updated_at.
    Student.objects.for_club(club).filter(id=student.id).update(status="at_risk")
    result = run_batch_chunk(club_id=club.id, batch_id=batch.id)
    assert result["counts"] == {"needs_review": 1}
    assert batch.items.get().errors[0]["code"] == "import_preview_stale"
    assert not Payment.objects.for_club(club).exists()


def test_cli_status_does_not_print_private_source_values(club, command, tmp_path, settings):
    actor, terms = command
    batch = prepare(club, actor, terms, tmp_path, settings)
    output = StringIO()
    call_command("student_opening_import", "validate", club_id=club.id, actor_user_id=actor.id,
                 batch_id=batch.id, stdout=output)
    assert terms.phone not in output.getvalue()
    assert terms.first_name not in output.getvalue()
    assert "selection_token" not in output.getvalue()
    assert '"ready": 1' in output.getvalue()


def test_partial_batch_can_correct_unapplied_row_without_editing_receipt(club, command, tmp_path, settings):
    actor, terms = command
    second = replace(terms, student_source_key="second-student", entitlement_source_key="second-package",
                     payment_source_key="second-payment", phone="+79990000222", first_name="Second")
    batch = prepare(club, actor, terms, tmp_path, settings, extra_terms=(second,))
    bad = batch.items.order_by("id").last()
    update_draft_item(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id,
                      item_id=bad.id, values={**bad.source_data, "Исторический долг": None})
    accept(club, actor, batch)
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["status"] == "partial"
    good = batch.items.order_by("id").first()
    with pytest.raises(BusinessLogicError):
        update_draft_item(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id,
                          item_id=good.id, values=good.source_data)
    update_draft_item(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id,
                      item_id=bad.id, values={**bad.source_data, "Исторический долг": 0})
    accept(club, actor, batch)
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["status"] == "completed"
    assert Payment.objects.for_club(club).count() == 2


@pytest.mark.django_db(transaction=True)
def test_postgres_duplicate_runners_wait_and_issue_once(club, command, tmp_path, settings):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    from time import monotonic, sleep

    from django.db import close_old_connections, connection, transaction

    from apps.clubs.models import Club

    if connection.vendor != "postgresql":
        pytest.skip("Requires independent PostgreSQL connections with observed blocking")
    actor, terms = command
    batch = prepare(club, actor, terms, tmp_path, settings)
    accept(club, actor, batch)
    locked, waiting, release = Event(), Event(), Event()
    pids = {}

    def run(first):
        close_old_connections()
        try:
            with transaction.atomic():
                with connection.cursor() as cursor:
                    cursor.execute("SELECT pg_backend_pid()")
                    pids[first] = cursor.fetchone()[0]
                if first:
                    Club.objects.select_for_update().get(id=club.id)
                    locked.set()
                    assert release.wait(15)
                else:
                    waiting.set()
                return run_batch_chunk(club_id=club.id, batch_id=batch.id)
        finally:
            close_old_connections()

    with ThreadPoolExecutor(max_workers=2) as executor:
        first = executor.submit(run, True)
        assert locked.wait(10)
        second = executor.submit(run, False)
        assert waiting.wait(10)
        try:
            observed, deadline = False, monotonic() + 10
            while monotonic() < deadline:
                with connection.cursor() as cursor:
                    cursor.execute("SELECT %s = ANY(pg_blocking_pids(%s))", [pids[True], pids[False]])
                    observed = cursor.fetchone()[0]
                if observed:
                    break
                sleep(0.02)
            assert observed, "Second runner must wait on the shared Club fence"
        finally:
            release.set()
        assert first.result(timeout=15)["status"] == second.result(timeout=15)["status"] == "completed"
    assert Payment.objects.for_club(club).count() == 1
    assert OpeningImportItemReceipt.objects.for_club(club).count() == 1


def test_old_runner_cannot_issue_after_partial_draft_revision_changes(club, command, tmp_path, settings):
    from apps.students.imports.runner import _apply_item

    actor, terms = command
    batch = prepare(club, actor, terms, tmp_path, settings)
    preview = accept(club, actor, batch)
    settings.STUDENT_OPENING_IMPORT_ENABLED = False
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["status"] == "partial"
    item = batch.items.get()
    update_draft_item(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id,
                      item_id=item.id, values={**item.source_data, "Комментарий": "Новая проверка"})
    settings.STUDENT_OPENING_IMPORT_ENABLED = True
    validate_batch(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id)
    with pytest.raises(BusinessLogicError) as error:
        _apply_item(club_id=club.id, batch_id=batch.id, preview_id=preview.id, entry=preview.selection[0])
    assert error.value.code == "import_runner_stale"
    assert not Payment.objects.for_club(club).exists()


def test_conflicting_student_key_blocks_every_affected_row(club, command, tmp_path, settings):
    actor, terms = command
    second = replace(terms, first_name="Other", entitlement_source_key="another", payment_source_key="another")
    batch = prepare(club, actor, terms, tmp_path, settings, extra_terms=(second,))
    preview = validate_batch(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id)
    assert preview.selection == []
    assert set(batch.items.values_list("status", flat=True)) == {"needs_review"}


def test_private_source_ttl_keeps_durable_financial_receipts(club, command, tmp_path, settings):
    import os
    from datetime import timedelta

    from django.utils import timezone

    from apps.students.imports.tasks import cleanup_expired_import_files

    actor, terms = command
    batch = prepare(club, actor, terms, tmp_path, settings)
    accept(club, actor, batch)
    run_batch_chunk(club_id=club.id, batch_id=batch.id)
    old = (timezone.now() - timedelta(days=8)).timestamp()
    for name in (batch.source_file, batch.prepared_file):
        os.utime(private_path(name=name), (old, old))
    assert cleanup_expired_import_files() == {"removed_files": 2}
    assert OpeningImportItemReceipt.objects.for_club(club).count() == 1
    assert OpeningEntitlementSnapshot.objects.for_club(club).count() == 1


def test_retention_task_is_registered_idempotently():
    from django_q.models import Schedule

    call_command("register_scheduled_tasks", stdout=StringIO())
    call_command("register_scheduled_tasks", stdout=StringIO())
    scheduled = Schedule.objects.get(name="cleanup_expired_student_import_files")
    assert scheduled.func == "apps.students.imports.tasks.cleanup_expired_import_files"
    assert scheduled.schedule_type == Schedule.HOURLY


def test_worker_processes_twenty_five_then_requeues_remaining(club, command, tmp_path, settings):
    from apps.students.imports.runner import run_batch_worker

    actor, terms = command
    additional = [replace(
        terms, student_source_key=f"student-{index}", entitlement_source_key=f"package-{index}",
        payment_source_key=f"payment-{index}", first_name=f"Synthetic{index}", phone=f"+7999000{index:04}",
    ) for index in range(26, 51)]
    batch = prepare(club, actor, terms, tmp_path, settings, extra_terms=additional)
    accept(club, actor, batch)
    with patch("django_q.tasks.async_task") as enqueue:
        first = run_batch_worker(club_id=club.id, batch_id=batch.id)
        assert first["counts"] == {"applied": 25, "ready": 1}
        enqueue.assert_called_once_with("apps.students.imports.runner.run_batch_worker",
                                        club_id=club.id, batch_id=batch.id)
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["counts"] == {"applied": 26}


def test_prepare_adult_without_dob_shares_generated_key_across_packages(club, command, tmp_path, settings):
    actor, terms = command
    terms = replace(terms, student_source_key="", distinct_payment_reference="first-source")
    second = replace(terms, entitlement_source_key="second-package", payment_source_key="second-payment",
                     distinct_payment_reference="second-source")
    batch = prepare(club, actor, terms, tmp_path, settings, extra_terms=(second,))
    assert len({item.source_data["student_key"] for item in batch.items.all()}) == 1
    accept(club, actor, batch)
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["counts"] == {"applied": 2}
    assert Student.objects.for_club(club).count() == 1


def test_export_preserves_generated_source_keys_and_card_link(club, command, tmp_path, settings):
    from openpyxl import load_workbook

    actor, terms = command
    terms = replace(terms, student_source_key="", entitlement_source_key="", payment_source_key="")
    batch = prepare(club, actor, terms, tmp_path, settings)
    accept(club, actor, batch)
    run_batch_chunk(club_id=club.id, batch_id=batch.id)
    name = export_result(club_id=club.id, actor_user_id=actor.id, batch_id=batch.id)
    book = load_workbook(private_path(name=name), read_only=True)
    try:
        assert book["Проверка"].cell(2, 5).value.endswith("/card/")
    finally:
        book.close()
    replay = prepare_batch(club_id=club.id, actor_user_id=actor.id, path=private_path(name=name))
    accept(club, actor, replay)
    assert run_batch_chunk(club_id=club.id, batch_id=replay.id)["counts"] == {"replayed": 1}


def test_other_club_cannot_read_or_validate_batch(club, command, tmp_path, settings):
    from apps.clubs.tests.factories import ClubFactory
    from apps.students.imports.selectors import batch_status

    actor, terms = command
    batch = prepare(club, actor, terms, tmp_path, settings)
    other = ClubFactory()
    ClubMembership.objects.create(club=other, user=actor, role="owner")
    for action in (batch_status, validate_batch):
        with pytest.raises(BusinessLogicError) as error:
            action(club_id=other.id, actor_user_id=actor.id, batch_id=batch.id)
        assert error.value.code == "target_not_available"


def test_private_cleanup_protects_applying_files_across_clubs(club, other_club, owner_user, tmp_path, settings):
    import os
    from datetime import timedelta

    from django.utils import timezone

    from apps.students.imports.storage import save_private
    from apps.students.imports.tasks import cleanup_expired_import_files
    from apps.students.models import OpeningImportBatch

    settings.STUDENT_IMPORT_PRIVATE_ROOT = tmp_path / "private"
    old = (timezone.now() - timedelta(days=8)).timestamp()
    protected = []
    for tenant in (club, other_club):
        name = save_private(content=b"synthetic active source", suffix="xlsx")
        os.utime(private_path(name=name), (old, old))
        OpeningImportBatch.objects.create(
            club=tenant, created_by=owner_user, source_namespace=f"ttl-{tenant.id}",
            source_file=name, status="applying", expires_at=timezone.now() + timedelta(days=1),
        )
        protected.append(name)
    expired = save_private(content=b"synthetic expired source", suffix="xlsx")
    os.utime(private_path(name=expired), (old, old))

    assert cleanup_expired_import_files() == {"removed_files": 1}
    assert all(private_path(name=name).is_file() for name in protected)
    assert not private_path(name=expired).exists()
    assert OpeningImportBatch.objects.for_club(club).count() == 1
    assert OpeningImportBatch.objects.for_club(other_club).count() == 1
