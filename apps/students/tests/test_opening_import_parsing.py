from datetime import date
from decimal import Decimal
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from apps.common.exceptions import BusinessLogicError
from apps.students.imports.parsing import (
    parse_date,
    parse_money,
    parse_phone,
    read_workbook,
    write_text,
)


def workbook(tmp_path, rows):
    path = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.title = "Абонементы"
    for row in rows:
        book.active.append(row)
    book.save(path)
    return path


def test_source_keys_and_dates_survive_workbook_read(tmp_path):
    path = workbook(tmp_path, [["student_key", "Начало"], ["opaque-key", date(2026, 9, 1)]])
    result = read_workbook(path=path)
    assert result.rows[0].values["student_key"] == "opaque-key"
    assert parse_date(result.rows[0].values["Начало"]) == date(2026, 9, 1)


def test_formula_rejected_without_private_value_in_error(tmp_path):
    path = workbook(tmp_path, [["Имя"], ['=HYPERLINK("https://private.invalid")']])
    with pytest.raises(BusinessLogicError) as error:
        read_workbook(path=path)
    assert "private.invalid" not in str(error.value)


def test_external_links_rejected(tmp_path):
    path = workbook(tmp_path, [["Имя"], ["Пример"]])
    with ZipFile(path, "a") as archive:
        archive.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
    with pytest.raises(BusinessLogicError):
        read_workbook(path=path)


def test_entitlement_limit_is_enforced(tmp_path):
    path = workbook(tmp_path, [["student_key"], *[[str(i)] for i in range(2001)]])
    with pytest.raises(BusinessLogicError, match="2 000"):
        read_workbook(path=path)


@pytest.mark.parametrize("value", [None, True, "NaN", "Infinity", "1.001", "9999999999999"])
def test_unknown_or_invalid_money_is_not_coerced(value):
    with pytest.raises(BusinessLogicError):
        parse_money(value)


def test_decimal_uses_source_decimal_representation():
    assert parse_money(6500.20) == Decimal("6500.20")
    assert parse_money("0") == Decimal("0.00")


def test_numeric_phone_requires_review():
    with pytest.raises(BusinessLogicError):
        parse_phone(79001234567)


@pytest.mark.parametrize("value", ["20260901", "01.09.2026", 46200, None])
def test_dates_are_not_guessed(value):
    with pytest.raises(BusinessLogicError):
        parse_date(value)


def test_exported_formula_like_source_remains_literal(tmp_path):
    path = tmp_path / "export.xlsx"
    book = Workbook()
    write_text(book.active.cell(1, 1), "=private-source")
    book.save(path)
    reopened = load_workbook(path, read_only=True)
    try:
        assert reopened.active.cell(1, 1).data_type == "s"
        assert reopened.active.cell(1, 1).value == "=private-source"
    finally:
        reopened.close()
