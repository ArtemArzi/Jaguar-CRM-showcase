import logging
import re
from pathlib import Path

from django.conf import settings
from django.core.paginator import Paginator
from django.db import IntegrityError
from django.db.models import Q
from django.http import FileResponse, Http404, HttpRequest, HttpResponse
from django.shortcuts import render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import content_disposition_header
from django.views.decorators.http import require_POST

from apps.attendance.selectors import get_student_attendance_summary
from apps.billing.models import Subscription
from apps.billing.selectors import get_student_subscriptions
from apps.common.exceptions import BusinessLogicError
from apps.common.permissions import management_view_required
from apps.common.phone import normalize_phone
from apps.documents.models import StudentDocument
from apps.documents.selectors import get_document_checklist, get_missing_documents_count
from apps.documents.services import mark_document_provided, upload_student_document
from apps.students.access_services import open_account_access_for_student, reset_account_access_for_student
from apps.students.duplicates import DuplicateStudentError, existing_contacts_for_club
from apps.students.models import AccountAccess, Student
from apps.students.parent_services import create_parent_invite
from apps.students.selectors import get_student_detail, get_students
from apps.students.services import (
    ALLOWED_TRANSITIONS,
    add_student_note,
    create_student,
    import_students_from_excel,
    transition_status,
    update_student,
)

PHONE_RE = re.compile(r"^\+?\d{11,15}$")

logger = logging.getLogger(__name__)


SAFE_DOCUMENT_CONTENT_TYPES = {
    ".pdf": "application/pdf",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}


def _allowed_status_choices(current_status: str) -> list[tuple[str, str]]:
    """Return only the status choices the student can transition to."""
    allowed_values = ALLOWED_TRANSITIONS.get(current_status, [])
    choices = dict(Student.Status.choices)
    return [(v, choices[v]) for v in allowed_values if v in choices]


STATUS_BADGE_COLORS = {
    "lead": "bg-gray-100 text-gray-800",
    "trial": "bg-blue-100 text-blue-800",
    "active": "bg-green-100 text-green-800",
    "at_risk": "bg-amber-100 text-amber-800",
    "churned": "bg-red-100 text-red-800",
    "lost": "bg-gray-300 text-gray-600",
}

PARENT_INVITE_ERROR_MESSAGES = {
    "not_child": "Родителя можно пригласить только для детской анкеты.",
    "parent_already_linked": "У ребёнка уже привязан родитель.",
}

ACCOUNT_ACCESS_ERROR_MESSAGES = {
    "account_access_not_open": "Сначала откройте доступ к личному кабинету.",
    "account_access_requires_active_student": "Кабинет открывается после оплаты и перевода ученика в активные.",
    "account_access_requires_paid_subscription": "Кабинет открывается только после активной оплаченной подписки.",
    "child_requires_parent_access": "Для детской анкеты доступ открывается родителю.",
    "invalid_phone": "Проверьте телефон в анкете: по нему создаётся логин.",
    "manual_review_required": "Номер уже связан с другим пользователем. Нужна ручная проверка.",
    "parent_phone_required": "Для детской анкеты укажите телефон родителя.",
    "phone_required": "В анкете нужен телефон: по нему создаётся логин.",
    "student_not_accessible": "Для потерянного ученика доступ открыть нельзя.",
    "student_not_found": "Ученик не найден.",
}


def _parent_invite_error_message(code: str | None) -> str:
    return PARENT_INVITE_ERROR_MESSAGES.get(
        code or "",
        "Не удалось создать приглашение. Проверьте анкету ребёнка и попробуйте ещё раз.",
    )


def _account_access_error_message(code: str | None) -> str:
    return ACCOUNT_ACCESS_ERROR_MESSAGES.get(
        code or "",
        "Не удалось изменить доступ. Проверьте анкету и попробуйте ещё раз.",
    )


def _account_access_for_student(*, club, student: Student) -> AccountAccess | None:
    role = AccountAccess.Role.PARENT if student.is_child else AccountAccess.Role.STUDENT
    return (
        AccountAccess.objects.for_club(club)
        .select_related("user")
        .filter(student=student, role=role)
        .first()
    )


def _has_paid_active_subscription(subscriptions) -> bool:
    now = timezone.now()
    return any(
        sub.status == Subscription.Status.ACTIVE
        and sub.paid_amount is not None
        and sub.paid_amount > 0
        and (sub.expires_at is None or sub.expires_at > now)
        and (sub.trainings_left is None or sub.trainings_left > 0)
        for sub in subscriptions
    )


def _subscription_card_item(*, subscription: Subscription, now) -> dict:
    trainings_limit = subscription.tariff.trainings_limit
    is_single_credit = trainings_limit == 1
    if not is_single_credit:
        remaining_text = None
        if subscription.trainings_left is not None and trainings_limit is not None:
            remaining_text = f"Осталось {subscription.trainings_left} из {trainings_limit}"
        return {
            "subscription": subscription,
            "is_single_credit": False,
            "is_current_entitlement": True,
            "kind_label": "АБОНЕМЕНТ",
            "state_label": "",
            "remaining_text": remaining_text,
        }

    if subscription.trainings_used >= 1:
        state_label = "ПОСЕЩЕНА"
    elif subscription.status == Subscription.Status.CANCELLED:
        state_label = "ОТМЕНЕНА"
    elif subscription.status == Subscription.Status.PENDING:
        state_label = "ОЖИДАЕТ ПОДТВЕРЖДЕНИЯ"
    elif subscription.status == Subscription.Status.FROZEN:
        state_label = "ЗАМОРОЖЕНА"
    elif subscription.status == Subscription.Status.EXPIRED or (
        subscription.status == Subscription.Status.ACTIVE
        and subscription.expires_at is not None
        and subscription.expires_at <= now
    ):
        state_label = "ИСТЕКЛА"
    elif (
        subscription.status == Subscription.Status.ACTIVE
        and subscription.trainings_left is not None
        and subscription.trainings_left > 0
    ):
        state_label = "ДОСТУПНА"
    else:
        state_label = "НЕДОСТУПНА"

    return {
        "subscription": subscription,
        "is_single_credit": True,
        "is_current_entitlement": state_label
        in {
            "ДОСТУПНА",
            "ОЖИДАЕТ ПОДТВЕРЖДЕНИЯ",
            "ЗАМОРОЖЕНА",
        },
        "kind_label": "РАЗОВАЯ ТРЕНИРОВКА",
        "state_label": state_label,
        "remaining_text": None,
    }


def _student_card_context(*, club, student) -> dict:
    """Context for the student card partial — used by card, edit, status, note views."""
    subscriptions = list(get_student_subscriptions(club=club, student_id=student.id))
    now = timezone.now()
    subscription_cards = []
    for subscription in subscriptions:
        card = _subscription_card_item(subscription=subscription, now=now)
        if card["is_current_entitlement"]:
            subscription_cards.append(card)
    attendance_summary = get_student_attendance_summary(
        club=club,
        student_id=student.id,
    )
    account_access = _account_access_for_student(club=club, student=student)
    has_paid_active_subscription = _has_paid_active_subscription(subscriptions)
    document_checklist = []
    for item in get_document_checklist(club=club, student_id=student.id):
        student_document = item.get("student_document")
        enriched = dict(item)
        if student_document and student_document.id:
            enriched["open_url"] = reverse(
                "student-document-open",
                kwargs={"student_id": student.id, "student_document_id": student_document.id},
            )
        document_checklist.append(enriched)
    from apps.billing.models import PaymentRefund
    from apps.students.operation_selectors import get_student_operation_sections

    existing_refunds = (PaymentRefund.objects.for_club(club).filter(payment__student=student)
                        .select_related("payment").order_by("-id"))
    operations_enabled = bool(settings.STUDENT_ADMIN_CORRECTIONS_ENABLED)
    return {
        "student_operations_enabled": operations_enabled,
        "existing_refunds": existing_refunds,
        **(get_student_operation_sections(club=club, student=student, subscriptions=subscriptions)
           if operations_enabled else {}),
        "student": student,
        "subscriptions": subscriptions,
        "subscription_cards": subscription_cards,
        "attendance_summary": attendance_summary,
        "document_checklist": document_checklist,
        "missing_documents_count": get_missing_documents_count(club=club, student_id=student.id),
        "documents_target": "#student-documents-block",
        "documents_mark_url": reverse("student-document-mark", kwargs={"student_id": student.id}),
        "documents_upload_url": reverse("student-document-upload", kwargs={"student_id": student.id}),
        "statuses": _allowed_status_choices(student.status),
        "badge_colors": STATUS_BADGE_COLORS,
        "account_access": account_access,
        "account_access_role_label": "родителя" if student.is_child else "ученика",
        "has_paid_active_subscription": has_paid_active_subscription,
        "can_issue_account_access": (
            account_access is None
            and student.status == Student.Status.ACTIVE
            and has_paid_active_subscription
        ),
    }


def _render_account_access_panel(
    *,
    request: HttpRequest,
    student_id: int,
    temporary_password: str | None = None,
    account_access_error: str | None = None,
    status: int = 200,
) -> HttpResponse:
    try:
        student = get_student_detail(club=request.club, student_id=student_id)
    except Student.DoesNotExist as exc:
        raise Http404 from exc
    context = _student_card_context(club=request.club, student=student)
    context["temporary_password"] = temporary_password
    context["account_access_error"] = account_access_error
    return render(request, "dashboard/students/_account_access.html", context, status=status)


def _render_student_documents_block(
    *,
    request: HttpRequest,
    student_id: int,
    status: int = 200,
    error: str | None = None,
) -> HttpResponse:
    try:
        student = get_student_detail(club=request.club, student_id=student_id)
    except Student.DoesNotExist as exc:
        raise Http404 from exc
    context = _student_card_context(club=request.club, student=student)
    if error:
        context["documents_error"] = error
    return render(request, "dashboard/students/_card_documents.html", context, status=status)


@management_view_required
def student_create(request: HttpRequest) -> HttpResponse:
    sources = Student.Source.choices

    if request.method == "GET":
        return render(request, "dashboard/students/_create_form.html", {"sources": sources})

    first_name = (request.POST.get("first_name") or "").strip()
    last_name = (request.POST.get("last_name") or "").strip()
    raw_phone = (request.POST.get("phone") or "").strip()
    # Normalize: strip spaces, dashes, parens → only digits and leading +
    phone = re.sub(r"[\s\-\(\)]", "", raw_phone)
    source = request.POST.get("source", "other")
    is_child = request.POST.get("is_child") == "1"

    note = (request.POST.get("note") or "").strip()

    form_data = {"first_name": first_name, "last_name": last_name, "phone": raw_phone, "note": note}

    if not first_name:
        return render(request, "dashboard/students/_create_form.html", {
            "sources": sources, "error": "Имя обязательно", "form_data": form_data,
        })

    if not phone or not PHONE_RE.match(phone):
        return render(request, "dashboard/students/_create_form.html", {
            "sources": sources,
            "error": "Введите корректный телефон (11-15 цифр, например +79031234567)",
            "form_data": form_data,
        })

    try:
        student = create_student(
            club_id=request.club.id,
            first_name=first_name,
            last_name=last_name,
            phone=phone,
            source=source,
            is_child=is_child,
        )
        if note:
            add_student_note(club_id=request.club.id, student_id=student.id, text=note, author_id=request.user.id)
    except BusinessLogicError as e:
        return render(request, "dashboard/students/_create_form.html", {
            "sources": sources, "error": e.message or str(e), "form_data": form_data,
        })

    # Success — re-render list and close slide-over
    from django.core.paginator import Paginator as _Paginator
    students = get_students(club=request.club)
    paginator = _Paginator(students, 20)
    context = {
        "page_title": "Ученики",
        "page_obj": paginator.get_page(1),
        "statuses": Student.Status.choices,
        "current_status": "",
        "search_query": "",
        "badge_colors": STATUS_BADGE_COLORS,
    }
    response = render(request, "dashboard/students/list.html#content", context)
    response["HX-Retarget"] = "#content"
    response["HX-Trigger"] = "closeSlideOver"
    return response


@management_view_required
def student_list(request: HttpRequest) -> HttpResponse:
    q = request.GET.get("q", "").strip()
    status = request.GET.get("status", "")
    page_num = request.GET.get("page", "1")

    students = get_students(club=request.club)

    if q:
        students = students.filter(
            Q(first_name__icontains=q) | Q(last_name__icontains=q) | Q(phone__icontains=q)
        )
    if status:
        students = students.filter(status=status)

    paginator = Paginator(students, 20)
    page_obj = paginator.get_page(page_num)

    context = {
        "page_title": "Ученики",
        "page_obj": page_obj,
        "statuses": Student.Status.choices,
        "current_status": status,
        "search_query": q,
        "badge_colors": STATUS_BADGE_COLORS,
    }
    if request.htmx:
        return render(request, "dashboard/students/list.html#content", context)
    return render(request, "dashboard/students/list.html", context)


@management_view_required
def student_card(request: HttpRequest, student_id: int) -> HttpResponse:
    try:
        student = get_student_detail(club=request.club, student_id=student_id)
    except Student.DoesNotExist:
        raise Http404
    return render(request, "dashboard/students/_card.html", _student_card_context(club=request.club, student=student))


@management_view_required
@require_POST
def student_parent_invite(request: HttpRequest, student_id: int) -> HttpResponse:
    try:
        student = Student.objects.for_club(request.club).get(id=student_id, deleted_at__isnull=True)
    except Student.DoesNotExist as exc:
        raise Http404 from exc

    try:
        invite = create_parent_invite(club_id=request.club.id, student_id=student.id)
    except BusinessLogicError as exc:
        return render(
            request,
            "dashboard/students/_parent_invite.html",
            {
                "student": student,
                "invite_error": _parent_invite_error_message(exc.code),
            },
            status=422,
        )

    invite_token = str(invite.token)
    return render(
        request,
        "dashboard/students/_parent_invite.html",
        {
            "student": student,
            "invite_url": request.build_absolute_uri(f"/parent-invite/{invite_token}"),
            "invite_token": invite_token,
            "expires_at": invite.expires_at,
        },
    )


@management_view_required
@require_POST
def student_account_access_open(request: HttpRequest, student_id: int) -> HttpResponse:
    try:
        result = open_account_access_for_student(
            club_id=request.club.id,
            student_id=student_id,
            parent_phone=(request.POST.get("parent_phone") or "").strip() or None,
            issued_by_id=request.user.id,
        )
    except BusinessLogicError as exc:
        return _render_account_access_panel(
            request=request,
            student_id=student_id,
            account_access_error=_account_access_error_message(exc.code),
        )

    return _render_account_access_panel(
        request=request,
        student_id=student_id,
        temporary_password=result.temporary_password,
    )


@management_view_required
@require_POST
def student_account_access_reset(request: HttpRequest, student_id: int) -> HttpResponse:
    try:
        result = reset_account_access_for_student(
            club_id=request.club.id,
            student_id=student_id,
            reset_by_id=request.user.id,
        )
    except BusinessLogicError as exc:
        return _render_account_access_panel(
            request=request,
            student_id=student_id,
            account_access_error=_account_access_error_message(exc.code),
        )

    return _render_account_access_panel(
        request=request,
        student_id=student_id,
        temporary_password=result.temporary_password,
    )


# --- Excel Import ---


@management_view_required
def import_upload(request: HttpRequest) -> HttpResponse:
    if request.method == "GET":
        return render(request, "dashboard/students/_import_form.html")

    # POST: parse Excel and show preview
    file = request.FILES.get("file")
    if not file:
        return render(request, "dashboard/students/_import_form.html", {"error": "Выберите файл"})

    # Validate size and content-type before openpyxl parses the file
    # (defence against zip-bombs and accidental wrong-file uploads).
    if file.size > 5 * 1024 * 1024:
        return render(request, "dashboard/students/_import_form.html",
                      {"error": "Файл слишком большой. Максимум 5 МБ."})
    valid_xlsx_types = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    )
    if file.content_type not in valid_xlsx_types and not file.name.lower().endswith(".xlsx"):
        return render(request, "dashboard/students/_import_form.html",
                      {"error": "Файл должен быть .xlsx"})

    from openpyxl import load_workbook

    try:
        wb = load_workbook(file, read_only=True)
    except Exception:
        return render(request, "dashboard/students/_import_form.html", {"error": "Некорректный Excel-файл"})

    ws = wb.active
    rows = []
    existing_contacts = existing_contacts_for_club(club_id=request.club.id)
    seen_phones: set[str] = set()
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(cell is None for cell in row):
            continue
        name = str(row[0]).strip() if row[0] else ""
        raw_phone = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        phone = normalize_phone(raw_phone) if raw_phone else ""
        row_errors: list[str] = []
        if not name:
            row_errors.append("Пустое имя")
        if not phone:
            row_errors.append("Пустой телефон")
        elif phone in existing_contacts:
            row_errors.append("Телефон уже существует")
        elif phone in seen_phones:
            row_errors.append("Дублирующийся телефон в файле")

        if phone:
            seen_phones.add(phone)
        rows.append({"row_num": row_num, "name": name, "phone": phone, "errors": row_errors})
    wb.close()

    # Save file to temp storage via session
    from django.core.files.storage import default_storage

    file.seek(0)
    temp_path = default_storage.save(f"tmp/import_{request.club.id}_{file.name}", file)
    request.session["import_file_path"] = temp_path

    errors_count = sum(1 for r in rows if r["errors"])
    valid_count = len(rows) - errors_count
    context = {
        "rows": rows,
        "total": len(rows),
        "errors_count": errors_count,
        "valid_count": valid_count,
    }
    return render(request, "dashboard/students/_import_preview.html", context)


@management_view_required
def import_confirm(request: HttpRequest) -> HttpResponse:
    temp_path = request.session.pop("import_file_path", None)
    if not temp_path:
        return HttpResponse(status=204, headers={"HX-Redirect": "/dashboard/students/"})

    from django.core.files.storage import default_storage

    try:
        with default_storage.open(temp_path, "rb") as f:
            result = import_students_from_excel(club_id=request.club.id, file=f)
        logger.info(
            "import_confirmed",
            extra={"club_id": request.club.id, "created_count": result.created, "skipped": result.skipped},
        )
    finally:
        try:
            default_storage.delete(temp_path)
        except Exception:
            pass

    return HttpResponse(status=204, headers={"HX-Redirect": "/dashboard/students/"})


@management_view_required
def student_add_note(request: HttpRequest, student_id: int) -> HttpResponse:
    text = request.POST.get("text", "").strip()
    if text:
        try:
            add_student_note(
                club_id=request.club.id,
                student_id=student_id,
                author_id=request.user.id,
                text=text,
            )
        except Student.DoesNotExist:
            raise Http404

    try:
        student = get_student_detail(club=request.club, student_id=student_id)
    except Student.DoesNotExist:
        raise Http404
    return render(request, "dashboard/students/_card.html", _student_card_context(club=request.club, student=student))


@management_view_required
def student_edit(request: HttpRequest, student_id: int) -> HttpResponse:
    try:
        student = get_student_detail(club=request.club, student_id=student_id)
    except Student.DoesNotExist:
        raise Http404

    if request.method == "POST":
        raw_phone = request.POST.get("phone", student.phone).strip()
        raw_guardian_phone = request.POST.get("guardian_phone", student.guardian_phone).strip()
        phone = re.sub(r"[\s\-\(\)]", "", raw_phone)
        guardian_phone = re.sub(r"[\s\-\(\)]", "", raw_guardian_phone)
        source = request.POST.get("source", "")
        fields = {
            "first_name": request.POST.get("first_name", student.first_name).strip(),
            "last_name": request.POST.get("last_name", student.last_name).strip(),
            "email": request.POST.get("email", student.email).strip(),
            "contraindications": request.POST.get("contraindications", student.contraindications).strip(),
            "is_child": student.is_child,
        }
        if source:
            fields["source"] = source
        if student.is_child:
            fields["guardian_phone"] = guardian_phone
            if phone:
                fields["phone"] = phone
        else:
            fields["phone"] = phone
        try:
            student = update_student(
                student_id=student_id,
                club_id=request.club.id,
                **fields,
            )
        except (DuplicateStudentError, BusinessLogicError) as e:
            return render(request, "dashboard/students/_edit_form.html", {
                "student": student,
                "sources": Student.Source.choices,
                "errors": [e.message or str(e)],
            })
        except IntegrityError:
            return render(request, "dashboard/students/_edit_form.html", {
                "student": student,
                "sources": Student.Source.choices,
                "errors": ["Контакт уже есть в CRM. Откройте существующую карточку."],
            })
        logger.info("student_edited", extra={"student_id": student_id, "club_id": request.club.id})

        return render(
            request,
            "dashboard/students/_card.html",
            _student_card_context(club=request.club, student=student),
        )

    context = {
        "student": student,
        "sources": Student.Source.choices,
    }
    return render(request, "dashboard/students/_edit_form.html", context)


@management_view_required
def student_status_change(request: HttpRequest, student_id: int) -> HttpResponse:
    new_status = request.POST.get("new_status", "")
    error = None
    changed = False

    if new_status:
        try:
            transition_status(
                student_id=student_id,
                club_id=request.club.id,
                new_status=new_status,
                actor_user_id=request.user.id,
                source="htmx_admin_student_status",
            )
            changed = True
        except Student.DoesNotExist:
            raise Http404
        except BusinessLogicError as exc:
            error = str(exc)

    try:
        student = get_student_detail(club=request.club, student_id=student_id)
    except Student.DoesNotExist:
        raise Http404
    context = _student_card_context(club=request.club, student=student)
    context["status_error"] = error
    response = render(request, "dashboard/students/_card.html", context)
    if changed:
        response["HX-Trigger"] = "studentUpdated"
    return response


@management_view_required
@require_POST
def student_document_mark(request: HttpRequest, student_id: int) -> HttpResponse:
    document_type_id = request.POST.get("document_type_id", "").strip()
    is_provided_raw = request.POST.get("is_provided", "true").strip().lower()
    notes = request.POST.get("notes", "").strip()

    if not document_type_id.isdigit():
        return _render_student_documents_block(
            request=request, student_id=student_id, status=422, error="Некорректный документ."
        )

    try:
        mark_document_provided(
            club_id=request.club.id,
            student_id=student_id,
            document_type_id=int(document_type_id),
            is_provided=is_provided_raw in {"1", "true", "on", "yes"},
            notes=notes,
        )
    except Student.DoesNotExist:
        raise Http404
    except BusinessLogicError as exc:
        return _render_student_documents_block(request=request, student_id=student_id, status=422, error=str(exc))

    return _render_student_documents_block(request=request, student_id=student_id)


@management_view_required
@require_POST
def student_document_upload(request: HttpRequest, student_id: int) -> HttpResponse:
    document_type_id = request.POST.get("document_type_id", "").strip()
    file = request.FILES.get("file")

    if not document_type_id.isdigit():
        return _render_student_documents_block(
            request=request, student_id=student_id, status=422, error="Некорректный документ."
        )
    if file is None:
        return _render_student_documents_block(
            request=request, student_id=student_id, status=422, error="Выберите файл для загрузки."
        )

    try:
        upload_student_document(
            club_id=request.club.id,
            student_id=student_id,
            document_type_id=int(document_type_id),
            file=file,
        )
    except Student.DoesNotExist:
        raise Http404
    except BusinessLogicError as exc:
        return _render_student_documents_block(request=request, student_id=student_id, status=422, error=str(exc))

    return _render_student_documents_block(request=request, student_id=student_id)


@management_view_required
def student_document_open(request: HttpRequest, student_id: int, student_document_id: int) -> HttpResponse:
    try:
        student = Student.objects.for_club(request.club).get(id=student_id, deleted_at__isnull=True)
    except Student.DoesNotExist:
        raise Http404

    student_document = (
        StudentDocument.objects.for_club(request.club)
        .select_related("student", "document_type")
        .filter(id=student_document_id, student_id=student.id, deleted_at__isnull=True)
        .first()
    )
    if not student_document or not student_document.file:
        raise Http404

    filename = Path(student_document.file.name).name
    extension = Path(filename).suffix.lower()
    response = FileResponse(
        student_document.file.open("rb"),
        content_type=SAFE_DOCUMENT_CONTENT_TYPES.get(extension, "application/octet-stream"),
    )
    response["Content-Disposition"] = content_disposition_header(as_attachment=True, filename=filename) or "attachment"
    response["X-Content-Type-Options"] = "nosniff"
    return response
