import io
from datetime import date
from decimal import Decimal, InvalidOperation
from urllib.parse import urlencode

from django.core.exceptions import ValidationError
from django.db.models import Sum
from django.http import Http404, HttpRequest, HttpResponse
from django.shortcuts import render

from apps.billing.models import Debt, Expense
from apps.billing.services import create_expense, delete_expense
from apps.common.permissions import management_view_required
from apps.dashboard.services import (
    get_attendance_metrics,
    get_business_metrics,
    get_dormant_count,
    get_pnl_report,
    get_trial_conversion,
)


def _parse_report_date_values(date_from_str: str, date_to_str: str) -> tuple[date, date]:
    today = date.today()
    try:
        report_from = date.fromisoformat(date_from_str) if date_from_str else today.replace(day=1)
    except ValueError:
        report_from = today.replace(day=1)
    try:
        report_to = date.fromisoformat(date_to_str) if date_to_str else today
    except ValueError:
        report_to = today
    if report_from > report_to:
        report_from, report_to = report_to, report_from
    return report_from, report_to


def _parse_report_dates(request: HttpRequest) -> tuple[date, date]:
    return _parse_report_date_values(
        request.GET.get("date_from", ""),
        request.GET.get("date_to", ""),
    )


def _report_redirect_url(*, date_from: str, date_to: str) -> str:
    report_from, report_to = _parse_report_date_values(date_from, date_to)
    query = urlencode({"date_from": report_from.isoformat(), "date_to": report_to.isoformat()})
    return f"/dashboard/reports/?{query}"


@management_view_required
def pnl_report(request: HttpRequest) -> HttpResponse:
    report_from, report_to = _parse_report_dates(request)

    pnl = get_pnl_report(club=request.club, date_from=report_from, date_to=report_to)
    metrics = get_business_metrics(club=request.club, date_from=report_from, date_to=report_to)
    attendance = get_attendance_metrics(club=request.club, date_from=report_from, date_to=report_to)
    trial = get_trial_conversion(club=request.club, date_from=report_from, date_to=report_to)
    dormant = get_dormant_count(club=request.club)

    total_debt = (
        Debt.objects.for_club(request.club)
        .filter(resolved_at__isnull=True, settlement_payment__isnull=True)
        .aggregate(total=Sum("tariff_price"))["total"]
        or 0
    )
    context = {
        "page_title": "P&L Report",
        "pnl": pnl,
        "metrics": metrics,
        "attendance": attendance,
        "trial": trial,
        "dormant": dormant,
        "total_debt": total_debt,
        "date_from": report_from.isoformat(),
        "date_to": report_to.isoformat(),
    }
    if request.htmx:
        return render(request, "dashboard/reports/pnl.html#content", context)
    return render(request, "dashboard/reports/pnl.html", context)


@management_view_required
def expense_create(request: HttpRequest) -> HttpResponse:
    """GET: show create form in slide-over. POST: create expense."""
    if request.method == "GET":
        report_from, report_to = _parse_report_dates(request)
        context = {
            "today": date.today().isoformat(),
            "date_from": report_from.isoformat(),
            "date_to": report_to.isoformat(),
        }
        return render(request, "dashboard/reports/_expense_form.html", context)

    # POST
    name = request.POST.get("name", "").strip()
    amount_str = request.POST.get("amount", "").strip()
    expense_date = request.POST.get("date", "").strip()
    category = request.POST.get("category", "").strip()
    is_recurring = request.POST.get("is_recurring") == "on"

    error = None
    if not name:
        error = "Укажите название"

    amount = None
    if not error:
        try:
            amount = Decimal(amount_str)
            if amount <= 0:
                error = "Сумма должна быть больше 0"
        except (InvalidOperation, ValueError):
            error = "Некорректная сумма"

    parsed_date = None
    if not error:
        try:
            parsed_date = date.fromisoformat(expense_date)
        except ValueError:
            error = "Некорректная дата"

    if error:
        context = {
            "error": error,
            "today": date.today().isoformat(),
            "date_from": request.POST.get("date_from", ""),
            "date_to": request.POST.get("date_to", ""),
        }
        return render(request, "dashboard/reports/_expense_form.html", context)

    try:
        create_expense(
            club_id=request.club.id,
            name=name,
            amount=amount,
            date=parsed_date,
            category=category,
            is_recurring=is_recurring,
        )
    except ValidationError as e:
        context = {
            "error": "; ".join(e.messages),
            "today": date.today().isoformat(),
            "date_from": request.POST.get("date_from", ""),
            "date_to": request.POST.get("date_to", ""),
        }
        return render(request, "dashboard/reports/_expense_form.html", context)

    redirect_url = _report_redirect_url(
        date_from=request.POST.get("date_from", ""),
        date_to=request.POST.get("date_to", ""),
    )
    return HttpResponse(status=204, headers={"HX-Redirect": redirect_url})


@management_view_required
def expense_delete(request: HttpRequest, expense_id: int) -> HttpResponse:
    """Soft-delete an expense."""
    if request.method != "POST":
        return HttpResponse(status=405, headers={"Allow": "POST"})

    try:
        delete_expense(expense_id=expense_id, club_id=request.club.id)
    except Expense.DoesNotExist:
        raise Http404

    redirect_url = _report_redirect_url(
        date_from=request.GET.get("date_from", ""),
        date_to=request.GET.get("date_to", ""),
    )
    return HttpResponse(status=204, headers={"HX-Redirect": redirect_url})


@management_view_required
def pnl_export_excel(request: HttpRequest) -> HttpResponse:
    """Export P&L report as Excel file."""
    import openpyxl  # lazy import — heavy optional dependency
    from openpyxl.styles import Font

    report_from, report_to = _parse_report_dates(request)

    pnl = get_pnl_report(club=request.club, date_from=report_from, date_to=report_to)
    metrics = get_business_metrics(club=request.club, date_from=report_from, date_to=report_to)
    attendance = get_attendance_metrics(club=request.club, date_from=report_from, date_to=report_to)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "P&L"

    bold = Font(bold=True)

    # Title
    ws.append([f"Финансовый отчёт: {report_from.strftime('%d.%m.%Y')} — {report_to.strftime('%d.%m.%Y')}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])

    # Summary
    ws.append(["СВОДКА"])
    ws[f"A{ws.max_row}"].font = bold
    ws.append(["Валовые оплаты", float(pnl["gross_income"])])
    ws.append(["Возвраты", float(pnl["refunded_income"])])
    ws.append(["Доход после возвратов", float(pnl["income"])])
    ws.append(["Зарплаты тренеров", float(pnl["salary_expenses"])])
    ws.append(["Прочие расходы", float(pnl["manual_expenses"])])
    ws.append(["Итого расходы", float(pnl["total_expenses"])])
    ws.append(["Маржа", float(pnl["margin"])])
    ws.append(["Маржа %", pnl["margin_percent"]])
    ws.append([])

    # Income breakdown
    ws.append(["РАЗБИВКА ДОХОДА"])
    ws[f"A{ws.max_row}"].font = bold
    for item in pnl["income_breakdown"]:
        ws.append([item["label"], float(item["amount"])])
    ws.append([])

    # Salary breakdown
    ws.append(["ЗАРПЛАТЫ ТРЕНЕРОВ"])
    ws[f"A{ws.max_row}"].font = bold
    for item in pnl["salary_breakdown"]:
        ws.append([item["trainer"], float(item["amount"])])
    ws.append([])

    # Expenses
    ws.append(["РАСХОДЫ"])
    ws[f"A{ws.max_row}"].font = bold
    ws.append(["Название", "Сумма", "Повтор", "Дата"])
    for item in pnl["expenses_breakdown"]:
        ws.append([
            item["name"],
            float(item["amount"]),
            "Да" if item["is_recurring"] else "Нет",
            item["date"].strftime("%d.%m.%Y") if item["date"] else "",
        ])
    ws.append([])

    # Metrics
    ws.append(["КЛЮЧЕВЫЕ МЕТРИКИ"])
    ws[f"A{ws.max_row}"].font = bold
    ws.append(["Отток", f"{metrics['churn_rate']}%" if metrics["churn_rate"] is not None else "—"])
    ws.append(["Удержание", f"{metrics['retention_rate']}%" if metrics["retention_rate"] is not None else "—"])
    ws.append(["ARPM", float(metrics["arpm"]) if metrics["arpm"] is not None else "—"])
    ws.append(["LTV", float(metrics["ltv"]) if metrics["ltv"] is not None else "—"])
    ws.append(["Посещаемость", attendance["total_checkins"]])
    ws.append(["Среднее в день", attendance["avg_per_day"]])

    # Column widths
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 12

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"pnl_{report_from.strftime('%Y%m%d')}_{report_to.strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
