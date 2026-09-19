import io
import logging
import re
from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal
from html.parser import HTMLParser
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from django.contrib.staticfiles import finders
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.utils import timezone
from openpyxl import Workbook

from apps.attendance.models import (
    Checkin,
    GroupSession,
    PersonalAvailabilitySlot,
    PersonalDropInBooking,
    Schedule,
    ScheduleEnrollment,
    ScheduleException,
)
from apps.attendance.services import (
    book_personal_drop_in,
    create_personal_booking_payment_reservation,
    create_personal_drop_in_payment,
)
from apps.attendance.tests.factories import CheckinFactory, ScheduleExceptionFactory, ScheduleFactory
from apps.attendance.tests.rollout import update_training_group_rollout_state_for_test
from apps.billing.models import (
    BankPaymentOrder,
    BankPaymentOrderReviewEvent,
    BankPaymentProviderEvent,
    BankPaymentReconciliationAttempt,
    DebtWriteOffEvent,
    Payment,
    PaymentRefund,
    PaymentRefundCase,
    Subscription,
    SubscriptionFreeze,
    Tariff,
    TariffComponent,
    TrainingType,
)
from apps.billing.services import verify_payment
from apps.billing.tests.factories import (
    DiscountFactory,
    ExpenseFactory,
    PaymentFactory,
    SubscriptionFactory,
    TariffComponentFactory,
    TariffFactory,
    TrainingTypeFactory,
)
from apps.clubs.models import ClubMembership
from apps.clubs.tests.factories import LocationFactory, UserFactory
from apps.documents.tests.factories import DocumentTypeFactory, StudentDocumentFactory
from apps.grades.tests.factories import GradeSystemFactory
from apps.notifications.models import MassNotification, NotificationPreference, NotificationTemplate
from apps.notifications.tests.factories import NotificationTemplateFactory, PushSubscriptionFactory
from apps.onboarding.models import OnboardingDraft
from apps.students.models import AccountAccess, ParentInvite, Student
from apps.students.tests.factories import StudentFactory
from apps.trainers.services import close_trainer_payroll_period
from apps.trainers.tests.factories import TrainerFactory, TrainerLocationFactory

pytestmark = pytest.mark.django_db

MOCK_METRICS = {
    "active_subscriptions": 42,
    "expiring_subscriptions": 5,
    "debtors": 3,
    "checkins": 18,
    "revenue": Decimal("150000"),
    "new_students": 7,
}

ZERO_METRICS = {
    "active_subscriptions": 0,
    "expiring_subscriptions": 0,
    "debtors": 0,
    "checkins": 0,
    "revenue": Decimal("0"),
    "new_students": 0,
}


def _future_schedule_date(schedule: Schedule) -> date:
    today = date.today()
    days_until_schedule = (schedule.day_of_week - today.weekday()) % 7
    if days_until_schedule == 0:
        days_until_schedule = 7
    return today + timedelta(days=days_until_schedule)


class _FormNestingParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self._form_depth = 0
        self.nested_forms: list[tuple[int, int]] = []

    def handle_starttag(self, tag, attrs):
        if tag != "form":
            return
        if self._form_depth:
            self.nested_forms.append(self.getpos())
        self._form_depth += 1

    def handle_endtag(self, tag):
        if tag == "form" and self._form_depth:
            self._form_depth -= 1


def _assert_no_nested_forms(html: str) -> None:
    parser = _FormNestingParser()
    parser.feed(html)
    assert parser.nested_forms == []


def _create_paid_active_subscription(*, club, student):
    training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
    tariff = TariffFactory(training_type=training_type, price=Decimal("5000"))
    return SubscriptionFactory(
        tariff=tariff,
        student=student,
        status=Subscription.Status.ACTIVE,
        paid_amount=Decimal("5000"),
    )


def _create_admin_personal_drop_in(*, club, owner_user, idempotency_key):
    training_type = TrainingTypeFactory(
        club=club,
        kind=TrainingType.Kind.PERSONAL,
        drop_in_price=Decimal("1000.00"),
    )
    tariff = TariffFactory(
        club=club,
        training_type=training_type,
        price=Decimal("1000.00"),
        trainings_limit=1,
        duration_days=30,
        scope=Tariff.Scope.CLUB,
        trainer_payout_policy=Tariff.PayoutPolicy.ON_CHECKIN,
    )
    TariffComponentFactory(
        club=club,
        tariff=tariff,
        training_type=training_type,
        entitlement_kind=TariffComponent.EntitlementKind.FINITE_CREDITS,
        credits_total=1,
        scope=Tariff.Scope.CLUB,
        location=None,
        trainer_payout_policy=Tariff.PayoutPolicy.ON_CHECKIN,
        paid_amount_basis=Decimal("1000.00"),
    )
    trainer = TrainerFactory(club=club)
    location = LocationFactory(club=club)
    TrainerLocationFactory(
        club=club,
        trainer=trainer,
        location=location,
        rate_personal=Decimal("50.00"),
    )
    student = StudentFactory(club=club, status=Student.Status.ACTIVE)
    target_date = timezone.localdate() + timedelta(days=3)
    starts_at = timezone.make_aware(datetime.combine(target_date, time(10, 0)))
    result = book_personal_drop_in(
        club_id=club.id,
        student_id=student.id,
        trainer_id=trainer.id,
        starts_at=starts_at,
        ends_at=starts_at + timedelta(hours=1),
        location_id=location.id,
        training_type_id=training_type.id,
        tariff_id=tariff.id,
        actor_user_id=owner_user.id,
        idempotency_key=idempotency_key,
    )
    return {
        "booking": result.booking,
        "schedule": result.schedule,
        "student": student,
        "starts_at": starts_at,
        "ends_at": starts_at + timedelta(hours=1),
    }


def _confirm_admin_drop_in_payment(*, club, owner_user, booking, idempotency_key):
    with patch("apps.attendance.services.async_task"), patch("django_q.tasks.async_task"):
        link = create_personal_drop_in_payment(
            club_id=club.id,
            booking_id=booking.id,
            payment_method=Payment.Method.CASH,
            created_by_id=owner_user.id,
            idempotency_key=idempotency_key,
        )
        verify_payment(
            payment_id=link.payment_id,
            club_id=club.id,
            verified_by_id=owner_user.id,
            action="confirm",
        )
    return Payment.objects.for_club(club).select_related("subscription").get(id=link.payment_id)


class TestDashboardHome:
    def test_unauthenticated_redirects_to_login(self, client: Client):
        response = client.get("/dashboard/")
        assert response.status_code == 302
        assert "/dashboard/login/" in response.url

    def test_authenticated_returns_full_page(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "<aside" in content  # sidebar present
        assert "Дашборд" in content
        assert "Ученики" in content
        assert "Расписание" in content
        assert "Оплаты" in content
        assert "Тренеры" in content
        assert "Аналитика" in content
        assert "Настройки" in content

    def test_htmx_request_returns_partial(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        content = response.content.decode()
        # Partial should NOT contain sidebar
        assert "<aside" not in content
        # But should contain dashboard content
        assert "МОЙ ДЕНЬ" in content


class TestTrainingGroupReconciliationDashboard:
    @pytest.fixture(autouse=True)
    def _enable_training_group_new_writes_for_confirmed_reconciliation(self, settings):
        settings.TRAINING_GROUP_NEW_WRITES_ENABLED = True

    def test_owner_sees_canonical_group_lifecycle_and_archive_preconditions(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.attendance.models import TrainingGroupMembership
        from apps.attendance.tests.factories import TrainingGroupFactory

        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        trainer = TrainerFactory(club=club)
        group = TrainingGroupFactory(
            club=club,
            name="Owner lifecycle group",
            training_type=training_type,
            responsible_trainer=trainer,
        )
        schedule = ScheduleFactory(
            club=club,
            training_group=group,
            training_type=training_type,
            trainer=trainer,
        )
        student = StudentFactory(club=club)
        TrainingGroupMembership.objects.create(
            club=club,
            student=student,
            training_group=group,
            starts_on=date(2026, 7, 1),
            source=TrainingGroupMembership.Source.MANUAL,
        )
        client.force_login(owner_user)

        response = client.get("/dashboard/training-groups/reconciliation/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Текущее состояние канонических групп" in content
        assert group.name in content
        assert str(trainer) in content
        assert str(schedule.id) in content
        assert "Режим rollout:" in content
        assert "Требует сверки" in content
        assert f'aria-label="Выбрать слот #{schedule.id}"' in content
        assert "Перед архивом закройте активные состояния" in content

    def test_owner_can_archive_a_group_only_through_visible_clean_preconditions(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.attendance.models import TrainingGroup
        from apps.attendance.tests.factories import TrainingGroupFactory

        group = TrainingGroupFactory(club=club, name="Ready for archive")
        client.force_login(owner_user)

        page = client.get("/dashboard/training-groups/reconciliation/")
        assert page.status_code == 200
        content = page.content.decode()
        assert "Предусловия архивации сейчас чистые" in content
        assert f"/dashboard/training-groups/{group.id}/archive/" in content

        response = client.post(
            f"/dashboard/training-groups/{group.id}/archive/",
            {"rationale": "Owner confirmed that every lifecycle root is closed."},
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 204
        assert response["HX-Redirect"] == "/dashboard/training-groups/reconciliation/"
        group.refresh_from_db()
        assert group.status == TrainingGroup.Status.ARCHIVED

    def test_owner_preview_renders_digest_without_apply_or_mutation(self, client: Client, owner_user, club):
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        location = LocationFactory(club=club)
        trainer = TrainerFactory(club=club)
        first_schedule = ScheduleFactory(
            club=club,
            training_type=training_type,
            location=location,
            trainer=trainer,
        )
        second_schedule = ScheduleFactory(
            club=club,
            training_type=training_type,
            location=location,
            trainer=trainer,
            day_of_week=2,
        )
        student = StudentFactory(club=club)
        ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=first_schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=date(2026, 7, 1),
            created_from=ScheduleEnrollment.CreatedFrom.MANUAL,
        )
        before_count = ScheduleEnrollment.objects.for_club(club).count()
        client.force_login(owner_user)

        response = client.post(
            "/dashboard/training-groups/reconciliation/",
            {
                "canonical_name": "Dashboard preview",
                "schedule_ids": [first_schedule.id, second_schedule.id],
            },
            HTTP_HX_REQUEST="true",
        )

        content = response.content.decode()
        assert response.status_code == 200
        assert "Предпросмотр #" in content
        assert "в режиме reconciling" in content
        assert f"student #{student.id}" in content
        assert "link_existing" in content
        assert "Подтвердить и применить digest" in content
        assert ScheduleEnrollment.objects.for_club(club).count() == before_count

    def test_owner_can_apply_preview_digest_and_retry_safely(self, client: Client, owner_user, club):
        from apps.attendance.models import TrainingGroupMappingEvent, TrainingGroupRolloutState

        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        location = LocationFactory(club=club)
        trainer = TrainerFactory(club=club)
        first_schedule = ScheduleFactory(
            club=club,
            training_type=training_type,
            location=location,
            trainer=trainer,
        )
        second_schedule = ScheduleFactory(
            club=club,
            training_type=training_type,
            location=location,
            trainer=trainer,
            day_of_week=2,
        )
        student = StudentFactory(club=club)
        ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=first_schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=date(2026, 7, 1),
            created_from=ScheduleEnrollment.CreatedFrom.MANUAL,
        )
        rollout_state = TrainingGroupRolloutState.objects.for_club(club).get()
        update_training_group_rollout_state_for_test(
            TrainingGroupRolloutState.objects.for_club(club).filter(id=rollout_state.id),
            mode=TrainingGroupRolloutState.Mode.RECONCILING
        )
        client.force_login(owner_user)
        preview_payload = {
            "canonical_name": "Dashboard apply",
            "schedule_ids": [first_schedule.id, second_schedule.id],
            "reconciliation_action": "preview",
        }
        preview_response = client.post(
            "/dashboard/training-groups/reconciliation/",
            preview_payload,
            HTTP_HX_REQUEST="true",
        )
        digest_match = re.search(
            r'name="preview_digest" value="([0-9a-f]{64})"',
            preview_response.content.decode(),
        )
        assert preview_response.status_code == 200
        assert digest_match is not None

        apply_payload = {
            **preview_payload,
            "reconciliation_action": "apply",
            "preview_digest": digest_match.group(1),
            "rationale": "Owner confirmed the dashboard reconciliation.",
            "idempotency_key": "htmx-reconciliation-apply",
        }
        response = client.post(
            "/dashboard/training-groups/reconciliation/",
            apply_payload,
            HTTP_HX_REQUEST="true",
        )
        event_count = TrainingGroupMappingEvent.objects.for_club(club).count()
        retry_response = client.post(
            "/dashboard/training-groups/reconciliation/",
            apply_payload,
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert "Сверка применена" in response.content.decode()
        assert retry_response.status_code == 200
        assert "Сверка применена" in retry_response.content.decode()
        assert TrainingGroupMappingEvent.objects.for_club(club).count() == event_count == 3

    def test_owner_apply_rebinds_explicit_start_date_and_rejects_stale_digest(
        self, client: Client, owner_user, club
    ):
        from apps.attendance.models import TrainingGroupMembership, TrainingGroupRolloutState

        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        location = LocationFactory(club=club)
        trainer = TrainerFactory(club=club)
        first_schedule = ScheduleFactory(
            club=club,
            training_type=training_type,
            location=location,
            trainer=trainer,
        )
        second_schedule = ScheduleFactory(
            club=club,
            training_type=training_type,
            location=location,
            trainer=trainer,
            day_of_week=2,
        )
        student = StudentFactory(club=club)
        ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=first_schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=date(2026, 7, 1),
            created_from=ScheduleEnrollment.CreatedFrom.MANUAL,
        )
        ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=second_schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=date(2026, 7, 8),
            created_from=ScheduleEnrollment.CreatedFrom.MANUAL,
        )
        rollout_state = TrainingGroupRolloutState.objects.for_club(club).get()
        update_training_group_rollout_state_for_test(
            TrainingGroupRolloutState.objects.for_club(club).filter(id=rollout_state.id),
            mode=TrainingGroupRolloutState.Mode.RECONCILING
        )
        client.force_login(owner_user)
        preview_payload = {
            "canonical_name": "Dashboard explicit start",
            "schedule_ids": [first_schedule.id, second_schedule.id],
            "reconciliation_action": "preview",
        }

        unresolved_response = client.post(
            "/dashboard/training-groups/reconciliation/",
            preview_payload,
            HTTP_HX_REQUEST="true",
        )
        assert unresolved_response.status_code == 200
        assert f'name="start_date_{student.id}"' in unresolved_response.content.decode()

        selected_start_date = "2026-07-03"
        selected_preview_response = client.post(
            "/dashboard/training-groups/reconciliation/",
            {**preview_payload, f"start_date_{student.id}": selected_start_date},
            HTTP_HX_REQUEST="true",
        )
        digest_match = re.search(
            r'name="preview_digest" value="([0-9a-f]{64})"',
            selected_preview_response.content.decode(),
        )
        assert digest_match is not None

        stale_response = client.post(
            "/dashboard/training-groups/reconciliation/",
            {
                **preview_payload,
                f"start_date_{student.id}": "2026-07-04",
                "reconciliation_action": "apply",
                "preview_digest": digest_match.group(1),
                "rationale": "Owner submitted an intentionally stale start-date digest.",
                "idempotency_key": "htmx-explicit-start-stale",
            },
            HTTP_HX_REQUEST="true",
        )
        assert stale_response.status_code == 200
        assert "stale" in stale_response.content.decode().lower()
        assert TrainingGroupMembership.objects.for_club(club).count() == 0

        apply_payload = {
            **preview_payload,
            f"start_date_{student.id}": selected_start_date,
            "reconciliation_action": "apply",
            "preview_digest": digest_match.group(1),
            "rationale": "Owner confirmed the explicit membership start date.",
            "idempotency_key": "htmx-explicit-start-apply",
        }
        response = client.post(
            "/dashboard/training-groups/reconciliation/",
            apply_payload,
            HTTP_HX_REQUEST="true",
        )
        retry_response = client.post(
            "/dashboard/training-groups/reconciliation/",
            apply_payload,
            HTTP_HX_REQUEST="true",
        )

        membership = TrainingGroupMembership.objects.for_club(club).get(student=student)
        assert response.status_code == retry_response.status_code == 200
        assert "Сверка применена" in response.content.decode()
        assert "Сверка применена" in retry_response.content.decode()
        assert membership.starts_on.isoformat() == selected_start_date


class TestAdminLogin:
    def test_login_page_renders(self, client: Client):
        response = client.get("/dashboard/login/")
        assert response.status_code == 200
        assert "Войти".encode() in response.content

    def test_admin_and_login_pages_link_valid_svg_favicon(
        self,
        client: Client,
        owner_user,
    ):
        login_response = client.get("/dashboard/login/")
        client.force_login(owner_user)
        admin_response = client.get("/dashboard/schedule/")
        favicon = Path("static/admin/favicon.svg")

        assert login_response.status_code == 200
        assert admin_response.status_code == 200
        assert 'rel="icon" type="image/svg+xml" sizes="any" href="/static/admin/favicon.svg"' in (
            login_response.content.decode()
        )
        assert 'rel="icon" type="image/svg+xml" sizes="any" href="/static/admin/favicon.svg"' in (
            admin_response.content.decode()
        )
        assert favicon.exists()
        assert favicon.read_text(encoding="utf-8").lstrip().startswith("<svg")
        collected_favicon = finders.find("admin/favicon.svg")
        assert collected_favicon is not None
        assert Path(collected_favicon).resolve() == favicon.resolve()

    def test_login_with_valid_credentials(self, client: Client, owner_user):
        owner_user.set_password("testpass123")
        owner_user.save()
        response = client.post(
            "/dashboard/login/",
            {
                "email": owner_user.email,
                "password": "testpass123",
            },
        )
        assert response.status_code == 302
        assert response.url == "/dashboard/"

    def test_login_with_invalid_credentials(self, client: Client):
        response = client.post(
            "/dashboard/login/",
            {
                "email": "wrong@example.com",
                "password": "wrong",
            },
        )
        assert response.status_code == 200
        assert "Неверный email или пароль".encode() in response.content

    def test_failed_login_logs_hashed_client_ip(
        self,
        client: Client,
    ):
        source_ip = "203.0.113.10"
        cache.clear()
        with patch("apps.htmx_admin.views.dashboard.logger.warning") as mock_warning:
            response = client.post(
                "/dashboard/login/",
                {
                    "email": "wrong@example.com",
                    "password": "wrong",
                },
                REMOTE_ADDR=source_ip,
            )

        assert response.status_code == 200
        mock_warning.assert_called_once()
        message, = mock_warning.call_args.args
        extra = mock_warning.call_args.kwargs["extra"]
        assert message == "login_failed"
        assert extra["client_ip_hash"]
        assert extra["client_ip_hash"] != source_ip
        assert "ip" not in extra
        cache.clear()

    def test_blocked_login_logs_hashed_client_ip(
        self,
        client: Client,
    ):
        source_ip = "203.0.113.20"
        cache.clear()
        cache.set("login_attempts:ip:203.0.113.20", 5, 300)
        with patch("apps.htmx_admin.views.dashboard.logger.warning") as mock_warning:
            response = client.post(
                "/dashboard/login/",
                {
                    "email": "blocked@example.com",
                    "password": "wrong",
                },
                REMOTE_ADDR=source_ip,
            )

        assert response.status_code == 200
        mock_warning.assert_called_once()
        message, = mock_warning.call_args.args
        extra = mock_warning.call_args.kwargs["extra"]
        assert message == "login_blocked"
        assert extra["client_ip_hash"]
        assert extra["client_ip_hash"] != source_ip
        assert "ip" not in extra
        cache.clear()


class TestAdminLogout:
    def test_logout_redirects_to_login(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/logout/")
        assert response.status_code == 302
        assert "/dashboard/login/" in response.url


class TestDashboardMetrics:
    """ADMIN-01: Dashboard shows 6 KPI metrics."""

    @patch("apps.htmx_admin.views.dashboard.get_dashboard_metrics")
    @patch("apps.htmx_admin.views.dashboard.get_attention_alerts")
    def test_dashboard_shows_metrics(self, mock_alerts, mock_metrics, client: Client, owner_user, club):
        mock_metrics.return_value = MOCK_METRICS
        mock_alerts.return_value = []
        client.force_login(owner_user)
        response = client.get("/dashboard/")
        content = response.content.decode()
        assert "42" in content  # active_subscriptions
        assert "18" in content  # checkins
        assert "150 000" in content  # revenue
        assert "7" in content  # new_students

    @patch("apps.htmx_admin.views.dashboard.get_dashboard_metrics")
    @patch("apps.htmx_admin.views.dashboard.get_attention_alerts")
    def test_period_switcher(self, mock_alerts, mock_metrics, client: Client, owner_user, club):
        mock_metrics.return_value = ZERO_METRICS
        mock_alerts.return_value = []
        client.force_login(owner_user)
        response = client.get("/dashboard/?period=week")
        assert response.status_code == 200
        content = response.content.decode()
        # Week button should be active
        assert "period=week" in content

    @patch("apps.htmx_admin.views.dashboard.get_dashboard_metrics")
    @patch("apps.htmx_admin.views.dashboard.get_attention_alerts")
    def test_period_month(self, mock_alerts, mock_metrics, client: Client, owner_user, club):
        mock_metrics.return_value = ZERO_METRICS
        mock_alerts.return_value = []
        client.force_login(owner_user)
        response = client.get("/dashboard/?period=month")
        assert response.status_code == 200

    @patch("apps.htmx_admin.views.dashboard.get_dashboard_metrics")
    @patch("apps.htmx_admin.views.dashboard.get_attention_alerts")
    def test_metrics_partial_endpoint(self, mock_alerts, mock_metrics, client: Client, owner_user, club):
        mock_metrics.return_value = {**ZERO_METRICS, "active_subscriptions": 10}
        mock_alerts.return_value = []
        client.force_login(owner_user)
        response = client.get("/dashboard/metrics/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert "10" in response.content.decode()

    @patch("apps.htmx_admin.views.dashboard.get_dashboard_metrics")
    @patch("apps.htmx_admin.views.dashboard.get_attention_alerts")
    def test_polling_element_present(self, mock_alerts, mock_metrics, client: Client, owner_user, club):
        mock_metrics.return_value = ZERO_METRICS
        mock_alerts.return_value = []
        client.force_login(owner_user)
        response = client.get("/dashboard/")
        content = response.content.decode()
        assert 'hx-trigger="every 60s"' in content

    @patch("apps.htmx_admin.views.dashboard.get_dashboard_metrics")
    @patch("apps.htmx_admin.views.dashboard.get_attention_alerts")
    def test_refresh_button_present(self, mock_alerts, mock_metrics, client: Client, owner_user, club):
        mock_metrics.return_value = ZERO_METRICS
        mock_alerts.return_value = []
        client.force_login(owner_user)
        response = client.get("/dashboard/")
        content = response.content.decode()
        assert "Обновить" in content

    @patch("apps.htmx_admin.views.dashboard.get_dashboard_metrics")
    @patch("apps.htmx_admin.views.dashboard.get_attention_alerts")
    def test_dashboard_today_sessions_use_effective_occurrences(
        self, mock_alerts, mock_metrics, client: Client, owner_user, club
    ):
        mock_metrics.return_value = ZERO_METRICS
        mock_alerts.return_value = []
        today = date.today()
        tomorrow = today + timedelta(days=1)

        moved_away_schedule = ScheduleFactory(
            club=club,
            day_of_week=today.weekday(),
            start_time=time(18, 0),
            end_time=time(19, 0),
            group_name="Moved Away PT",
        )
        ScheduleExceptionFactory(
            schedule=moved_away_schedule,
            date=today,
            exception_type="rescheduled",
            new_date=tomorrow,
            new_start_time=time(20, 0),
            new_end_time=time(21, 0),
        )

        moved_in_schedule = ScheduleFactory(
            club=club,
            day_of_week=tomorrow.weekday(),
            start_time=time(12, 0),
            end_time=time(13, 0),
            group_name="Moved In PT",
        )
        ScheduleExceptionFactory(
            schedule=moved_in_schedule,
            date=tomorrow,
            exception_type="rescheduled",
            new_date=today,
            new_start_time=time(21, 0),
            new_end_time=time(22, 0),
        )

        client.force_login(owner_user)
        response = client.get("/dashboard/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Moved Away PT" not in content
        assert "Moved In PT" in content
        assert "21:00" in content


class TestDashboardAlerts:
    """ADMIN-02: Alerts appear with correct counts."""

    @patch("apps.htmx_admin.views.dashboard.get_dashboard_metrics")
    @patch("apps.htmx_admin.views.dashboard.get_attention_alerts")
    def test_alerts_render(self, mock_alerts, mock_metrics, client: Client, owner_user, club):
        mock_metrics.return_value = ZERO_METRICS
        mock_alerts.return_value = [
            {"type": "expiring_subscriptions", "count": 5},
            {"type": "at_risk_students", "count": 3},
        ]
        client.force_login(owner_user)
        response = client.get("/dashboard/")
        content = response.content.decode()
        assert "Истекающие абонементы" in content
        assert "Ученики в зоне риска" in content
        assert "5" in content
        assert "3" in content
        assert "ТРЕБУЕТ ВНИМАНИЯ" in content

    @patch("apps.htmx_admin.views.dashboard.get_dashboard_metrics")
    @patch("apps.htmx_admin.views.dashboard.get_attention_alerts")
    def test_no_alerts_section_when_empty(self, mock_alerts, mock_metrics, client: Client, owner_user, club):
        mock_metrics.return_value = ZERO_METRICS
        mock_alerts.return_value = []
        client.force_login(owner_user)
        response = client.get("/dashboard/")
        content = response.content.decode()
        # With Pencil design, the two-column layout is always shown.
        # When no alerts, the "Всё в порядке" message appears.
        assert "Всё в порядке" in content

    @patch("apps.htmx_admin.views.dashboard.get_dashboard_metrics")
    @patch("apps.htmx_admin.views.dashboard.get_attention_alerts")
    def test_alert_links_correct(self, mock_alerts, mock_metrics, client: Client, owner_user, club):
        mock_metrics.return_value = ZERO_METRICS
        mock_alerts.return_value = [
            {"type": "unconfirmed_payments", "count": 2},
        ]
        client.force_login(owner_user)
        response = client.get("/dashboard/")
        content = response.content.decode()
        assert "/dashboard/billing/payments/" in content


class TestStudentList:
    """ADMIN-03: Student CRUD, filter, search, pagination."""

    def test_student_list_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/students/")
        assert response.status_code == 200
        assert "УЧЕНИКИ".encode() in response.content

    def test_student_list_htmx_partial(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/students/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert b"<aside" not in response.content

    def test_student_list_search(self, client: Client, owner_user, club):
        StudentFactory(club=club, first_name="UniqueSearchName")
        client.force_login(owner_user)
        response = client.get("/dashboard/students/?q=UniqueSearchName")
        assert response.status_code == 200
        assert b"UniqueSearchName" in response.content

    def test_student_list_search_by_phone(self, client: Client, owner_user, club):
        StudentFactory(club=club, phone="+79991112233", first_name="PhoneTest")
        client.force_login(owner_user)
        response = client.get("/dashboard/students/?q=9991112233")
        assert response.status_code == 200
        assert b"PhoneTest" in response.content

    def test_student_list_status_filter(self, client: Client, owner_user, club):
        StudentFactory(club=club, status="active", first_name="ActiveStudent")
        StudentFactory(club=club, status="lead", first_name="LeadStudent")
        client.force_login(owner_user)
        response = client.get("/dashboard/students/?status=active")
        assert response.status_code == 200
        content = response.content.decode()
        assert "ActiveStudent" in content
        assert "LeadStudent" not in content

    def test_student_card_renders(self, client: Client, owner_user, club):
        student = StudentFactory(club=club)
        client.force_login(owner_user)
        response = client.get(f"/dashboard/students/{student.id}/card/")
        assert response.status_code == 200
        assert student.first_name.encode() in response.content
        content = response.content.decode()
        assert 'data-attendance-metric="total"' in content
        assert 'data-value="0"' in content
        assert "ВСЕГО" in content
        assert "ПОСЕЩЕНО" in content
        assert "ПРОПУЩЕНО" in content

    def test_student_card_unifies_completed_single_credits_into_attendance(
        self,
        client: Client,
        owner_user,
        club,
    ):
        student = StudentFactory(club=club)
        training_type = TrainingTypeFactory(
            club=club,
            kind=TrainingType.Kind.PERSONAL,
        )
        tariff_names = []
        for offset in range(2):
            tariff = TariffFactory(
                club=club,
                training_type=training_type,
                trainings_limit=1,
                name=f"Проведённая разовая {offset + 1}",
            )
            tariff_names.append(tariff.name)
            subscription = SubscriptionFactory(
                club=club,
                student=student,
                tariff=tariff,
                status=Subscription.Status.EXPIRED,
                trainings_left=0,
                trainings_used=1,
            )
            schedule = ScheduleFactory(
                club=club,
                training_type=training_type,
            )
            CheckinFactory(
                club=club,
                student=student,
                schedule=schedule,
                training_type=training_type,
                trainer=schedule.trainer,
                location=schedule.location,
                subscription=subscription,
                date=timezone.localdate() - timedelta(days=offset),
            )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/students/{student.id}/card/")
        content = response.content.decode()

        assert response.status_code == 200
        assert "РАЗОВАЯ ТРЕНИРОВКА" not in content
        assert all(tariff_name not in content for tariff_name in tariff_names)
        assert "0 из 1" not in content
        assert "100%" in content
        assert re.search(r'data-attendance-metric="total"\s+data-value="2"', content)
        assert re.search(r'data-attendance-metric="attended"\s+data-value="2"', content)
        assert re.search(r'data-attendance-metric="missed"\s+data-value="0"', content)

    def test_student_card_shows_only_current_single_credits_and_multi_credit_remaining(
        self,
        client: Client,
        owner_user,
        club,
    ):
        student = StudentFactory(club=club)
        training_type = TrainingTypeFactory(
            club=club,
            kind=TrainingType.Kind.GROUP,
        )
        single_tariff = TariffFactory(
            club=club,
            training_type=training_type,
            trainings_limit=1,
            name="Разовый групповой вход",
        )
        SubscriptionFactory(
            club=club,
            student=student,
            tariff=single_tariff,
            status=Subscription.Status.ACTIVE,
            trainings_left=1,
            trainings_used=0,
            expires_at=timezone.now() - timedelta(minutes=1),
        )
        unavailable_tariff = TariffFactory(
            club=club,
            training_type=training_type,
            trainings_limit=1,
            name="Неконсистентная разовая",
        )
        SubscriptionFactory(
            club=club,
            student=student,
            tariff=unavailable_tariff,
            status=Subscription.Status.ACTIVE,
            trainings_left=0,
            trainings_used=0,
        )
        available_tariff = TariffFactory(
            club=club,
            training_type=training_type,
            trainings_limit=1,
            name="Доступная разовая",
        )
        SubscriptionFactory(
            club=club,
            student=student,
            tariff=available_tariff,
            status=Subscription.Status.ACTIVE,
            trainings_left=1,
            trainings_used=0,
        )
        pending_tariff = TariffFactory(
            club=club,
            training_type=training_type,
            trainings_limit=1,
            name="Разовая на подтверждении",
        )
        SubscriptionFactory(
            club=club,
            student=student,
            tariff=pending_tariff,
            status=Subscription.Status.PENDING,
            trainings_left=1,
            trainings_used=0,
        )
        multi_tariff = TariffFactory(
            club=club,
            training_type=training_type,
            trainings_limit=8,
            name="Групповой абонемент",
        )
        SubscriptionFactory(
            club=club,
            student=student,
            tariff=multi_tariff,
            status=Subscription.Status.ACTIVE,
            trainings_left=5,
            trainings_used=3,
        )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/students/{student.id}/card/")
        content = response.content.decode()

        assert response.status_code == 200
        assert content.count("РАЗОВАЯ ТРЕНИРОВКА") == 2
        assert "Разовый групповой вход" not in content
        assert "Неконсистентная разовая" not in content
        assert "ИСТЕКЛА" not in content
        assert "НЕДОСТУПНА" not in content
        assert "ДОСТУПНА" in content
        assert "ОЖИДАЕТ ПОДТВЕРЖДЕНИЯ" in content
        assert "Осталось 5 из 8" in content

    def test_student_edit_duplicate_phone_renders_inline_error(self, client: Client, owner_user, club):
        existing = StudentFactory(club=club, phone="+79991112233", first_name="Existing")
        edited = StudentFactory(club=club, phone="+79994445566", first_name="Edited")

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/students/{edited.id}/edit/",
            {
                "first_name": edited.first_name,
                "last_name": edited.last_name,
                "phone": existing.phone,
                "email": edited.email,
                "source": edited.source,
                "contraindications": edited.contraindications,
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        content = response.content.decode()
        assert "Ученик с таким телефоном уже существует" in content
        edited.refresh_from_db()
        assert edited.phone == "+79994445566"

    def test_student_card_status_options_are_allowed_transitions(self, client: Client, owner_user, club):
        # Arrange
        student = StudentFactory(club=club, status="active")

        # Act
        client.force_login(owner_user)
        response = client.get(f"/dashboard/students/{student.id}/card/")

        # Assert
        assert response.status_code == 200
        content = response.content.decode()
        assert '<option value="active" selected>Активен</option>' in content
        assert '<option value="at_risk">В риске</option>' in content
        assert '<option value="lead">Лид</option>' not in content
        assert '<option value="trial">Пробное</option>' not in content
        assert '<option value="churned">Ушёл</option>' not in content
        assert '<option value="lost">Потерян</option>' not in content

    def test_student_status_change_calls_transition_service_and_triggers_update(
        self,
        client: Client,
        owner_user,
        club,
        monkeypatch: pytest.MonkeyPatch,
    ):
        # Arrange
        from apps.htmx_admin.views import students as student_views

        student = StudentFactory(club=club, status="lead")
        calls = []

        def fake_transition_status(
            *,
            student_id: int,
            club_id: int,
            new_status: str,
            actor_user_id: int | None = None,
            source: str = "",
        ):
            calls.append(
                {
                    "student_id": student_id,
                    "club_id": club_id,
                    "new_status": new_status,
                    "actor_user_id": actor_user_id,
                    "source": source,
                }
            )
            student.status = new_status
            student.save(update_fields=["status"])
            return student

        monkeypatch.setattr(student_views, "transition_status", fake_transition_status, raising=False)

        # Act
        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/students/{student.id}/status/",
            {"new_status": "trial"},
            HTTP_HX_REQUEST="true",
        )

        # Assert
        assert response.status_code == 200
        assert calls == [
            {
                "student_id": student.id,
                "club_id": club.id,
                "new_status": "trial",
                "actor_user_id": owner_user.id,
                "source": "htmx_admin_student_status",
            }
        ]
        student.refresh_from_db()
        assert student.status == "trial"
        assert response["HX-Trigger"] == "studentUpdated"

    def test_student_status_change_invalid_transition_shows_error_without_trigger(
        self,
        client: Client,
        owner_user,
        club,
    ):
        # Arrange
        student = StudentFactory(club=club, status="active")

        # Act
        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/students/{student.id}/status/",
            {"new_status": "lost"},
            HTTP_HX_REQUEST="true",
        )

        # Assert
        assert response.status_code == 200
        student.refresh_from_db()
        assert student.status == "active"
        content = response.content.decode()
        assert "Нельзя перевести из" in content
        assert "studentUpdated" not in response.headers.get("HX-Trigger", "")

    def test_student_tenant_isolation(self, client: Client, owner_user, club, other_club):
        StudentFactory(club=club, first_name="MyStudent")
        StudentFactory(club=other_club, first_name="OtherStudent")
        client.force_login(owner_user)
        response = client.get("/dashboard/students/")
        content = response.content.decode()
        assert "MyStudent" in content
        assert "OtherStudent" not in content

    def test_student_card_tenant_isolation(self, client: Client, owner_user, club, other_club):
        other_student = StudentFactory(club=other_club)
        client.force_login(owner_user)
        response = client.get(f"/dashboard/students/{other_student.id}/card/")
        assert response.status_code == 404

    def test_unauthenticated_redirects(self, client: Client):
        response = client.get("/dashboard/students/")
        assert response.status_code == 302

    def test_student_card_shows_parent_invite_action_for_child_only(self, client: Client, owner_user, club):
        child = StudentFactory(club=club, is_child=True)
        adult = StudentFactory(club=club, is_child=False)
        client.force_login(owner_user)

        child_response = client.get(f"/dashboard/students/{child.id}/card/")
        adult_response = client.get(f"/dashboard/students/{adult.id}/card/")

        assert child_response.status_code == 200
        assert adult_response.status_code == 200
        assert "Связать родителя" in child_response.content.decode()
        assert "Связать родителя" not in adult_response.content.decode()

    def test_owner_can_generate_parent_invite_for_child_without_logging_token(
        self,
        client: Client,
        owner_user,
        club,
        caplog: pytest.LogCaptureFixture,
    ):
        child = StudentFactory(club=club, is_child=True)
        caplog.set_level(logging.INFO)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/students/{child.id}/parent-invite/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        invite = ParentInvite.objects.for_club(club).get(student=child)
        content = response.content.decode()
        assert "Ссылка для привязки родителя" in content
        assert "не выдаёт пароль" in content
        assert f"/parent-invite/{invite.token}" in content
        assert str(invite.token) in content
        assert str(invite.token) not in caplog.text

    def test_admin_can_generate_parent_invite_for_child(self, client: Client, admin_user, club):
        child = StudentFactory(club=club, is_child=True)
        client.force_login(admin_user)

        response = client.post(
            f"/dashboard/students/{child.id}/parent-invite/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        invite = ParentInvite.objects.for_club(club).get(student=child)
        assert str(invite.token) in response.content.decode()

    def test_parent_invite_rejects_non_child_with_safe_error(self, client: Client, owner_user, club):
        adult = StudentFactory(club=club, is_child=False)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/students/{adult.id}/parent-invite/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 422
        assert "Родителя можно пригласить только для детской анкеты." in response.content.decode()
        assert not ParentInvite.objects.for_club(club).filter(student=adult).exists()

    def test_parent_invite_foreign_student_is_not_visible(self, client: Client, owner_user, club, other_club):
        foreign_child = StudentFactory(club=other_club, is_child=True)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/students/{foreign_child.id}/parent-invite/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 404
        assert not ParentInvite.objects.for_club(other_club).filter(student=foreign_child).exists()

    @pytest.mark.parametrize("user_fixture", ["trainer_user", "parent_user"])
    def test_parent_invite_htmx_action_denies_non_management_roles(
        self,
        request: pytest.FixtureRequest,
        client: Client,
        user_fixture: str,
        club,
    ):
        child = StudentFactory(club=club, is_child=True)
        client.force_login(request.getfixturevalue(user_fixture))

        response = client.post(
            f"/dashboard/students/{child.id}/parent-invite/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 403
        assert not ParentInvite.objects.for_club(club).filter(student=child).exists()

    def test_student_card_shows_account_access_panel_and_requires_paid_subscription(
        self,
        client: Client,
        owner_user,
        club,
    ):
        paid = StudentFactory(
            club=club,
            is_child=False,
            status="active",
            phone="8 900 123 45 67",
        )
        unpaid = StudentFactory(club=club, is_child=False, status="active")
        _create_paid_active_subscription(club=club, student=paid)
        client.force_login(owner_user)

        paid_response = client.get(f"/dashboard/students/{paid.id}/card/")
        unpaid_response = client.get(f"/dashboard/students/{unpaid.id}/card/")

        assert paid_response.status_code == 200
        assert unpaid_response.status_code == 200
        paid_content = paid_response.content.decode()
        unpaid_content = unpaid_response.content.decode()
        assert "ЛИЧНЫЙ КАБИНЕТ" in paid_content
        assert f'hx-post="/dashboard/students/{paid.id}/account-access/open/"' in paid_content
        assert "Доступ откроется после активной оплаченной подписки." not in paid_content
        assert "Доступ откроется после активной оплаченной подписки." in unpaid_content

    def test_owner_can_open_student_account_access_without_logging_password(
        self,
        client: Client,
        owner_user,
        club,
        caplog: pytest.LogCaptureFixture,
    ):
        generated_password = "temporary-access-password-for-test"
        student = StudentFactory(
            club=club,
            is_child=False,
            status="active",
            phone="8 900 123 45 67",
        )
        _create_paid_active_subscription(club=club, student=student)
        client.force_login(owner_user)
        caplog.set_level(logging.INFO)

        with patch(
            "apps.students.access_services._generate_temporary_password",
            return_value=generated_password,
        ):
            response = client.post(
                f"/dashboard/students/{student.id}/account-access/open/",
                HTTP_HX_REQUEST="true",
            )

        assert response.status_code == 200
        student.refresh_from_db()
        assert student.user_id is not None
        assert AccountAccess.objects.for_club(club).filter(student=student).count() == 1
        content = response.content.decode()
        assert "ВРЕМЕННЫЙ ПАРОЛЬ" in content
        assert generated_password in content
        assert generated_password not in caplog.text

    def test_account_access_open_requires_paid_subscription_with_safe_error(
        self,
        client: Client,
        owner_user,
        club,
    ):
        student = StudentFactory(
            club=club,
            is_child=False,
            status="active",
            phone="8 900 123 45 67",
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/students/{student.id}/account-access/open/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert "Кабинет открывается только после активной оплаченной подписки." in response.content.decode()
        assert not AccountAccess.objects.for_club(club).filter(student=student).exists()

    def test_owner_can_open_child_parent_account_access(self, client: Client, owner_user, club):
        child = StudentFactory(
            club=club,
            is_child=True,
            status="active",
            phone="8 900 123 45 67",
        )
        _create_paid_active_subscription(club=club, student=child)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/students/{child.id}/account-access/open/",
            {"parent_phone": "8 901 222 33 44"},
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        child.refresh_from_db()
        assert child.user_id is None
        assert child.parent_user_id is not None
        assert AccountAccess.objects.for_club(club).filter(
            student=child,
            role=AccountAccess.Role.PARENT,
        ).count() == 1

    def test_child_account_access_form_does_not_default_to_child_phone(self, client: Client, owner_user, club):
        child = StudentFactory(
            club=club,
            is_child=True,
            status="active",
            phone="8 900 123 45 67",
        )
        _create_paid_active_subscription(club=club, student=child)
        client.force_login(owner_user)

        response = client.get(f"/dashboard/students/{child.id}/card/")

        assert response.status_code == 200
        content = response.content.decode()
        assert 'name="parent_phone"' in content
        assert f'value="{child.phone}"' not in content

    def test_owner_can_open_child_access_for_linked_parent_without_parent_phone(
        self,
        client: Client,
        owner_user,
        club,
    ):
        parent = UserFactory(username="+79012223344")
        parent.set_unusable_password()
        parent.save(update_fields=["password"])
        child = StudentFactory(
            club=club,
            is_child=True,
            status="active",
            phone="8 900 123 45 67",
            parent_user=parent,
        )
        _create_paid_active_subscription(club=club, student=child)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/students/{child.id}/account-access/open/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert ClubMembership.objects.filter(
            user=parent,
            club=club,
            role=ClubMembership.Role.PARENT,
            is_active=True,
        ).count() == 1
        assert AccountAccess.objects.for_club(club).filter(
            student=child,
            user=parent,
            role=AccountAccess.Role.PARENT,
        ).count() == 1

    def test_owner_can_reset_student_account_access_and_see_new_one_time_password(
        self,
        client: Client,
        owner_user,
        club,
    ):
        student = StudentFactory(
            club=club,
            is_child=False,
            status="active",
            phone="8 900 123 45 67",
        )
        _create_paid_active_subscription(club=club, student=student)
        client.force_login(owner_user)

        with patch(
            "apps.students.access_services._generate_temporary_password",
            side_effect=["first-temporary-password", "second-temporary-password"],
        ):
            open_response = client.post(
                f"/dashboard/students/{student.id}/account-access/open/",
                HTTP_HX_REQUEST="true",
            )
            reset_response = client.post(
                f"/dashboard/students/{student.id}/account-access/reset/",
                HTTP_HX_REQUEST="true",
            )

        assert open_response.status_code == 200
        assert reset_response.status_code == 200
        content = reset_response.content.decode()
        assert "ВРЕМЕННЫЙ ПАРОЛЬ" in content
        assert "second-temporary-password" in content

    def test_account_access_foreign_student_is_not_visible(self, client: Client, owner_user, club, other_club):
        student = StudentFactory(club=other_club, is_child=False, status="active")
        _create_paid_active_subscription(club=other_club, student=student)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/students/{student.id}/account-access/open/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 404
        assert not AccountAccess.objects.for_club(other_club).filter(student=student).exists()

    @pytest.mark.parametrize("user_fixture", ["trainer_user", "parent_user"])
    def test_account_access_htmx_action_denies_non_management_roles(
        self,
        request: pytest.FixtureRequest,
        client: Client,
        user_fixture: str,
        club,
    ):
        student = StudentFactory(club=club, is_child=False, status="active")
        _create_paid_active_subscription(club=club, student=student)
        client.force_login(request.getfixturevalue(user_fixture))

        response = client.post(
            f"/dashboard/students/{student.id}/account-access/open/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 403
        assert not AccountAccess.objects.for_club(club).filter(student=student).exists()


class TestExcelImport:
    """ADMIN-04: Excel import preview and confirm."""

    def test_import_form_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/students/import/")
        assert response.status_code == 200
        assert "ИМПОРТ ИЗ EXCEL".encode() in response.content

    def test_import_no_file_shows_error(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.post("/dashboard/students/import/")
        assert response.status_code == 200
        assert "Выберите файл".encode() in response.content

    def test_import_confirm_creates_students_without_logging_collision(self, client: Client, owner_user, club):
        wb = Workbook()
        ws = wb.active
        ws.append(["name", "phone"])
        ws.append(["Import Confirm Student", "+79990000123"])
        buffer = io.BytesIO()
        wb.save(buffer)
        wb.close()
        buffer.seek(0)

        client.force_login(owner_user)
        upload_response = client.post(
            "/dashboard/students/import/",
            {
                "file": SimpleUploadedFile(
                    "students.xlsx",
                    buffer.read(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            HTTP_HX_REQUEST="true",
        )
        assert upload_response.status_code == 200

        response = client.post("/dashboard/students/import/confirm/", HTTP_HX_REQUEST="true")

        assert response.status_code == 204
        assert response.headers["HX-Redirect"] == "/dashboard/students/"
        student = Student.objects.for_club(club).get(first_name="Import", last_name="Confirm Student")
        assert student.status == Student.Status.LEAD

    def test_import_preview_flags_existing_guardian_phone(self, client: Client, owner_user, club):
        StudentFactory(
            club=club,
            first_name="Existing Child",
            is_child=True,
            phone="",
            guardian_phone="+79990000124",
        )
        wb = Workbook()
        ws = wb.active
        ws.append(["name", "phone"])
        ws.append(["Import Duplicate", "8 (999) 000-01-24"])
        buffer = io.BytesIO()
        wb.save(buffer)
        wb.close()
        buffer.seek(0)

        client.force_login(owner_user)
        response = client.post(
            "/dashboard/students/import/",
            {
                "file": SimpleUploadedFile(
                    "students.xlsx",
                    buffer.read(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert "Телефон уже существует".encode() in response.content
        assert not Student.objects.for_club(club).filter(phone="+79990000124").exists()


class TestBillingSettingsTrainingTypes:
    def test_training_type_form_includes_grade_system_selector(self, client: Client, owner_user, club):
        grade_system = GradeSystemFactory(club=club, discipline="BJJ")
        client.force_login(owner_user)

        response = client.get("/dashboard/settings/billing/training-types/form/")

        assert response.status_code == 200
        content = response.content.decode()
        assert 'name="kind"' in content
        assert 'value="personal"' in content
        assert 'name="grade_system_id"' in content
        assert f'value="{grade_system.id}"' in content
        assert "BJJ" in content

    def test_training_type_form_persists_kind_and_grade_system(self, client: Client, owner_user, club):
        grade_system = GradeSystemFactory(club=club, discipline="BJJ")
        client.force_login(owner_user)

        response = client.post(
            "/dashboard/settings/billing/training-types/form/",
            {
                "name": "BJJ",
                "kind": TrainingType.Kind.PERSONAL,
                "grade_system_id": str(grade_system.id),
            },
        )

        assert response.status_code == 204
        training_type = TrainingType.objects.for_club(club).get(name="BJJ")
        assert training_type.kind == TrainingType.Kind.PERSONAL
        assert training_type.grade_system_id == grade_system.id

    def test_used_training_type_form_locks_kind_and_preserves_it_on_post(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.billing.tests.factories import TariffFactory

        training_type = TrainingTypeFactory(
            club=club,
            name="Personal",
            kind=TrainingType.Kind.PERSONAL,
        )
        TariffFactory(training_type=training_type)
        client.force_login(owner_user)

        response = client.get(f"/dashboard/settings/billing/training-types/{training_type.id}/form/")

        assert response.status_code == 200
        content = response.content.decode()
        assert 'name="kind" value="personal"' in content
        assert "disabled aria-describedby" in content
        assert "Формат нельзя менять" in content

        post_response = client.post(
            f"/dashboard/settings/billing/training-types/{training_type.id}/form/",
            {
                "name": "Personal renamed",
                "grade_system_id": "",
            },
        )

        assert post_response.status_code == 204
        training_type.refresh_from_db()
        assert training_type.name == "Personal renamed"
        assert training_type.kind == TrainingType.Kind.PERSONAL


class TestBillingSettingsTariffs:
    def test_tariff_form_includes_component_builder(self, client: Client, owner_user, club):
        TrainingTypeFactory(club=club, name="Группа", kind=TrainingType.Kind.GROUP)
        client.force_login(owner_user)

        response = client.get("/dashboard/settings/billing/tariffs/form/")

        assert response.status_code == 200
        content = response.content.decode()
        assert 'name="trainer_payout_policy"' in content
        assert 'name="use_components"' in content
        assert 'name="component_paid_amount_basis_0"' in content
        assert 'name="component_weekly_limit_0"' in content
        assert 'name="is_personal_booking_default"' in content

    def test_tariff_form_sets_personal_booking_default_and_shows_badge(
        self, client: Client, owner_user, club
    ):
        personal_type = TrainingTypeFactory(
            club=club,
            name="Персонал",
            kind=TrainingType.Kind.PERSONAL,
            slug="settings-personal-default",
        )
        client.force_login(owner_user)

        response = client.post(
            "/dashboard/settings/billing/tariffs/form/",
            {
                "name": "Разовая персоналка",
                "training_type_id": str(personal_type.id),
                "price": "2500",
                "trainings_limit": "1",
                "duration_days": "1",
                "description": "",
                "scope": Tariff.Scope.CLUB,
                "trainer_payout_policy": Tariff.PayoutPolicy.ON_CHECKIN,
                "is_personal_booking_default": "on",
            },
        )

        assert response.status_code == 204
        tariff = Tariff.objects.for_club(club).get(name="Разовая персоналка")
        assert tariff.is_personal_booking_default is True

        list_response = client.get("/dashboard/settings/billing/")
        assert list_response.status_code == 200
        assert "Текущая цена персоналки" in list_response.content.decode()

    def test_tariff_form_configures_and_edits_trainer_specific_personal_price(
        self, client: Client, owner_user, club
    ):
        personal_type = TrainingTypeFactory(
            club=club,
            name="Персонал",
            kind=TrainingType.Kind.PERSONAL,
            slug="settings-trainer-personal-default",
        )
        trainer = TrainerFactory(
            club=club,
            first_name="Алмаз",
            last_name="Тренер",
            is_active=True,
        )
        other_trainer = TrainerFactory(
            club=club,
            first_name="Ренат",
            last_name="Тренер",
            is_active=True,
        )
        client.force_login(owner_user)

        form = client.get("/dashboard/settings/billing/tariffs/form/")
        assert form.status_code == 200
        content = form.content.decode()
        assert 'name="personal_booking_trainer_id"' in content
        assert f'value="{trainer.id}"' in content

        invalid = client.post(
            "/dashboard/settings/billing/tariffs/form/",
            {
                "name": "Неполная персоналка",
                "training_type_id": str(personal_type.id),
                "price": "2000",
                "trainings_limit": "1",
                "duration_days": "",
                "scope": Tariff.Scope.CLUB,
                "trainer_payout_policy": Tariff.PayoutPolicy.ON_CHECKIN,
                "personal_booking_trainer_id": str(trainer.id),
            },
        )
        assert invalid.status_code == 200
        assert f'<option value="{trainer.id}" selected>' in invalid.content.decode()
        assert not Tariff.objects.for_club(club).filter(name="Неполная персоналка").exists()

        created = client.post(
            "/dashboard/settings/billing/tariffs/form/",
            {
                "name": "Персоналка Алмаза",
                "training_type_id": str(personal_type.id),
                "price": "2000",
                "trainings_limit": "1",
                "duration_days": "1",
                "description": "",
                "scope": Tariff.Scope.CLUB,
                "trainer_payout_policy": Tariff.PayoutPolicy.ON_CHECKIN,
                "is_personal_booking_default": "on",
                "personal_booking_trainer_id": str(trainer.id),
            },
        )
        assert created.status_code == 204
        tariff = Tariff.objects.for_club(club).get(name="Персоналка Алмаза")
        assert tariff.personal_booking_trainer_id == trainer.id
        assert tariff.is_personal_booking_default is True

        edited = client.post(
            f"/dashboard/settings/billing/tariffs/{tariff.id}/form/",
            {
                "name": tariff.name,
                "training_type_id": str(personal_type.id),
                "price": "2000",
                "trainings_limit": "1",
                "duration_days": "1",
                "description": "",
                "scope": Tariff.Scope.CLUB,
                "trainer_payout_policy": Tariff.PayoutPolicy.ON_CHECKIN,
                "personal_booking_trainer_id": str(other_trainer.id),
            },
        )
        assert edited.status_code == 204, edited.content.decode()
        tariff.refresh_from_db()
        assert tariff.personal_booking_trainer_id == other_trainer.id
        assert tariff.price == Decimal("2000")
        assert tariff.is_personal_booking_default is False

        reassigned = client.post(
            f"/dashboard/settings/billing/tariffs/{tariff.id}/form/",
            {
                "name": tariff.name,
                "training_type_id": str(personal_type.id),
                "price": "2000",
                "trainings_limit": "1",
                "duration_days": "1",
                "description": "",
                "scope": Tariff.Scope.CLUB,
                "trainer_payout_policy": Tariff.PayoutPolicy.ON_CHECKIN,
                "is_personal_booking_default": "on",
                "personal_booking_trainer_id": str(other_trainer.id),
            },
        )
        assert reassigned.status_code == 204
        tariff.refresh_from_db()
        assert tariff.is_personal_booking_default is True

        list_response = client.get("/dashboard/settings/billing/")
        assert "Ренат Тренер" in list_response.content.decode()

    def test_tariff_form_treats_single_custom_component_policy_as_component_builder(
        self, client: Client, owner_user, club
    ):
        group_type = TrainingTypeFactory(club=club, name="Группа", kind=TrainingType.Kind.GROUP)
        tariff = TariffFactory(
            training_type=group_type,
            name="Группа",
            price=Decimal("5000"),
            trainings_limit=8,
            trainer_payout_policy="",
        )
        TariffComponentFactory(
            tariff=tariff,
            name="Группа",
            training_type=group_type,
            entitlement_kind=TariffComponent.EntitlementKind.FINITE_CREDITS,
            credits_total=8,
            weekly_limit=None,
            paid_amount_basis=Decimal("5000"),
            trainer_payout_policy=Tariff.PayoutPolicy.ON_CHECKIN,
        )
        client.force_login(owner_user)

        response = client.get(f"/dashboard/settings/billing/tariffs/{tariff.id}/form/")

        assert response.status_code == 200
        content = response.content.decode()
        checkbox_index = content.index('name="use_components"')
        assert "checked" in content[checkbox_index : checkbox_index + 120]

    def test_tariff_form_creates_hybrid_components(self, client: Client, owner_user, club):
        group_type = TrainingTypeFactory(
            club=club,
            name="Группа",
            kind=TrainingType.Kind.GROUP,
            slug="settings-group",
        )
        personal_type = TrainingTypeFactory(
            club=club,
            name="Персонал",
            kind=TrainingType.Kind.PERSONAL,
            slug="settings-personal",
        )
        client.force_login(owner_user)

        response = client.post(
            "/dashboard/settings/billing/tariffs/form/",
            {
                "name": "Гибрид Оптимум",
                "training_type_id": str(group_type.id),
                "price": "9500",
                "trainings_limit": "",
                "duration_days": "30",
                "description": "",
                "trainer_payout_policy": "",
                "use_components": "on",
                "component_name_0": "Группа",
                "component_training_type_id_0": str(group_type.id),
                "component_entitlement_kind_0": TariffComponent.EntitlementKind.WEEKLY_LIMIT,
                "component_weekly_limit_0": "2",
                "component_paid_amount_basis_0": "3500",
                "component_trainer_payout_policy_0": Tariff.PayoutPolicy.ON_PAYMENT,
                "component_name_1": "Персоналки",
                "component_training_type_id_1": str(personal_type.id),
                "component_entitlement_kind_1": TariffComponent.EntitlementKind.FINITE_CREDITS,
                "component_credits_total_1": "3",
                "component_paid_amount_basis_1": "6000",
                "component_trainer_payout_policy_1": Tariff.PayoutPolicy.ON_CHECKIN,
            },
        )

        assert response.status_code == 204
        tariff = Tariff.objects.for_club(club).get(name="Гибрид Оптимум")
        components = list(
            TariffComponent.objects.for_club(club)
            .filter(tariff=tariff, is_active=True)
            .order_by("sort_order")
        )
        assert len(components) == 2
        assert components[0].training_type_id == group_type.id
        assert components[0].entitlement_kind == TariffComponent.EntitlementKind.WEEKLY_LIMIT
        assert components[0].weekly_limit == 2
        assert components[0].paid_amount_basis == Decimal("3500")
        assert components[0].trainer_payout_policy == Tariff.PayoutPolicy.ON_PAYMENT
        assert components[1].training_type_id == personal_type.id
        assert components[1].credits_total == 3
        assert components[1].paid_amount_basis == Decimal("6000")
        assert components[1].trainer_payout_policy == Tariff.PayoutPolicy.ON_CHECKIN

    def test_tariff_form_carries_location_scope_to_components(self, client: Client, owner_user, club):
        location = LocationFactory(club=club, name="Main Hall")
        group_type = TrainingTypeFactory(club=club, name="Группа", kind=TrainingType.Kind.GROUP)
        personal_type = TrainingTypeFactory(club=club, name="Персонал", kind=TrainingType.Kind.PERSONAL)
        client.force_login(owner_user)

        response = client.post(
            "/dashboard/settings/billing/tariffs/form/",
            {
                "name": "Гибрид Main",
                "training_type_id": str(group_type.id),
                "price": "9500",
                "trainings_limit": "",
                "duration_days": "30",
                "description": "",
                "scope": Tariff.Scope.LOCATION,
                "location_id": str(location.id),
                "trainer_payout_policy": "",
                "use_components": "on",
                "component_name_0": "Группа",
                "component_training_type_id_0": str(group_type.id),
                "component_entitlement_kind_0": TariffComponent.EntitlementKind.WEEKLY_LIMIT,
                "component_weekly_limit_0": "2",
                "component_paid_amount_basis_0": "3500",
                "component_trainer_payout_policy_0": Tariff.PayoutPolicy.ON_PAYMENT,
                "component_name_1": "Персоналки",
                "component_training_type_id_1": str(personal_type.id),
                "component_entitlement_kind_1": TariffComponent.EntitlementKind.FINITE_CREDITS,
                "component_credits_total_1": "3",
                "component_paid_amount_basis_1": "6000",
                "component_trainer_payout_policy_1": Tariff.PayoutPolicy.ON_CHECKIN,
            },
        )

        assert response.status_code == 204
        tariff = Tariff.objects.for_club(club).get(name="Гибрид Main")
        assert tariff.scope == Tariff.Scope.LOCATION
        components = TariffComponent.objects.for_club(club).filter(tariff=tariff, is_active=True)
        assert components.count() == 2
        assert {component.scope for component in components} == {Tariff.Scope.LOCATION}
        assert {component.location_id for component in components} == {location.id}

    def test_tariff_form_rejects_policy_only_partial_component_row(self, client: Client, owner_user, club):
        group_type = TrainingTypeFactory(club=club, name="Группа", kind=TrainingType.Kind.GROUP)
        client.force_login(owner_user)

        response = client.post(
            "/dashboard/settings/billing/tariffs/form/",
            {
                "name": "Bad Hybrid",
                "training_type_id": str(group_type.id),
                "price": "9500",
                "trainings_limit": "",
                "duration_days": "30",
                "description": "",
                "trainer_payout_policy": "",
                "use_components": "on",
                "component_trainer_payout_policy_0": Tariff.PayoutPolicy.ON_PAYMENT,
            },
        )

        assert response.status_code == 200
        assert "Компонент 1: выберите тип тренировки" in response.content.decode()
        assert not Tariff.objects.for_club(club).filter(name="Bad Hybrid").exists()

    def test_tariff_list_shows_component_summary(self, client: Client, owner_user, club):
        group_type = TrainingTypeFactory(club=club, name="Группа", kind=TrainingType.Kind.GROUP)
        personal_type = TrainingTypeFactory(club=club, name="Персонал", kind=TrainingType.Kind.PERSONAL)
        tariff = TariffFactory(training_type=group_type, name="Гибрид PRO", price=Decimal("12500"))
        TariffComponentFactory(
            tariff=tariff,
            name="Группа",
            training_type=group_type,
            entitlement_kind=TariffComponent.EntitlementKind.WEEKLY_LIMIT,
            credits_total=None,
            weekly_limit=5,
            paid_amount_basis=Decimal("4500"),
            trainer_payout_policy=Tariff.PayoutPolicy.ON_PAYMENT,
            sort_order=0,
        )
        TariffComponentFactory(
            tariff=tariff,
            name="Персоналки",
            training_type=personal_type,
            credits_total=4,
            paid_amount_basis=Decimal("8000"),
            trainer_payout_policy=Tariff.PayoutPolicy.ON_CHECKIN,
            sort_order=1,
        )
        client.force_login(owner_user)

        response = client.get("/dashboard/settings/billing/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Группа: 5/нед., 4500" in content
        assert "Персоналки: 4 занятий, 8000" in content
        assert "После оплаты / За посещение" in content

    def test_active_hybrid_tariff_allows_name_update_when_components_unchanged(
        self,
        client: Client,
        owner_user,
        club,
    ):
        group_type = TrainingTypeFactory(club=club, name="Группа", kind=TrainingType.Kind.GROUP)
        personal_type = TrainingTypeFactory(club=club, name="Персонал", kind=TrainingType.Kind.PERSONAL)
        tariff = TariffFactory(training_type=group_type, name="Гибрид", price=Decimal("9500"), trainings_limit=None)
        TariffComponentFactory(
            tariff=tariff,
            name="Группа",
            training_type=group_type,
            entitlement_kind=TariffComponent.EntitlementKind.WEEKLY_LIMIT,
            credits_total=None,
            weekly_limit=2,
            paid_amount_basis=Decimal("3500"),
            trainer_payout_policy=Tariff.PayoutPolicy.ON_PAYMENT,
            sort_order=0,
        )
        TariffComponentFactory(
            tariff=tariff,
            name="Персоналки",
            training_type=personal_type,
            credits_total=3,
            paid_amount_basis=Decimal("6000"),
            trainer_payout_policy=Tariff.PayoutPolicy.ON_CHECKIN,
            sort_order=1,
        )
        SubscriptionFactory(tariff=tariff, status=Subscription.Status.ACTIVE)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/settings/billing/tariffs/{tariff.id}/form/",
            {
                "name": "Гибрид Оптимум",
                "training_type_id": str(group_type.id),
                "price": "9500",
                "trainings_limit": "",
                "duration_days": "30",
                "description": "",
                "trainer_payout_policy": "",
                "use_components": "on",
                "component_name_0": "Группа",
                "component_training_type_id_0": str(group_type.id),
                "component_entitlement_kind_0": TariffComponent.EntitlementKind.WEEKLY_LIMIT,
                "component_weekly_limit_0": "2",
                "component_paid_amount_basis_0": "3500",
                "component_trainer_payout_policy_0": Tariff.PayoutPolicy.ON_PAYMENT,
                "component_name_1": "Персоналки",
                "component_training_type_id_1": str(personal_type.id),
                "component_entitlement_kind_1": TariffComponent.EntitlementKind.FINITE_CREDITS,
                "component_credits_total_1": "3",
                "component_paid_amount_basis_1": "6000",
                "component_trainer_payout_policy_1": Tariff.PayoutPolicy.ON_CHECKIN,
            },
        )

        assert response.status_code == 204
        tariff.refresh_from_db()
        assert tariff.name == "Гибрид Оптимум"


class TestScheduleManagement:
    """ADMIN-05: Schedule grid, cancel/reschedule/substitute."""

    def test_schedule_create_form_includes_training_type(self, client: Client, owner_user, club):
        training_type = TrainingTypeFactory(club=club, name="Group")
        client.force_login(owner_user)

        response = client.get("/dashboard/schedule/create/")

        assert response.status_code == 200
        content = response.content.decode()
        assert 'name="training_type_id"' in content
        assert f'value="{training_type.id}"' in content

    def test_schedule_create_persists_training_type(self, client: Client, owner_user, club):
        trainer = TrainerFactory(club=club)
        location = LocationFactory(club=club)
        training_type = TrainingTypeFactory(club=club)
        client.force_login(owner_user)

        response = client.post(
            "/dashboard/schedule/create/",
            {
                "group_name": "HTMX Kiosk Ready",
                "day_of_week": "0",
                "start_time": "10:00",
                "end_time": "11:00",
                "trainer_id": str(trainer.id),
                "location_id": str(location.id),
                "training_type_id": str(training_type.id),
            },
        )

        assert response.status_code == 204
        schedule = Schedule.objects.for_club(club).get(group_name="HTMX Kiosk Ready")
        assert schedule.training_type_id == training_type.id

    def test_schedule_create_can_add_slot_to_exact_existing_group_without_fragmentation(
        self,
        client: Client,
        owner_user,
        club,
        settings,
    ):
        from apps.attendance.models import TrainingGroup, TrainingGroupRolloutState
        from apps.attendance.tests.factories import TrainingGroupFactory

        settings.TRAINING_GROUP_NEW_WRITES_ENABLED = True
        trainer = TrainerFactory(club=club)
        group = TrainingGroupFactory(club=club, responsible_trainer=trainer)
        rollout = TrainingGroupRolloutState.objects.for_club(club).get()
        update_training_group_rollout_state_for_test(
            TrainingGroupRolloutState.objects.for_club(club).filter(id=rollout.id),
            mode=TrainingGroupRolloutState.Mode.SHADOW
        )
        group_count = TrainingGroup.objects.for_club(club).count()
        client.force_login(owner_user)

        form = client.get("/dashboard/schedule/create/")
        assert form.status_code == 200
        assert f'value="{group.id}"' in form.content.decode()
        response = client.post(
            "/dashboard/schedule/create/",
            {
                "training_group_id": str(group.id),
                "group_name": "",
                "day_of_week": "2",
                "start_time": "18:00",
                "end_time": "19:00",
                "trainer_id": str(trainer.id),
                "location_id": str(group.location_id),
                "training_type_id": str(group.training_type_id),
            },
        )

        assert response.status_code == 204
        schedule = Schedule.objects.for_club(club).get(
            training_group=group,
            day_of_week=2,
        )
        assert schedule.group_name == group.name
        assert TrainingGroup.objects.for_club(club).count() == group_count

    def test_schedule_create_rejects_unavailable_training_type(self, client: Client, owner_user, club, other_club):
        trainer = TrainerFactory(club=club)
        location = LocationFactory(club=club)
        foreign_type = TrainingTypeFactory(club=other_club)
        inactive_type = TrainingTypeFactory(club=club, is_active=False)
        client.force_login(owner_user)

        for training_type in (foreign_type, inactive_type):
            response = client.post(
                "/dashboard/schedule/create/",
                {
                    "group_name": f"Unavailable Type {training_type.id}",
                    "day_of_week": "0",
                    "start_time": "10:00",
                    "end_time": "11:00",
                    "trainer_id": str(trainer.id),
                    "location_id": str(location.id),
                    "training_type_id": str(training_type.id),
                },
            )

            assert response.status_code == 200
            assert (
                not Schedule.objects.for_club(club)
                .filter(
                    group_name=f"Unavailable Type {training_type.id}",
                )
                .exists()
            )

    def test_schedule_edit_form_includes_selected_training_type(self, client: Client, owner_user, club):
        current_type = TrainingTypeFactory(club=club, name="Group")
        other_type = TrainingTypeFactory(club=club, name="Personal")
        schedule = ScheduleFactory(club=club, training_type=current_type)
        client.force_login(owner_user)

        response = client.get(f"/dashboard/schedule/{schedule.id}/edit/")

        assert response.status_code == 200
        content = response.content.decode()
        assert 'name="training_type_id"' in content
        assert f'value="{current_type.id}" selected' in content
        assert f'value="{other_type.id}"' in content

    def test_schedule_edit_persists_training_type(self, client: Client, owner_user, club):
        trainer = TrainerFactory(club=club)
        location = LocationFactory(club=club)
        old_type = TrainingTypeFactory(club=club)
        new_type = TrainingTypeFactory(club=club)
        schedule = ScheduleFactory(club=club, trainer=trainer, location=location, training_type=old_type)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/{schedule.id}/edit/",
            {
                "group_name": "Updated HTMX Kiosk Ready",
                "day_of_week": str(schedule.day_of_week),
                "start_time": "10:30",
                "end_time": "11:30",
                "trainer_id": str(trainer.id),
                "location_id": str(location.id),
                "training_type_id": str(new_type.id),
            },
        )

        assert response.status_code == 204
        schedule.refresh_from_db()
        assert schedule.training_type_id == new_type.id
        assert schedule.group_name == "Updated HTMX Kiosk Ready"

    def test_linked_schedule_edit_keeps_identity_readonly_and_allows_safe_changes(
        self,
        client: Client,
        owner_user,
        club,
        settings,
    ):
        from apps.attendance.models import TrainingGroupRolloutState
        from apps.attendance.tests.factories import TrainingGroupFactory

        settings.TRAINING_GROUP_NEW_WRITES_ENABLED = True
        trainer = TrainerFactory(club=club)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        location = LocationFactory(club=club)
        other_location = LocationFactory(club=club)
        group = TrainingGroupFactory(
            club=club,
            name="Canonical linked group",
            training_type=training_type,
            location=location,
            responsible_trainer=trainer,
        )
        schedule = ScheduleFactory(
            club=club,
            group_name="Legacy linked snapshot",
            training_group=group,
            trainer=trainer,
            location=location,
            training_type=training_type,
            start_time=time(10, 0),
            end_time=time(11, 0),
        )
        Schedule.objects.for_club(club).filter(id=schedule.id).update(
            group_name="Legacy linked snapshot"
        )
        rollout = TrainingGroupRolloutState.objects.for_club(club).get()
        update_training_group_rollout_state_for_test(
            TrainingGroupRolloutState.objects.for_club(club).filter(id=rollout.id),
            mode=TrainingGroupRolloutState.Mode.SHADOW
        )
        schedule.refresh_from_db()
        client.force_login(owner_user)

        form = client.get(f"/dashboard/schedule/{schedule.id}/edit/")

        assert form.status_code == 200
        content = form.content.decode()
        assert "Canonical linked group" in content
        assert "name=\"group_name\" value=\"Legacy linked snapshot\"" in content
        assert "<select name=\"training_type_id\"" not in content
        assert "<select name=\"location_id\"" not in content
        assert f"name=\"training_type_id\" value=\"{training_type.id}\"" in content
        assert f"name=\"location_id\" value=\"{location.id}\"" in content

        response = client.post(
            f"/dashboard/schedule/{schedule.id}/edit/",
            {
                "group_name": schedule.group_name,
                "day_of_week": str(schedule.day_of_week),
                "start_time": "10:30",
                "end_time": "11:30",
                "trainer_id": str(trainer.id),
                "location_id": str(location.id),
                "training_type_id": str(training_type.id),
            },
        )

        assert response.status_code == 204
        schedule.refresh_from_db()
        assert schedule.start_time == time(10, 30)
        assert schedule.end_time == time(11, 30)
        assert schedule.group_name == "Legacy linked snapshot"

        response = client.post(
            f"/dashboard/schedule/{schedule.id}/edit/",
            {
                "group_name": schedule.group_name,
                "day_of_week": str(schedule.day_of_week),
                "start_time": "10:30",
                "end_time": "11:30",
                "trainer_id": str(trainer.id),
                "location_id": str(other_location.id),
                "training_type_id": str(training_type.id),
            },
        )

        assert response.status_code == 200
        assert "Use a future replacement slot" in response.content.decode()
        schedule.refresh_from_db()
        assert schedule.location_id == location.id

    def test_schedule_edit_rejects_unavailable_training_type(self, client: Client, owner_user, club, other_club):
        trainer = TrainerFactory(club=club)
        location = LocationFactory(club=club)
        current_type = TrainingTypeFactory(club=club)
        foreign_type = TrainingTypeFactory(club=other_club)
        inactive_type = TrainingTypeFactory(club=club, is_active=False)
        schedule = ScheduleFactory(club=club, trainer=trainer, location=location, training_type=current_type)
        client.force_login(owner_user)

        for training_type in (foreign_type, inactive_type):
            response = client.post(
                f"/dashboard/schedule/{schedule.id}/edit/",
                {
                    "group_name": f"Rejected {training_type.id}",
                    "day_of_week": str(schedule.day_of_week),
                    "start_time": "10:30",
                    "end_time": "11:30",
                    "trainer_id": str(trainer.id),
                    "location_id": str(location.id),
                    "training_type_id": str(training_type.id),
                },
            )

            assert response.status_code == 200
            schedule.refresh_from_db()
            assert schedule.training_type_id == current_type.id
            assert schedule.group_name != f"Rejected {training_type.id}"

    def test_schedule_renders(self, client: Client, owner_user, club):
        ScheduleFactory(club=club, day_of_week=0, group_name="Mon Boxing")
        client.force_login(owner_user)
        response = client.get("/dashboard/schedule/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "РАСПИСАНИЕ" in content
        assert "Mon Boxing" in content

    def test_schedule_htmx_partial(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/schedule/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert b"<aside" not in response.content

    def test_schedule_week_hides_old_slot_and_shows_rescheduled_slot(self, client: Client, owner_user, club):
        today = date.today()
        week_start = today - timedelta(days=today.weekday())
        original_date = week_start
        new_date = week_start + timedelta(days=2)
        schedule = ScheduleFactory(
            club=club,
            day_of_week=0,
            start_time=time(18, 0),
            end_time=time(19, 0),
            group_name="Moved PT",
        )
        ScheduleExceptionFactory(
            schedule=schedule,
            date=original_date,
            exception_type="rescheduled",
            new_date=new_date,
            new_start_time=time(20, 0),
            new_end_time=time(21, 0),
        )

        client.force_login(owner_user)
        response = client.get("/dashboard/schedule/")

        assert response.status_code == 200
        content = response.content.decode()
        assert f"/dashboard/schedule/{schedule.id}/detail/?date={original_date.isoformat()}" not in content
        assert f"/dashboard/schedule/{schedule.id}/detail/?date={new_date.isoformat()}" in content
        assert "20:00–21:00" in content
        assert "ПЕРЕНОС" in content

    def test_schedule_week_keeps_parallel_sessions_as_separate_cards(self, client: Client, owner_user, club):
        today = date.today()
        week_start = today - timedelta(days=today.weekday())
        target_date = week_start + timedelta(days=3)
        first_schedule = ScheduleFactory(
            club=club,
            day_of_week=3,
            start_time=time(19, 0),
            end_time=time(20, 0),
            group_name="Parallel PT 1",
        )
        second_schedule = ScheduleFactory(
            club=club,
            day_of_week=3,
            start_time=time(19, 0),
            end_time=time(20, 0),
            group_name="Parallel PT 2",
        )

        client.force_login(owner_user)
        response = client.get("/dashboard/schedule/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Parallel PT 1" in content
        assert "Parallel PT 2" in content
        assert f"/dashboard/schedule/{first_schedule.id}/detail/?date={target_date.isoformat()}" in content
        assert f"/dashboard/schedule/{second_schedule.id}/detail/?date={target_date.isoformat()}" in content

    def test_session_detail_renders(self, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        client.force_login(owner_user)
        today = date.today()
        response = client.get(f"/dashboard/schedule/{schedule.id}/detail/?date={today.isoformat()}")
        assert response.status_code == 200
        content = response.content.decode()
        assert schedule.group_name in content
        assert "ОТМЕНИТЬ ЗАНЯТИЕ" in content

    def test_session_detail_shows_schedule_enrollment_controls(self, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        target_schedule = ScheduleFactory(club=club, group_name="Transfer Target")
        enrolled = StudentFactory(club=club, first_name="Roster", last_name="Student")
        StudentFactory(club=club, first_name="Candidate", last_name="Student")
        enrollment = ScheduleEnrollment.objects.create(
            club=club,
            student=enrolled,
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=date.today(),
        )
        client.force_login(owner_user)

        response = client.get(f"/dashboard/schedule/{schedule.id}/detail/?date={date.today().isoformat()}")

        assert response.status_code == 200
        content = response.content.decode()
        assert "ЗАПИСЬ УЧЕНИКОВ" in content
        assert "Roster Student" in content
        assert "Candidate Student" in content
        assert 'name="student_id"' in content
        assert f"/dashboard/schedule/{schedule.id}/enroll/" in content
        assert f"/dashboard/schedule/enrollments/{enrollment.id}/transfer/" in content
        assert "ПЕРЕВЕСТИ" in content
        assert 'name="transfer_date"' in content
        assert f'<option value="{target_schedule.id}">' in content
        assert f'<option value="{schedule.id}">' not in content

    def test_session_detail_shows_checked_in_status_for_personal_booking(
        self,
        client: Client,
        owner_user,
        club,
    ):
        today = date.today()
        personal_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        schedule = ScheduleFactory(
            club=club,
            day_of_week=today.weekday(),
            one_time_date=today,
            group_name="Personal PT",
            training_type=personal_type,
        )
        student = StudentFactory(club=club, first_name="Checked", last_name="Personal")
        enrollment = ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=today,
            ends_on=today,
            created_from=ScheduleEnrollment.CreatedFrom.PERSONAL_BOOKING,
        )
        CheckinFactory(
            club=club,
            student=student,
            schedule=schedule,
            training_type=personal_type,
            trainer=schedule.trainer,
            location=schedule.location,
            date=today,
        )
        client.force_login(owner_user)

        response = client.get(f"/dashboard/schedule/{schedule.id}/detail/?date={today.isoformat()}")

        assert response.status_code == 200
        rows_by_enrollment_id = {
            row["enrollment"].id: row for row in response.context["enrollment_rows"]
        }
        assert rows_by_enrollment_id[enrollment.id]["is_checked_in"] is True
        content = response.content.decode()
        assert "Checked Personal" in content
        assert 'data-checkin-status="checked-in"' in content
        assert "ОТМЕЧЕН" in content

    def test_session_detail_ignores_cancelled_personal_checkin(
        self,
        client: Client,
        owner_user,
        club,
    ):
        today = date.today()
        personal_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        schedule = ScheduleFactory(
            club=club,
            day_of_week=today.weekday(),
            one_time_date=today,
            group_name="Cancelled Checkin PT",
            training_type=personal_type,
        )
        student = StudentFactory(club=club, first_name="Unchecked", last_name="Personal")
        enrollment = ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=today,
            ends_on=today,
            created_from=ScheduleEnrollment.CreatedFrom.PERSONAL_BOOKING,
        )
        CheckinFactory(
            club=club,
            student=student,
            schedule=schedule,
            training_type=personal_type,
            trainer=schedule.trainer,
            location=schedule.location,
            date=today,
            deleted_at=timezone.now(),
        )
        client.force_login(owner_user)

        response = client.get(f"/dashboard/schedule/{schedule.id}/detail/?date={today.isoformat()}")

        assert response.status_code == 200
        rows_by_enrollment_id = {
            row["enrollment"].id: row for row in response.context["enrollment_rows"]
        }
        assert rows_by_enrollment_id[enrollment.id]["is_checked_in"] is False
        content = response.content.decode()
        assert "Unchecked Personal" in content
        assert 'data-checkin-status="not-checked-in"' in content
        assert "НЕ ОТМЕЧЕН" in content

    def test_session_detail_shows_completed_drop_in_attendance_action_and_financial_preview(
        self,
        client: Client,
        owner_user,
        club,
    ):
        context = _create_admin_personal_drop_in(
            club=club,
            owner_user=owner_user,
            idempotency_key="htmx-attendance-preview",
        )
        _confirm_admin_drop_in_payment(
            club=club,
            owner_user=owner_user,
            booking=context["booking"],
            idempotency_key="htmx-attendance-preview-payment",
        )
        client.force_login(owner_user)

        with patch(
            "apps.attendance.selectors.timezone.now",
            return_value=context["ends_at"] + timedelta(minutes=5),
        ):
            response = client.get(
                f"/dashboard/schedule/{context['schedule'].id}/detail/"
                f"?date={context['schedule'].one_time_date.isoformat()}",
            )

        content = response.content.decode()
        assert response.status_code == 200
        assert "ПОСЕЩЕНИЕ" in content
        assert "ЗАЧЕСТЬ ПОСЕЩЕНИЕ" in content
        assert "Останется: 0" in content
        assert 'value="Тренер не отметил"' in content
        assert (
            f"/dashboard/schedule/personal-drop-ins/{context['booking'].id}/attendance/"
            in content
        )
        _assert_no_nested_forms(content)

    def test_session_detail_blocks_drop_in_attendance_before_session_end(
        self,
        client: Client,
        owner_user,
        club,
    ):
        context = _create_admin_personal_drop_in(
            club=club,
            owner_user=owner_user,
            idempotency_key="htmx-attendance-too-early",
        )
        client.force_login(owner_user)

        response = client.get(
            f"/dashboard/schedule/{context['schedule'].id}/detail/"
            f"?date={context['schedule'].one_time_date.isoformat()}",
        )

        content = response.content.decode()
        assert response.status_code == 200
        assert "ПОСЕЩЕНИЕ" in content
        assert "Посещение можно зачесть после окончания занятия" in content
        assert "ЗАЧЕСТЬ ПОСЕЩЕНИЕ" not in content

    def test_session_detail_defensively_blocks_frozen_drop_in_enrollment(
        self,
        client: Client,
        owner_user,
        club,
    ):
        context = _create_admin_personal_drop_in(
            club=club,
            owner_user=owner_user,
            idempotency_key="htmx-attendance-frozen-enrollment",
        )
        context["booking"].enrollment.status = ScheduleEnrollment.Status.FROZEN
        context["booking"].enrollment.save(update_fields=["status", "updated_at"])
        client.force_login(owner_user)

        with patch(
            "apps.attendance.selectors.timezone.now",
            return_value=context["ends_at"] + timedelta(minutes=5),
        ):
            response = client.get(
                f"/dashboard/schedule/{context['schedule'].id}/detail/"
                f"?date={context['schedule'].one_time_date.isoformat()}",
            )

        content = response.content.decode()
        assert response.status_code == 200
        assert "Запись ученика заморожена" in content
        assert "ЗАЧЕСТЬ ПОСЕЩЕНИЕ" not in content

    def test_one_time_personal_schedule_without_drop_in_booking_has_no_correction_action(
        self,
        client: Client,
        owner_user,
        club,
    ):
        target_date = date.today() - timedelta(days=1)
        personal_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        schedule = ScheduleFactory(
            club=club,
            day_of_week=target_date.weekday(),
            one_time_date=target_date,
            training_type=personal_type,
        )
        ScheduleEnrollment.objects.create(
            club=club,
            student=StudentFactory(club=club, status=Student.Status.ACTIVE),
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=target_date,
            ends_on=target_date,
            created_from=ScheduleEnrollment.CreatedFrom.PERSONAL_BOOKING,
        )
        client.force_login(owner_user)

        response = client.get(
            f"/dashboard/schedule/{schedule.id}/detail/?date={target_date.isoformat()}",
        )

        content = response.content.decode()
        assert response.status_code == 200
        assert "ЗАЧЕСТЬ ПОСЕЩЕНИЕ" not in content
        assert "data-personal-attendance-panel" not in content

    @pytest.mark.parametrize("role", ["owner", "admin"])
    def test_management_can_record_drop_in_attendance_from_session_detail(
        self,
        client: Client,
        owner_user,
        admin_user,
        club,
        role,
    ):
        context = _create_admin_personal_drop_in(
            club=club,
            owner_user=owner_user,
            idempotency_key=f"htmx-attendance-success-{role}",
        )
        payment = _confirm_admin_drop_in_payment(
            club=club,
            owner_user=owner_user,
            booking=context["booking"],
            idempotency_key=f"htmx-attendance-success-payment-{role}",
        )
        client.force_login(owner_user if role == "owner" else admin_user)
        completed_at = context["ends_at"] + timedelta(minutes=5)

        with (
            patch("apps.attendance.selectors.timezone.now", return_value=completed_at),
            patch("apps.attendance.services.checkin.timezone.now", return_value=completed_at),
            patch("apps.attendance.services.async_task"),
        ):
            response = client.post(
                f"/dashboard/schedule/personal-drop-ins/{context['booking'].id}/attendance/",
                {"reason": "  Тренер   не отметил  "},
                HTTP_HX_REQUEST="true",
            )

        context["booking"].refresh_from_db()
        payment.subscription.refresh_from_db()
        checkin = Checkin.objects.for_club(club).get(id=context["booking"].checkin_id)
        session = GroupSession.objects.for_club(club).get(
            schedule=context["schedule"],
            date=context["schedule"].one_time_date,
        )
        content = response.content.decode()
        assert response.status_code == 200
        assert "Посещение зачтено" in content
        assert 'data-attendance-correction-status="checked-in"' in content
        assert context["booking"].state == PersonalDropInBooking.State.ATTENDED
        assert checkin.source == Checkin.Source.BATCH
        assert payment.subscription.trainings_left == 0
        assert session.notes.endswith("Тренер не отметил")

    def test_drop_in_attendance_endpoint_is_post_only_and_validates_reason(
        self,
        client: Client,
        owner_user,
        club,
    ):
        context = _create_admin_personal_drop_in(
            club=club,
            owner_user=owner_user,
            idempotency_key="htmx-attendance-method-reason",
        )
        client.force_login(owner_user)
        path = (
            f"/dashboard/schedule/personal-drop-ins/{context['booking'].id}/attendance/"
        )

        get_response = client.get(path)
        completed_at = context["ends_at"] + timedelta(minutes=5)
        with (
            patch("apps.attendance.selectors.timezone.now", return_value=completed_at),
            patch("apps.attendance.services.checkin.timezone.now", return_value=completed_at),
            patch("apps.attendance.services.async_task"),
        ):
            post_response = client.post(path, {"reason": "   "}, HTTP_HX_REQUEST="true")

        assert get_response.status_code == 405
        assert post_response.status_code == 200
        assert "Укажите причину ручного подтверждения посещения" in post_response.content.decode()
        assert not Checkin.objects.for_club(club).exists()
        assert not GroupSession.objects.for_club(club).exists()

    def test_drop_in_attendance_endpoint_hides_foreign_booking_and_denies_trainer(
        self,
        client: Client,
        owner_user,
        trainer_user,
        club,
        other_club,
    ):
        local_context = _create_admin_personal_drop_in(
            club=club,
            owner_user=owner_user,
            idempotency_key="htmx-attendance-trainer-denied",
        )
        foreign_context = _create_admin_personal_drop_in(
            club=other_club,
            owner_user=owner_user,
            idempotency_key="htmx-attendance-foreign",
        )

        client.force_login(owner_user)
        foreign_response = client.post(
            f"/dashboard/schedule/personal-drop-ins/{foreign_context['booking'].id}/attendance/",
            {"reason": "Тренер не отметил"},
        )
        client.force_login(trainer_user)
        trainer_response = client.post(
            f"/dashboard/schedule/personal-drop-ins/{local_context['booking'].id}/attendance/",
            {"reason": "Тренер не отметил"},
        )

        assert foreign_response.status_code == 404
        assert trainer_response.status_code == 403
        assert not Checkin.objects.for_club(club).exists()
        assert not Checkin.objects.for_club(other_club).exists()

    def test_session_detail_explains_closed_payroll_block_for_paid_drop_in(
        self,
        client: Client,
        owner_user,
        club,
    ):
        context = _create_admin_personal_drop_in(
            club=club,
            owner_user=owner_user,
            idempotency_key="htmx-attendance-closed-payroll",
        )
        _confirm_admin_drop_in_payment(
            club=club,
            owner_user=owner_user,
            booking=context["booking"],
            idempotency_key="htmx-attendance-closed-payroll-payment",
        )
        close_trainer_payroll_period(
            club_id=club.id,
            period_start=context["schedule"].one_time_date,
            period_end=context["schedule"].one_time_date,
            reason="Выплаты закрыты",
            actor_user_id=owner_user.id,
        )
        client.force_login(owner_user)

        with patch(
            "apps.attendance.selectors.timezone.now",
            return_value=context["ends_at"] + timedelta(minutes=5),
        ):
            response = client.get(
                f"/dashboard/schedule/{context['schedule'].id}/detail/"
                f"?date={context['schedule'].one_time_date.isoformat()}",
            )

        content = response.content.decode()
        assert response.status_code == 200
        assert "Период выплат уже закрыт" in content
        assert "ЗАЧЕСТЬ ПОСЕЩЕНИЕ" not in content

    def test_session_detail_roster_marks_frozen_enrollment_with_blocked_alert(
        self,
        client: Client,
        owner_user,
        club,
    ):
        today = date.today()
        schedule = ScheduleFactory(club=club, day_of_week=today.weekday())
        target_schedule = ScheduleFactory(club=club, group_name="Transfer Target")
        student = StudentFactory(club=club, status="active", first_name="Frozen", last_name="Roster")
        enrollment = ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=schedule,
            status=ScheduleEnrollment.Status.FROZEN,
            starts_on=today - timedelta(days=7),
        )
        client.force_login(owner_user)

        response = client.get(f"/dashboard/schedule/{schedule.id}/detail/?date={today.isoformat()}")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Frozen Roster" in content
        assert "Заморожен" in content
        assert "Нельзя отметить" in content
        assert "Новичок" in content
        assert f"/dashboard/schedule/enrollments/{enrollment.id}/transfer/" in content
        assert f'<option value="{target_schedule.id}">' in content
        assert f"/dashboard/schedule/enrollments/{enrollment.id}/unfreeze/" in content
        assert f"/dashboard/schedule/enrollments/{enrollment.id}/cancel/" in content

    def test_session_detail_uses_effective_open_enrollment_for_status_and_actions(
        self,
        client: Client,
        owner_user,
        club,
    ):
        today = date.today()
        schedule = ScheduleFactory(club=club, day_of_week=today.weekday())
        student = StudentFactory(club=club, status="active", first_name="Reenrolled", last_name="Student")
        transferred = ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=schedule,
            status=ScheduleEnrollment.Status.TRANSFERRED,
            starts_on=today - timedelta(days=14),
            ends_on=today,
            created_from=ScheduleEnrollment.CreatedFrom.MANUAL,
        )
        active = ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=today,
            created_from=ScheduleEnrollment.CreatedFrom.MANUAL,
        )
        client.force_login(owner_user)

        response = client.get(f"/dashboard/schedule/{schedule.id}/detail/?date={today.isoformat()}")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Reenrolled Student" in content
        assert "Активен" in content
        assert f"/dashboard/schedule/enrollments/{active.id}/transfer/" in content
        assert f"/dashboard/schedule/enrollments/{active.id}/freeze/" in content
        assert f"/dashboard/schedule/enrollments/{active.id}/cancel/" in content
        assert f"/dashboard/schedule/enrollments/{transferred.id}/transfer/" not in content
        assert f"/dashboard/schedule/enrollments/{transferred.id}/freeze/" not in content
        assert f"/dashboard/schedule/enrollments/{transferred.id}/cancel/" not in content

    def test_session_detail_shows_enrollment_transfer_form(self, client: Client, owner_user, club, other_club):
        today = date.today()
        schedule = ScheduleFactory(club=club, group_name="Base Group")
        target_trainer = TrainerFactory(club=club, first_name="Target", last_name="Coach")
        target_location = LocationFactory(club=club, name="Blue Room")
        target_schedule = ScheduleFactory(
            club=club,
            day_of_week=today.weekday(),
            start_time=time(20, 0),
            end_time=time(21, 30),
            group_name="Advanced Grappling",
            trainer=target_trainer,
            location=target_location,
        )
        foreign_schedule = ScheduleFactory(club=other_club, group_name="Foreign Group")
        enrolled = StudentFactory(club=club, first_name="Transfer", last_name="Candidate")
        ScheduleEnrollment.objects.create(
            club=club,
            student=enrolled,
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=today - timedelta(days=7),
        )
        client.force_login(owner_user)

        response = client.get(f"/dashboard/schedule/{schedule.id}/detail/?date={today.isoformat()}")

        assert response.status_code == 200
        content = response.content.decode()
        assert "ПЕРЕВЕСТИ" in content
        assert 'name="target_schedule_id"' in content
        assert 'name="transfer_date"' in content
        assert "hx-confirm" in content
        assert "/dashboard/schedule/enrollments/" in content
        assert f'<option value="{target_schedule.id}">' in content
        assert "Advanced Grappling" in content
        assert "Target Coach" in content
        assert "Blue Room" in content
        assert "20:00" in content
        assert f'<option value="{schedule.id}">' not in content
        assert foreign_schedule.group_name not in content

    def test_session_enroll_adds_student_to_schedule(self, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        student = StudentFactory(club=club, first_name="New", last_name="Member")
        today = date.today()
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/{schedule.id}/enroll/",
            {
                "date": today.isoformat(),
                "student_id": str(student.id),
                "status": ScheduleEnrollment.Status.ACTIVE,
            },
        )

        assert response.status_code == 200
        enrollment = ScheduleEnrollment.objects.for_club(club).get(student=student, schedule=schedule)
        assert enrollment.status == ScheduleEnrollment.Status.ACTIVE
        assert enrollment.starts_on == today
        assert "New Member" in response.content.decode()

    def test_session_enroll_records_actor_on_mapped_group_membership_event(
        self,
        client: Client,
        owner_user,
        club,
        settings,
    ):
        from apps.attendance.models import TrainingGroupMembershipEvent, TrainingGroupRolloutState
        from apps.attendance.tests.factories import TrainingGroupFactory

        rollout_state = TrainingGroupRolloutState.objects.for_club(club).get()
        update_training_group_rollout_state_for_test(
            TrainingGroupRolloutState.objects.filter(id=rollout_state.id),
            mode=TrainingGroupRolloutState.Mode.SHADOW
        )
        settings.TRAINING_GROUP_NEW_WRITES_ENABLED = True
        trainer = TrainerFactory(club=club)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        location = LocationFactory(club=club)
        group = TrainingGroupFactory(
            club=club,
            training_type=training_type,
            location=location,
            responsible_trainer=trainer,
        )
        schedule = ScheduleFactory(
            club=club,
            training_group=group,
            trainer=trainer,
            training_type=training_type,
            location=location,
        )
        student = StudentFactory(club=club)
        today = date.today()
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/{schedule.id}/enroll/",
            {
                "date": today.isoformat(),
                "student_id": str(student.id),
                "status": ScheduleEnrollment.Status.ACTIVE,
            },
        )

        assert response.status_code == 200
        event = TrainingGroupMembershipEvent.objects.for_club(club).get(
            action="created",
            membership__student=student,
            membership__training_group=group,
        )
        assert event.actor_id == owner_user.id

    def test_session_enrollment_cancel_removes_student_for_selected_date(self, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        student = StudentFactory(club=club, first_name="Cancel", last_name="Member")
        today = date.today()
        enrollment = ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=today - timedelta(days=7),
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/enrollments/{enrollment.id}/cancel/",
            {"date": today.isoformat()},
        )

        assert response.status_code == 200
        enrollment.refresh_from_db()
        assert enrollment.status == ScheduleEnrollment.Status.CANCELLED
        assert enrollment.ends_on == today - timedelta(days=1)
        assert "Cancel Member" in response.content.decode()

    def test_session_enrollment_freeze_and_unfreeze(self, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        student = StudentFactory(club=club, first_name="Freeze", last_name="Member")
        enrollment = ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=date.today(),
        )
        client.force_login(owner_user)

        freeze_response = client.post(
            f"/dashboard/schedule/enrollments/{enrollment.id}/freeze/",
            {"date": date.today().isoformat()},
        )
        enrollment.refresh_from_db()
        unfreeze_response = client.post(
            f"/dashboard/schedule/enrollments/{enrollment.id}/unfreeze/",
            {"date": date.today().isoformat()},
        )

        assert freeze_response.status_code == 200
        assert "Заморожен" in freeze_response.content.decode()
        assert enrollment.status == ScheduleEnrollment.Status.FROZEN
        assert unfreeze_response.status_code == 200
        enrollment.refresh_from_db()
        assert enrollment.status == ScheduleEnrollment.Status.ACTIVE

    def test_session_enrollment_transfer_moves_student_without_gap_or_overlap(self, client: Client, owner_user, club):
        transfer_date = date.today()
        schedule = ScheduleFactory(club=club, group_name="Old Group")
        target_schedule = ScheduleFactory(club=club, group_name="New Group")
        student = StudentFactory(club=club, first_name="Move", last_name="Member")
        enrollment = ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=transfer_date - timedelta(days=14),
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/enrollments/{enrollment.id}/transfer/",
            {
                "date": transfer_date.isoformat(),
                "schedule_id": str(schedule.id),
                "target_schedule_id": str(target_schedule.id),
                "transfer_date": transfer_date.isoformat(),
                "confirm_transfer": "on",
            },
        )

        assert response.status_code == 200
        assert response.headers["HX-Trigger"] == "scheduleUpdated"
        enrollment.refresh_from_db()
        opened = ScheduleEnrollment.objects.for_club(club).get(
            student=student,
            schedule=target_schedule,
        )
        assert enrollment.status == ScheduleEnrollment.Status.TRANSFERRED
        assert enrollment.ends_on == transfer_date - timedelta(days=1)
        assert opened.status == ScheduleEnrollment.Status.ACTIVE
        assert opened.starts_on == transfer_date
        assert enrollment.ends_on + timedelta(days=1) == opened.starts_on

    def test_session_enrollment_transfer_rejects_same_schedule(self, client: Client, owner_user, club):
        transfer_date = date.today()
        schedule = ScheduleFactory(club=club)
        student = StudentFactory(club=club, first_name="Same", last_name="Schedule")
        enrollment = ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=transfer_date - timedelta(days=14),
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/enrollments/{enrollment.id}/transfer/",
            {
                "date": transfer_date.isoformat(),
                "schedule_id": str(schedule.id),
                "target_schedule_id": str(schedule.id),
                "transfer_date": transfer_date.isoformat(),
                "confirm_transfer": "on",
            },
        )

        assert response.status_code == 200
        assert response.headers["HX-Trigger"] == "scheduleUpdated"
        enrollment.refresh_from_db()
        assert enrollment.status == ScheduleEnrollment.Status.ACTIVE
        assert "Выберите другую группу" in response.content.decode()

    def test_session_enrollment_transfer_rejects_foreign_target_schedule(
        self,
        client: Client,
        owner_user,
        club,
        other_club,
    ):
        transfer_date = date.today()
        schedule = ScheduleFactory(club=club)
        foreign_schedule = ScheduleFactory(club=other_club, group_name="Foreign Group")
        student = StudentFactory(club=club, first_name="Foreign", last_name="Target")
        enrollment = ScheduleEnrollment.objects.create(
            club=club,
            student=student,
            schedule=schedule,
            status=ScheduleEnrollment.Status.ACTIVE,
            starts_on=transfer_date - timedelta(days=14),
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/enrollments/{enrollment.id}/transfer/",
            {
                "date": transfer_date.isoformat(),
                "schedule_id": str(schedule.id),
                "target_schedule_id": str(foreign_schedule.id),
                "transfer_date": transfer_date.isoformat(),
                "confirm_transfer": "on",
            },
        )

        assert response.status_code == 200
        assert response.headers["HX-Trigger"] == "scheduleUpdated"
        enrollment.refresh_from_db()
        assert enrollment.status == ScheduleEnrollment.Status.ACTIVE
        assert (
            not ScheduleEnrollment.objects.for_club(club)
            .filter(
                student=student,
                schedule_id=foreign_schedule.id,
            )
            .exists()
        )
        assert "Выберите расписание этого клуба" in response.content.decode()

    def test_session_detail_with_exception_hides_actions(self, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        today = date.today()
        ScheduleExceptionFactory(schedule=schedule, date=today, exception_type="cancelled")
        client.force_login(owner_user)
        response = client.get(f"/dashboard/schedule/{schedule.id}/detail/?date={today.isoformat()}")
        assert response.status_code == 200
        content = response.content.decode()
        assert "Занятие отменено" in content
        assert "ОТМЕНИТЬ ЗАНЯТИЕ" not in content

    def test_session_detail_for_moved_occurrence_uses_effective_slot_context(self, client: Client, owner_user, club):
        today = date.today()
        week_start = today - timedelta(days=today.weekday())
        original_date = week_start
        new_date = week_start + timedelta(days=2)
        schedule = ScheduleFactory(
            club=club,
            day_of_week=0,
            start_time=time(18, 0),
            end_time=time(19, 0),
            group_name="Moved Detail PT",
        )
        ScheduleExceptionFactory(
            schedule=schedule,
            date=original_date,
            exception_type="rescheduled",
            new_date=new_date,
            new_start_time=time(20, 0),
            new_end_time=time(21, 0),
        )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/schedule/{schedule.id}/detail/?date={new_date.isoformat()}")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Moved Detail PT" in content
        assert "20:00 — 21:00" in content
        assert "18:00 — 19:00" not in content
        assert f'name="date" value="{original_date.isoformat()}"' in content
        assert "Перенесено с" in content
        assert "ОТМЕНИТЬ ЗАНЯТИЕ" not in content

    def test_session_cancel(self, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        target_date = _future_schedule_date(schedule)
        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/schedule/{schedule.id}/cancel/",
            {"date": target_date.isoformat(), "reason": "Trainer sick"},
        )
        assert response.status_code == 200
        assert "Занятие отменено".encode() in response.content

    def test_session_cancel_with_checkins_renders_error_without_exception(self, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        target_date = _future_schedule_date(schedule)
        CheckinFactory(
            club=club,
            schedule=schedule,
            training_type=schedule.training_type,
            date=target_date,
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/{schedule.id}/cancel/",
            {"date": target_date.isoformat(), "reason": "Trainer sick"},
        )

        assert response.status_code == 200
        assert "уже есть отметки посещения".encode() in response.content
        assert "HX-Trigger" not in response.headers
        assert not ScheduleException.objects.filter(schedule=schedule, date=target_date).exists()

    def test_session_reschedule_triggers_schedule_refresh(self, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        target_date = _future_schedule_date(schedule)
        new_date = target_date + timedelta(days=1)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/{schedule.id}/reschedule/",
            {
                "date": target_date.isoformat(),
                "new_date": new_date.isoformat(),
                "new_time": "12:30",
            },
        )

        assert response.status_code == 200
        assert response.headers["HX-Trigger"] == "scheduleUpdated"
        assert ScheduleException.objects.filter(
            schedule=schedule,
            date=target_date,
            exception_type=ScheduleException.ExceptionType.RESCHEDULED,
            new_date=new_date,
        ).exists()

    def test_session_reschedule_with_checkins_renders_error_without_exception(self, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        target_date = _future_schedule_date(schedule)
        new_date = target_date + timedelta(days=1)
        CheckinFactory(
            club=club,
            schedule=schedule,
            training_type=schedule.training_type,
            date=target_date,
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/{schedule.id}/reschedule/",
            {
                "date": target_date.isoformat(),
                "new_date": new_date.isoformat(),
                "new_time": "12:30",
            },
        )

        assert response.status_code == 200
        assert "уже есть отметки посещения".encode() in response.content
        assert "HX-Trigger" not in response.headers
        assert not ScheduleException.objects.filter(schedule=schedule, date=target_date).exists()

    def test_session_substitute(self, client: Client, owner_user, club):
        from apps.trainers.tests.factories import TrainerFactory

        schedule = ScheduleFactory(club=club)
        sub_trainer = TrainerFactory(club=club, first_name="SubTrainer")
        target_date = _future_schedule_date(schedule)
        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/schedule/{schedule.id}/substitute/",
            {"date": target_date.isoformat(), "substitute_trainer_id": sub_trainer.id},
        )
        assert response.status_code == 200
        assert response.headers["HX-Trigger"] == "scheduleUpdated"
        assert "Замена тренера".encode() in response.content

    @patch("apps.attendance.services.checkin.async_task")
    def test_checkin_cancel_admin_triggers_schedule_refresh(self, _mock_async_task, client: Client, owner_user, club):
        schedule = ScheduleFactory(club=club)
        target_date = _future_schedule_date(schedule)
        checkin = CheckinFactory(
            club=club,
            schedule=schedule,
            training_type=schedule.training_type,
            date=target_date,
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/checkin/{checkin.id}/cancel/",
            {"schedule_id": str(schedule.id), "date": target_date.isoformat()},
        )

        assert response.status_code == 200
        assert response.headers["HX-Trigger"] == "scheduleUpdated"
        checkin.refresh_from_db()
        assert checkin.deleted_at is not None
        assert "Нет посещений" in response.content.decode()

    def test_session_exception_revert_restores_original_session(self, client: Client, owner_user, club):
        from apps.trainers.tests.factories import TrainerFactory

        schedule = ScheduleFactory(club=club)
        sub_trainer = TrainerFactory(club=club, first_name="RevertSub")
        target_date = _future_schedule_date(schedule)
        exception = ScheduleExceptionFactory(
            schedule=schedule,
            date=target_date,
            exception_type=ScheduleException.ExceptionType.SUBSTITUTE,
            substitute_trainer=sub_trainer,
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/schedule/{schedule.id}/revert/",
            {"date": target_date.isoformat()},
        )

        assert response.status_code == 200
        assert response.headers["HX-Trigger"] == "scheduleUpdated"
        assert not ScheduleException.objects.filter(id=exception.id).exists()
        content = response.content.decode()
        assert "ОТМЕНИТЬ ЗАНЯТИЕ" in content
        assert "ОТМЕНИТЬ ИЗМЕНЕНИЕ" not in content

    def test_schedule_tenant_isolation(self, client: Client, owner_user, club, other_club):
        ScheduleFactory(club=club, group_name="MyClubSession")
        ScheduleFactory(club=other_club, group_name="OtherClubSession")
        client.force_login(owner_user)
        response = client.get("/dashboard/schedule/")
        content = response.content.decode()
        assert "MyClubSession" in content
        assert "OtherClubSession" not in content

    def test_session_detail_tenant_isolation(self, client: Client, owner_user, club, other_club):
        other_schedule = ScheduleFactory(club=other_club)
        client.force_login(owner_user)
        response = client.get(f"/dashboard/schedule/{other_schedule.id}/detail/?date={date.today().isoformat()}")
        assert response.status_code == 404

    def test_unauthenticated_redirects(self, client: Client):
        response = client.get("/dashboard/schedule/")
        assert response.status_code == 302


class TestBillingViews:
    """ADMIN-06: Debtors with filters, payment verification."""

    def _create_personal_manual_review_order(self, *, settings, club, owner_user):
        settings.PAYMENT_PROVIDER = BankPaymentOrder.Provider.MOCK
        starts_at = timezone.now() + timedelta(days=14)
        ends_at = starts_at + timedelta(hours=1)
        trainer = TrainerFactory(club=club)
        location = LocationFactory(club=club)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        TrainerLocationFactory(club=club, trainer=trainer, location=location)
        student = StudentFactory(club=club, status=Student.Status.ACTIVE)
        tariff = TariffFactory(training_type=training_type)
        slot = PersonalAvailabilitySlot.objects.create(
            club=club,
            trainer=trainer,
            location=location,
            training_type=training_type,
            starts_at=starts_at,
            ends_at=ends_at,
            status=PersonalAvailabilitySlot.Status.PUBLISHED,
        )
        reservation = create_personal_booking_payment_reservation(
            club_id=club.id,
            student_id=student.id,
            trainer_id=trainer.id,
            starts_at=starts_at,
            ends_at=ends_at,
            location_id=location.id,
            training_type_id=training_type.id,
            tariff_id=tariff.id,
            availability_slot_id=slot.id,
            created_by_id=owner_user.id,
            source=BankPaymentOrder.Source.OWNER,
            idempotency_key=f"htmx-manual-review-{student.id}",
        )
        order = reservation.bank_payment_order
        order.status = BankPaymentOrder.Status.MANUAL_REVIEW
        order.last_error_code = "bank_payment_amount_mismatch"
        order.last_error_message = "Сумма provider webhook не совпадает с заказом"
        order.save(update_fields=["status", "last_error_code", "last_error_message", "updated_at"])
        reservation.status = "manual_review"
        reservation.last_error_code = order.last_error_code
        reservation.last_error_message = order.last_error_message
        reservation.save(update_fields=["status", "last_error_code", "last_error_message", "updated_at"])
        return order, reservation, slot

    def test_debtor_list_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/billing/")
        assert response.status_code == 200
        assert "ДОЛЖНИКИ".encode() in response.content

    def test_debtor_list_with_filters(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/billing/?student_status=active&date_from=2026-01-01")
        assert response.status_code == 200

    def test_debtor_list_htmx_partial(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/billing/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert b"<aside" not in response.content

    def test_payment_list_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/billing/payments/")
        assert response.status_code == 200
        assert "Верификация оплат".encode() in response.content

    def test_finance_workspace_separates_manual_online_and_history(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        manual_student = StudentFactory(club=club, first_name="ManualQueue")
        manual = PaymentFactory(
            club=club,
            tariff=TariffFactory(club=club),
            student=manual_student,
            payment_method=Payment.Method.CASH,
            status=Payment.Status.PENDING,
            recorded_by=owner_user,
        )
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        order.student.first_name = "OnlineReview"
        order.student.save(update_fields=["first_name", "updated_at"])
        history_student = StudentFactory(club=club, first_name="OnlineHistory")
        confirmed_online = PaymentFactory(
            club=club,
            tariff=TariffFactory(club=club),
            student=history_student,
            payment_method=Payment.Method.ONLINE,
            status=Payment.Status.CONFIRMED,
            recorded_by=owner_user,
            verified_by=owner_user,
            verified_at=timezone.now(),
        )
        client.force_login(owner_user)

        manual_page = client.get("/dashboard/billing/payments/").content.decode()
        online_page = client.get("/dashboard/billing/payments/?queue=online").content.decode()
        history_page = client.get("/dashboard/billing/payments/?queue=history").content.decode()

        assert "ManualQueue" in manual_page
        assert f"payments/{manual.id}/verify/" in manual_page
        assert "OnlineReview" not in manual_page
        assert "OnlineHistory" not in manual_page
        assert "Проблемы онлайн-оплаты" in online_page
        assert "OnlineReview" in online_page
        assert f"payments/{order.payment_id}/verify/" not in online_page
        assert "OnlineHistory" in history_page
        assert f"payments/{confirmed_online.id}/verify/" not in history_page

    def test_legacy_all_status_remains_read_only_and_shows_all_payment_states(
        self,
        client: Client,
        owner_user,
        club,
    ):
        pending_student = StudentFactory(club=club, first_name="LegacyPending")
        confirmed_student = StudentFactory(club=club, first_name="LegacyConfirmed")
        pending = PaymentFactory(
            club=club,
            student=pending_student,
            tariff=TariffFactory(club=club),
            payment_method=Payment.Method.CASH,
            status=Payment.Status.PENDING,
            recorded_by=owner_user,
        )
        PaymentFactory(
            club=club,
            student=confirmed_student,
            tariff=TariffFactory(club=club),
            payment_method=Payment.Method.ONLINE,
            status=Payment.Status.CONFIRMED,
            recorded_by=owner_user,
        )
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/?status=all")

        content = response.content.decode()
        assert response.status_code == 200
        assert "LegacyPending" in content
        assert "LegacyConfirmed" in content
        assert f"payments/{pending.id}/verify/" not in content
        assert "Ожидает" in content

    def test_finance_workspace_filters_and_safe_context_links(
        self,
        client: Client,
        owner_user,
        club,
    ):
        trainer = TrainerFactory(club=club, first_name="Filter", last_name="Trainer")
        other_trainer = TrainerFactory(club=club, first_name="Other", last_name="Trainer")
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        tariff = TariffFactory(club=club, training_type=training_type)
        schedule = ScheduleFactory(club=club, trainer=trainer, training_type=training_type)
        student = StudentFactory(club=club, first_name="ExactContext")
        target_date = date.today() + timedelta(days=7)
        payment = PaymentFactory(
            club=club,
            tariff=tariff,
            student=student,
            seller_trainer=trainer,
            target_schedule=schedule,
            target_start_date=target_date,
            target_group_name_snapshot="Exact group",
            payment_method=Payment.Method.TRANSFER,
            status=Payment.Status.PENDING,
            recorded_by=owner_user,
        )
        PaymentFactory(
            club=club,
            tariff=tariff,
            student=StudentFactory(club=club, first_name="WrongContext"),
            seller_trainer=other_trainer,
            payment_method=Payment.Method.CASH,
            status=Payment.Status.PENDING,
            recorded_by=owner_user,
        )
        client.force_login(owner_user)

        response = client.get(
            "/dashboard/billing/payments/",
            {
                "queue": "manual",
                "trainer_id": trainer.id,
                "method": Payment.Method.TRANSFER,
                "context": "group",
            },
        )

        content = response.content.decode()
        assert response.status_code == 200
        assert "ExactContext" in content
        assert "WrongContext" not in content
        assert f'/dashboard/students/{student.id}/card/' in content
        assert (
            f'/dashboard/schedule/{schedule.id}/detail/?date={target_date.isoformat()}'
            in content
        )
        resource_path = f'queue=manual&amp;payment_id={payment.id}#payment-{payment.id}'
        assert resource_path in content
        assert "Открыть операцию" in content

        resource_response = client.get(
            "/dashboard/billing/payments/",
            {"queue": "manual", "payment_id": payment.id},
        )
        assert resource_response.status_code == 200
        assert "ExactContext" in resource_response.content.decode()

    def test_finance_workspace_paginates_queue_and_preserves_scope(
        self,
        client: Client,
        owner_user,
        club,
    ):
        tariff = TariffFactory(club=club)
        for index in range(21):
            PaymentFactory(
                club=club,
                tariff=tariff,
                student=StudentFactory(club=club, first_name=f"Queue{index:02d}"),
                payment_method=Payment.Method.CASH,
                status=Payment.Status.PENDING,
                recorded_by=owner_user,
            )
        client.force_login(owner_user)

        first_page = client.get(
            "/dashboard/billing/payments/?queue=manual&method=cash"
        ).content.decode()
        second_page = client.get(
            "/dashboard/billing/payments/?queue=manual&method=cash&page=2"
        ).content.decode()

        assert "Дальше" in first_page
        assert "queue=manual&amp;method=cash&amp;page=2" in first_page
        assert "2 из 2" in second_page

    def test_finance_workspace_normalizes_method_when_switching_online_to_manual(
        self,
        client: Client,
        owner_user,
        club,
    ):
        payment = PaymentFactory(
            club=club,
            tariff=TariffFactory(club=club),
            student=StudentFactory(club=club, first_name="MethodBoundary"),
            payment_method=Payment.Method.CASH,
            status=Payment.Status.PENDING,
            recorded_by=owner_user,
        )
        client.force_login(owner_user)

        online_response = client.get(
            "/dashboard/billing/payments/",
            {"queue": "online", "method": Payment.Method.ONLINE},
        )
        manual_response = client.get(
            "/dashboard/billing/payments/",
            {"queue": "manual", "method": Payment.Method.ONLINE},
        )
        history_response = client.get(
            "/dashboard/billing/payments/",
            {"queue": "history", "method": Payment.Method.ONLINE},
        )

        assert online_response.status_code == 200
        assert 'href="/dashboard/billing/payments/?queue=manual"' in (
            online_response.content.decode()
        )
        assert manual_response.status_code == 200
        manual_content = manual_response.content.decode()
        assert f'id="payment-{payment.id}"' in manual_content
        assert '<option value="online"' not in manual_content
        assert history_response.status_code == 200
        assert '<option value="online" selected' in history_response.content.decode()

    def test_online_workspace_filters_and_paginates_refund_action_sections(
        self,
        client: Client,
        owner_user,
        club,
    ):
        trainer = TrainerFactory(club=club, first_name="RefundTarget")
        other_trainer = TrainerFactory(club=club, first_name="RefundOther")
        tariff = TariffFactory(club=club)

        def create_refund_case(*, label: str, seller, resolved: bool = False):
            student = StudentFactory(club=club, first_name=label)
            subscription = SubscriptionFactory(
                club=club,
                student=student,
                tariff=tariff,
                status=Subscription.Status.ACTIVE,
            )
            payment = PaymentFactory(
                club=club,
                student=student,
                tariff=tariff,
                subscription=subscription,
                seller_trainer=seller,
                payment_method=Payment.Method.ONLINE,
                status=Payment.Status.CONFIRMED,
                recorded_by=owner_user,
            )
            order = BankPaymentOrder.objects.create(
                club=club,
                payment=payment,
                subscription=subscription,
                student=student,
                provider=BankPaymentOrder.Provider.MOCK,
                source=BankPaymentOrder.Source.OWNER,
                status=BankPaymentOrder.Status.REFUNDED_PARTIALLY,
                amount_snapshot=payment.amount,
                purpose_snapshot=label,
                expires_at=timezone.now() + timedelta(days=1),
                created_by=owner_user,
            )
            refund_case = PaymentRefundCase.objects.create(
                club=club,
                order=order,
                refund_kind=PaymentRefundCase.Kind.PARTIAL,
                status=(
                    PaymentRefundCase.Status.RESOLVED
                    if resolved
                    else PaymentRefundCase.Status.RECONCILIATION_REQUIRED
                ),
                provider_refunded_at=timezone.now(),
                resolved_at=timezone.now() if resolved else None,
                resolved_by=owner_user if resolved else None,
            )
            return order, payment, subscription, refund_case

        for index in range(21):
            create_refund_case(label=f"RefundTarget{index:02d}", seller=trainer)
        create_refund_case(label="RefundOtherCase", seller=other_trainer)
        _, payment, subscription, refund_case = create_refund_case(
            label="PayrollTarget",
            seller=trainer,
            resolved=True,
        )
        PaymentRefund.objects.create(
            club=club,
            refund_case=refund_case,
            order=refund_case.order,
            payment=payment,
            subscription=subscription,
            approved_by=owner_user,
            amount=Decimal("100.00"),
            refund_kind=PaymentRefund.Kind.PARTIAL,
            accounting_date=date.today(),
            idempotency_key="finance-workspace-payroll-target",
            reason="Closed payroll period",
            entitlement_disposition=PaymentRefund.EntitlementDisposition.KEPT_PARTIAL,
            status=PaymentRefund.Status.PAYROLL_ACTION_REQUIRED,
        )
        _, other_payment, other_subscription, other_case = create_refund_case(
            label="PayrollOther",
            seller=other_trainer,
            resolved=True,
        )
        PaymentRefund.objects.create(
            club=club,
            refund_case=other_case,
            order=other_case.order,
            payment=other_payment,
            subscription=other_subscription,
            approved_by=owner_user,
            amount=Decimal("100.00"),
            refund_kind=PaymentRefund.Kind.PARTIAL,
            accounting_date=date.today(),
            idempotency_key="finance-workspace-payroll-other",
            reason="Closed payroll period",
            entitlement_disposition=PaymentRefund.EntitlementDisposition.KEPT_PARTIAL,
            status=PaymentRefund.Status.PAYROLL_ACTION_REQUIRED,
        )
        client.force_login(owner_user)

        first_page = client.get(
            "/dashboard/billing/payments/",
            {"queue": "online", "trainer_id": trainer.id},
        ).content.decode()
        second_page = client.get(
            "/dashboard/billing/payments/",
            {"queue": "online", "trainer_id": trainer.id, "refund_page": 2},
        ).content.decode()
        future_page = client.get(
            "/dashboard/billing/payments/",
            {
                "queue": "online",
                "trainer_id": trainer.id,
                "date_from": (date.today() + timedelta(days=1)).isoformat(),
            },
        ).content.decode()

        assert "RefundTarget00" in first_page
        assert "RefundTarget20" not in first_page
        assert "RefundTarget20" in second_page
        assert "RefundOtherCase" not in first_page
        assert "PayrollTarget" in first_page
        assert "PayrollOther" not in first_page
        assert (
            f"trainer_id={trainer.id}&amp;method=online&amp;refund_page=2"
            in first_page
        )
        assert "Проблемы онлайн-оплаты</span><span>24</span>" in first_page
        assert "RefundTarget00" not in future_page
        assert "PayrollTarget" not in future_page

    def test_finance_workspace_is_owner_admin_only_and_tenant_scoped(
        self,
        client: Client,
        admin_user,
        trainer_user,
        owner_user,
        club,
        other_club,
    ):
        PaymentFactory(
            club=club,
            tariff=TariffFactory(club=club),
            student=StudentFactory(club=club, first_name="LocalFinance"),
            status=Payment.Status.PENDING,
            recorded_by=owner_user,
        )
        PaymentFactory(
            club=other_club,
            tariff=TariffFactory(club=other_club),
            student=StudentFactory(club=other_club, first_name="ForeignFinance"),
            status=Payment.Status.PENDING,
            recorded_by=owner_user,
        )

        client.force_login(admin_user)
        admin_response = client.get("/dashboard/billing/payments/")
        client.force_login(trainer_user)
        trainer_response = client.get("/dashboard/billing/payments/")

        assert admin_response.status_code == 200
        assert "LocalFinance" in admin_response.content.decode()
        assert "ForeignFinance" not in admin_response.content.decode()
        assert trainer_response.status_code == 403

    def test_payment_list_shows_pending(self, client: Client, owner_user, club):
        from apps.billing.tests.factories import PaymentFactory

        PaymentFactory(club=club, status="pending")
        client.force_login(owner_user)
        response = client.get("/dashboard/billing/payments/")
        assert response.status_code == 200
        assert "Подтвердить".encode() in response.content
        assert "Отклонить".encode() in response.content
        assert b'name="rejection_reason"' in response.content

    def test_payment_list_renders_applied_discount_breakdown(self, client: Client, owner_user, club):
        payment = PaymentFactory(
            club=club,
            amount=Decimal("4500"),
            original_amount=Decimal("5000"),
        )
        discount = DiscountFactory(
            club=club,
            name="Семейная скидка",
            discount_type="percent",
            value=Decimal("10"),
        )
        payment.applied_discounts.add(discount)
        discount.name = "Семейная программа"
        discount.value = Decimal("20")
        discount.save(update_fields=["name", "value"])
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Базовая сумма: 5 000₽" in content
        assert "Скидка: Семейная программа" in content
        assert "Фактически вычтено: 500₽" in content
        assert "Итог: 4 500₽" in content
        assert "20%" not in content

    def test_payment_list_does_not_render_foreign_club_discount_relation(
        self,
        client: Client,
        owner_user,
        club,
        other_club,
    ):
        payment = PaymentFactory(club=club, amount=Decimal("4500"), original_amount=Decimal("5000"))
        foreign_discount = DiscountFactory(
            club=other_club,
            name="Чужая скидка",
            discount_type="fixed",
            value=Decimal("500"),
        )
        payment.applied_discounts.add(foreign_discount)
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Чужая скидка" not in content
        assert "Базовая сумма:" not in content

    def test_manual_payment_queue_excludes_online_pending_payment(self, client: Client, owner_user, club):
        from apps.billing.tests.factories import PaymentFactory

        payment = PaymentFactory(
            club=club,
            status=Payment.Status.PENDING,
            payment_method=Payment.Method.ONLINE,
            recorded_by=owner_user,
        )
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/")

        assert response.status_code == 200
        content = response.content.decode()
        assert str(payment.student) not in content
        assert f'/dashboard/billing/payments/{payment.id}/verify/' not in content
        assert "Да, подтвердить" not in content
        assert "Да, отклонить" not in content

    def test_payment_list_shows_canonical_group_identity_for_owner_verification(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.attendance.models import TrainingGroupMembership
        from apps.attendance.tests.factories import TrainingGroupFactory

        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        trainer = TrainerFactory(club=club, first_name="Responsible", last_name="Trainer")
        group = TrainingGroupFactory(
            club=club,
            name="Owner canonical group",
            training_type=training_type,
            responsible_trainer=trainer,
        )
        schedule = ScheduleFactory(
            club=club,
            training_group=group,
            training_type=training_type,
            trainer=trainer,
            start_time=time(18, 30),
            end_time=time(19, 45),
        )
        student = StudentFactory(club=club)
        membership = TrainingGroupMembership.objects.create(
            club=club,
            student=student,
            training_group=group,
            starts_on=date(2026, 7, 1),
            source=TrainingGroupMembership.Source.MANUAL,
        )
        payment = PaymentFactory(
            club=club,
            student=student,
            status=Payment.Status.PENDING,
            target_schedule=schedule,
            target_training_group=group,
            target_group_membership=membership,
            target_start_date=date(2026, 7, 8),
            target_group_name_snapshot=group.name,
            target_trainer_name_snapshot=str(trainer),
            target_location_name_snapshot=group.location.name,
            seller_trainer=trainer,
            recorded_by=owner_user,
        )
        group.name = "Renamed after payment"
        group.save(update_fields=["name"])
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Каноническая группа" in content
        assert "Owner canonical group" in content
        assert "Renamed after payment" not in content
        assert "Тренер занятия" in content
        assert "Время: 18:30–19:45" in content
        assert "Ответственный" in content
        assert f"Каноническое членство #{membership.id}" in content
        assert f"Принял: {owner_user.get_full_name() or owner_user.username}" in content
        assert payment.id

    def test_manual_payment_queue_shows_exact_personal_booking_context(
        self,
        client: Client,
        owner_user,
        club,
    ):
        context = _create_admin_personal_drop_in(
            club=club,
            owner_user=owner_user,
            idempotency_key="finance-workspace-personal-booking",
        )
        with patch("apps.attendance.services.async_task"), patch("django_q.tasks.async_task"):
            link = create_personal_drop_in_payment(
                club_id=club.id,
                booking_id=context["booking"].id,
                payment_method=Payment.Method.CASH,
                created_by_id=owner_user.id,
                idempotency_key="finance-workspace-personal-payment",
            )
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/?queue=manual")

        content = response.content.decode()
        schedule = context["schedule"]
        assert response.status_code == 200
        assert f'id="payment-{link.payment_id}"' in content
        assert "Точная персональная запись" in content
        assert context["starts_at"].strftime("%d.%m.%Y") in content
        assert "10:00–11:00" in content
        assert str(schedule.trainer) in content
        assert str(schedule.location) in content
        assert str(schedule.training_type) in content
        assert (
            f"/dashboard/schedule/{schedule.id}/detail/"
            f"?date={context['booking'].enrollment.starts_on.isoformat()}"
            in content
        )

    def test_online_review_shows_exact_group_context_and_safe_provider_evidence(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.attendance.tests.factories import TrainingGroupFactory

        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        trainer = TrainerFactory(club=club, first_name="Provider", last_name="Trainer")
        location = LocationFactory(club=club, name="Provider Hall")
        group = TrainingGroupFactory(
            club=club,
            name="Provider Exact Group",
            training_type=training_type,
            location=location,
            responsible_trainer=trainer,
        )
        schedule = ScheduleFactory(
            club=club,
            training_group=group,
            training_type=training_type,
            trainer=trainer,
            location=location,
            start_time=time(20, 15),
            end_time=time(21, 30),
        )
        target_date = _future_schedule_date(schedule)
        student = StudentFactory(club=club, first_name="ProviderEvidence")
        tariff = TariffFactory(club=club, training_type=training_type)
        subscription = SubscriptionFactory(
            club=club,
            student=student,
            tariff=tariff,
            status=Subscription.Status.PENDING,
        )
        payment = PaymentFactory(
            club=club,
            student=student,
            tariff=tariff,
            subscription=subscription,
            payment_method=Payment.Method.ONLINE,
            status=Payment.Status.PENDING,
            target_schedule=schedule,
            target_training_group=group,
            target_start_date=target_date,
            target_group_name_snapshot=group.name,
            target_location_name_snapshot=location.name,
            target_trainer_name_snapshot=str(trainer),
            seller_trainer=trainer,
            recorded_by=owner_user,
        )
        order = BankPaymentOrder.objects.create(
            club=club,
            payment=payment,
            subscription=subscription,
            student=student,
            provider=BankPaymentOrder.Provider.MOCK,
            source=BankPaymentOrder.Source.OWNER,
            status=BankPaymentOrder.Status.MANUAL_REVIEW,
            amount_snapshot=payment.amount,
            purpose_snapshot="Generic immutable provider purpose",
            expires_at=timezone.now() + timedelta(days=1),
            created_by=owner_user,
        )
        operation_id = "provider-operation-reference-1234567890"
        BankPaymentProviderEvent.objects.create(
            club=club,
            order=order,
            provider=BankPaymentOrder.Provider.MOCK,
            event_type="acquiringInternetPayment",
            provider_event_id="provider-evidence-event",
            provider_operation_id=operation_id,
            provider_status="APPROVED",
            normalized_status_snapshot="approved",
            received_at=timezone.now(),
            processed_at=timezone.now(),
            processing_status=BankPaymentProviderEvent.ProcessingStatus.FAILED,
            failure_code="private_failure_code",
            failure_message="private provider failure detail",
            redacted_payload_metadata={"safe": "must-not-render"},
        )
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/?queue=online")

        content = response.content.decode()
        fragment = content.split(
            f'data-manual-review-order-id="{order.id}"',
            maxsplit=1,
        )[1].split("</section>", maxsplit=1)[0]
        assert response.status_code == 200
        assert "Групповое обучение" in fragment
        assert group.name in fragment
        assert target_date.strftime("%d.%m.%Y") in fragment
        assert "Время: 20:15–21:30" in fragment
        assert location.name in fragment
        assert str(trainer) in fragment
        assert (
            f"/dashboard/schedule/{schedule.id}/detail/?date={target_date.isoformat()}"
            in fragment
        )
        assert "Провайдер: Тестовый провайдер" in fragment
        assert "Статус банка: Оплата подтверждена" in fragment
        assert operation_id[-12:] in fragment
        assert operation_id not in fragment
        assert "private_failure_code" not in content
        assert "private provider failure detail" not in content
        assert "must-not-render" not in content

    def test_payment_list_shows_bank_manual_review_queue_for_personal_reservation(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        order.provider = BankPaymentOrder.Provider.TOCHKA
        order.save(update_fields=["provider", "updated_at"])
        immutable_purpose = order.purpose_snapshot
        order.payment.tariff.name = "Переименованный после создания тариф"
        order.payment.tariff.save(update_fields=["name", "updated_at"])
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/?queue=online")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Спорные онлайн-оплаты" in content
        assert "bank_payment_amount_mismatch" not in content
        assert "Сумма provider webhook" not in content
        assert "Автоматическая проверка не завершена" in content
        manual_review_fragment = content.split(
            f'data-manual-review-order-id="{order.id}"',
            maxsplit=1,
        )[1].split("</section>", maxsplit=1)[0]
        assert immutable_purpose in manual_review_fragment
        assert "Переименованный после создания тариф" not in manual_review_fragment
        assert "Срок ссылки до" in manual_review_fragment
        assert "Персональная бронь" in content
        assert str(order.personal_payment_reservation.location) in manual_review_fragment
        assert str(order.personal_payment_reservation.training_type) in manual_review_fragment
        assert "Подтверждённых данных провайдера нет" in manual_review_fragment
        assert f"/dashboard/billing/bank-payment-orders/{order.id}/reconcile/" in content
        assert "СВЕРИТЬ С БАНКОМ" in content
        assert 'value="confirm_paid"' not in content
        assert 'value="reject"' not in content
        assert f"queue=online&amp;review_order={order.id}#bank-order-{order.id}" in content

        resource_response = client.get(
            "/dashboard/billing/payments/",
            {"queue": "online", "review_order": order.id},
        )
        assert resource_response.status_code == 200
        assert f'id="bank-order-{order.id}"' in resource_response.content.decode()

    def test_live_bank_order_deep_link_keeps_manual_review_backlog_visible(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        review_order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        from apps.billing.services import create_bank_payment_order

        settings.PAYMENT_PROVIDER = BankPaymentOrder.Provider.MOCK
        settings.MOCK_PAYMENT_ORDER_CREATION_ENABLED = True
        settings.ONLINE_PAYMENT_ORDER_CREATION_ENABLED = True
        settings.MOCK_PAYMENT_BASE_URL = "https://pay.example.test"
        live_order = create_bank_payment_order(
            club_id=club.id,
            student_id=StudentFactory(club=club).id,
            tariff_id=TariffFactory(
                club=club,
                training_type=TrainingTypeFactory(
                    club=club,
                    kind=TrainingType.Kind.GROUP,
                ),
            ).id,
            source=BankPaymentOrder.Source.OWNER,
            created_by_id=owner_user.id,
        )
        client.force_login(owner_user)

        response = client.get(
            "/dashboard/billing/payments/",
            {"bank_order": live_order.id},
        )

        content = response.content.decode()
        assert response.status_code == 200
        assert f'data-bank-payment-order-id="{live_order.id}"' in content
        assert f'data-manual-review-order-id="{review_order.id}"' in content

    def test_bank_payment_order_reconcile_action_enqueues_provider_check(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        order.provider = BankPaymentOrder.Provider.TOCHKA
        order.provider_operation_id = "operation-htmx-manual-review"
        order.save(update_fields=["provider", "provider_operation_id", "updated_at"])
        from apps.billing.service_modules.provider_events import request_provider_reconciliation

        attempt = request_provider_reconciliation(
            club_id=club.id,
            order_id=order.id,
            provider_event_id=None,
        )
        attempt.status = BankPaymentReconciliationAttempt.Status.MANUAL_REVIEW
        attempt.last_error_code = order.last_error_code
        attempt.retry_at = None
        attempt.save(update_fields=["status", "last_error_code", "retry_at", "updated_at"])
        client.force_login(owner_user)

        with patch(
            "apps.billing.service_modules.provider_events.enqueue_provider_reconciliation"
        ) as enqueue_reconciliation:
            response = client.post(
                f"/dashboard/billing/bank-payment-orders/{order.id}/reconcile/",
                HTTP_HX_REQUEST="true",
            )

        assert response.status_code == 204
        assert response.headers["HX-Redirect"] == "/dashboard/billing/payments/?queue=online"
        attempt.refresh_from_db()
        assert attempt.status == BankPaymentReconciliationAttempt.Status.PENDING
        assert attempt.last_error_code == ""
        assert (
            BankPaymentOrderReviewEvent.objects.for_club(club)
            .filter(
                order=order,
                actor=owner_user,
                resolution=BankPaymentOrderReviewEvent.Resolution.RETRY_RECONCILIATION,
            )
            .exists()
        )
        enqueue_reconciliation.assert_called_once_with(club_id=club.id, order_id=order.id)

    def test_bank_payment_order_reconcile_action_rejects_non_tochka_order(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        client.force_login(owner_user)

        with (
            patch(
                "apps.billing.service_modules.provider_events.request_provider_reconciliation"
            ) as request_reconciliation,
            patch(
                "apps.billing.service_modules.provider_events.enqueue_provider_reconciliation"
            ) as enqueue_reconciliation,
        ):
            response = client.post(
                f"/dashboard/billing/bank-payment-orders/{order.id}/reconcile/",
                HTTP_HX_REQUEST="true",
            )

        assert response.status_code == 400
        assert "Сверка с банком недоступна" in response.content.decode()
        request_reconciliation.assert_not_called()
        enqueue_reconciliation.assert_not_called()

    def test_bank_payment_order_reconcile_action_is_tenant_scoped(
        self,
        client: Client,
        settings,
        owner_user,
        club,
        other_club,
    ):
        other_order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=other_club,
            owner_user=owner_user,
        )
        other_order.provider = BankPaymentOrder.Provider.TOCHKA
        other_order.save(update_fields=["provider", "updated_at"])
        client.force_login(owner_user)

        with (
            patch(
                "apps.billing.service_modules.provider_events.request_provider_reconciliation"
            ) as request_reconciliation,
            patch(
                "apps.billing.service_modules.provider_events.enqueue_provider_reconciliation"
            ) as enqueue_reconciliation,
        ):
            response = client.post(
                f"/dashboard/billing/bank-payment-orders/{other_order.id}/reconcile/",
                HTTP_HX_REQUEST="true",
            )

        assert response.status_code == 404
        request_reconciliation.assert_not_called()
        enqueue_reconciliation.assert_not_called()

    def test_payment_list_shows_safe_staff_actions_for_live_bank_order(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        from apps.billing.services import create_bank_payment_order
        from apps.billing.tests.factories import TariffFactory

        settings.PAYMENT_PROVIDER = BankPaymentOrder.Provider.MOCK
        settings.MOCK_PAYMENT_ORDER_CREATION_ENABLED = True
        settings.ONLINE_PAYMENT_ORDER_CREATION_ENABLED = True
        settings.MOCK_PAYMENT_BASE_URL = "https://pay.example.test"
        student = StudentFactory(club=club)
        tariff = TariffFactory(
            training_type=TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        )
        order = create_bank_payment_order(
            club_id=club.id,
            student_id=student.id,
            tariff_id=tariff.id,
            source=BankPaymentOrder.Source.OWNER,
            created_by_id=owner_user.id,
        )
        client.force_login(owner_user)

        queue_response = client.get("/dashboard/billing/payments/?queue=online")
        response = client.get(f"/dashboard/billing/payments/?bank_order={order.id}")

        assert "Активные ссылки СБП" not in queue_response.content.decode()
        assert response.status_code == 200
        content = response.content.decode()
        assert "Активные ссылки СБП" in content
        assert "ОТПРАВИТЬ ССЫЛКУ" in content
        assert "СКОПИРОВАТЬ ССЫЛКУ" in content
        assert "ОТКРЫТЬ ПРЕДПРОСМОТР" in content
        assert order.provider_payment_url not in re.sub(r'(href|data-payment-url)="[^"]+"', '', content)

    def test_bank_payment_order_review_reject_action_releases_personal_reservation(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, reservation, slot = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/billing/bank-payment-orders/{order.id}/review/",
            {
                "resolution": BankPaymentOrderReviewEvent.Resolution.REJECT,
                "reason": "Выписка не подтверждает оплату",
            },
            HTTP_HX_REQUEST="true",
        )

        order.refresh_from_db()
        order.payment.refresh_from_db()
        reservation.refresh_from_db()
        slot.refresh_from_db()
        assert response.status_code == 204
        assert response.headers["HX-Redirect"] == "/dashboard/billing/payments/?queue=online"
        assert order.status == BankPaymentOrder.Status.FAILED
        assert order.payment.status == Payment.Status.REJECTED
        assert reservation.status == "cancelled"
        assert slot.status == PersonalAvailabilitySlot.Status.PUBLISHED
        assert BankPaymentOrderReviewEvent.objects.for_club(club).filter(
            order=order,
            resolution=BankPaymentOrderReviewEvent.Resolution.REJECT,
        ).exists()

    def test_duplicate_online_review_action_is_safe_and_does_not_duplicate_audit(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        client.force_login(owner_user)
        payload = {
            "resolution": BankPaymentOrderReviewEvent.Resolution.REJECT,
            "reason": "Проверено по выписке",
        }

        first = client.post(
            f"/dashboard/billing/bank-payment-orders/{order.id}/review/?queue=online",
            payload,
            HTTP_HX_REQUEST="true",
        )
        second = client.post(
            f"/dashboard/billing/bank-payment-orders/{order.id}/review/?queue=online",
            payload,
            HTTP_HX_REQUEST="true",
        )
        mismatch = client.post(
            f"/dashboard/billing/bank-payment-orders/{order.id}/review/?queue=online",
            {**payload, "reason": "Другая причина"},
            HTTP_HX_REQUEST="true",
        )

        order.refresh_from_db()
        assert first.status_code == 204
        assert second.status_code == 204
        assert mismatch.status_code == 400
        assert order.status == BankPaymentOrder.Status.FAILED
        assert (
            BankPaymentOrderReviewEvent.objects.for_club(club)
            .filter(order=order, resolution=BankPaymentOrderReviewEvent.Resolution.REJECT)
            .count()
            == 1
        )

    def test_bank_payment_order_review_confirm_action_requires_visible_reason(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/billing/bank-payment-orders/{order.id}/review/",
            {"resolution": BankPaymentOrderReviewEvent.Resolution.CONFIRM_PAID},
            HTTP_HX_REQUEST="true",
        )

        order.refresh_from_db()
        assert response.status_code == 400
        assert "Укажите комментарий".encode() in response.content
        assert order.status == BankPaymentOrder.Status.MANUAL_REVIEW

    def test_bank_payment_order_review_confirm_action_books_personal_reservation(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, reservation, slot = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/billing/bank-payment-orders/{order.id}/review/",
            {
                "resolution": BankPaymentOrderReviewEvent.Resolution.CONFIRM_PAID,
                "reason": "Проверено по выписке",
            },
            HTTP_HX_REQUEST="true",
        )

        order.refresh_from_db()
        order.payment.refresh_from_db()
        reservation.refresh_from_db()
        slot.refresh_from_db()
        assert response.status_code == 204
        assert order.status == BankPaymentOrder.Status.APPROVED
        assert order.payment.status == Payment.Status.CONFIRMED
        assert reservation.status == "booked"
        assert slot.status == PersonalAvailabilitySlot.Status.BOOKED

    def test_payment_list_shows_refund_review_actions(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        order.payment.status = Payment.Status.CONFIRMED
        order.payment.save(update_fields=["status", "updated_at"])
        order.last_error_code = "bank_payment_refunded_partially_requires_review"
        order.last_error_message = "Provider reported partial refund"
        order.save(update_fields=["last_error_code", "last_error_message", "updated_at"])
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/?queue=online")

        assert response.status_code == 200
        content = response.content.decode()
        assert 'value="mark_refunded_partially"' in content
        assert 'name="refund_amount"' in content
        assert "ЧАСТИЧНЫЙ ВОЗВРАТ" in content
        assert 'value="confirm_paid"' not in content

    def test_bank_payment_order_review_partial_refund_action_records_review_event(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        order.payment.status = Payment.Status.CONFIRMED
        order.payment.save(update_fields=["status", "updated_at"])
        order.last_error_code = "bank_payment_refunded_partially_requires_review"
        order.save(update_fields=["last_error_code", "updated_at"])
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/billing/bank-payment-orders/{order.id}/review/",
            {
                "resolution": BankPaymentOrderReviewEvent.Resolution.MARK_REFUNDED_PARTIALLY,
                "reason": "Частичный возврат в provider",
                "refund_amount": "1000.00",
            },
            HTTP_HX_REQUEST="true",
        )

        order.refresh_from_db()
        assert response.status_code == 204
        assert order.status == BankPaymentOrder.Status.REFUNDED_PARTIALLY
        event = BankPaymentOrderReviewEvent.objects.for_club(club).get(order=order)
        refund = PaymentRefund.objects.for_club(club).get(order=order)
        assert event.resolution == BankPaymentOrderReviewEvent.Resolution.MARK_REFUNDED_PARTIALLY
        assert event.evidence_metadata["refund_amount"] == "1000.00"
        assert event.evidence_metadata["refund_id"] == str(refund.id)
        assert refund.amount == Decimal("1000.00")

    def test_bank_payment_order_review_full_refund_action_records_review_event(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        order.payment.status = Payment.Status.CONFIRMED
        order.payment.save(update_fields=["status", "updated_at"])
        order.last_error_code = "bank_payment_refunded_requires_review"
        order.save(update_fields=["last_error_code", "updated_at"])
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/billing/bank-payment-orders/{order.id}/review/",
            {
                "resolution": BankPaymentOrderReviewEvent.Resolution.MARK_REFUNDED,
                "reason": "Полный возврат в provider",
                "entitlement_action": PaymentRefund.EntitlementDisposition.KEEP_CLUB_ABSORBS,
            },
            HTTP_HX_REQUEST="true",
        )

        order.refresh_from_db()
        assert response.status_code == 204
        assert order.status == BankPaymentOrder.Status.REFUNDED
        assert order.last_error_code == ""
        assert order.last_error_message == "Полный возврат в provider"
        event = BankPaymentOrderReviewEvent.objects.for_club(club).get(order=order)
        refund = PaymentRefund.objects.for_club(club).get(order=order)
        assert event.resolution == BankPaymentOrderReviewEvent.Resolution.MARK_REFUNDED
        assert event.reason == "Полный возврат в provider"
        assert event.evidence_metadata["refund_id"] == str(refund.id)
        assert refund.entitlement_disposition == PaymentRefund.EntitlementDisposition.KEEP_CLUB_ABSORBS

    def test_bank_payment_order_review_full_refund_requires_entitlement_action(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        order.payment.status = Payment.Status.CONFIRMED
        order.payment.save(update_fields=["status", "updated_at"])
        order.last_error_code = "bank_payment_refunded_requires_review"
        order.save(update_fields=["last_error_code", "updated_at"])
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/billing/bank-payment-orders/{order.id}/review/",
            {
                "resolution": BankPaymentOrderReviewEvent.Resolution.MARK_REFUNDED,
                "reason": "Полный возврат без решения по правам",
            },
            HTTP_HX_REQUEST="true",
        )

        order.refresh_from_db()
        assert response.status_code == 400
        assert order.status == BankPaymentOrder.Status.MANUAL_REVIEW
        assert not PaymentRefund.objects.for_club(club).filter(order=order).exists()

    def test_payment_list_keeps_legacy_refund_case_actionable_after_order_left_manual_review(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        order.payment.status = Payment.Status.CONFIRMED
        order.payment.save(update_fields=["status", "updated_at"])
        order.status = BankPaymentOrder.Status.REFUNDED_PARTIALLY
        order.save(update_fields=["status", "updated_at"])
        refund_case = PaymentRefundCase.objects.create(
            club=club,
            order=order,
            refund_kind=PaymentRefundCase.Kind.PARTIAL,
            status=PaymentRefundCase.Status.RECONCILIATION_REQUIRED,
            provider_refunded_at=timezone.now(),
        )
        client.force_login(owner_user)

        page = client.get("/dashboard/billing/payments/?queue=online")
        response = client.post(
            f"/dashboard/billing/payment-refund-cases/{refund_case.id}/approve/",
            {
                "refund_kind": PaymentRefund.Kind.PARTIAL,
                "amount": "100.00",
                "reason": "Сверено по старой выписке",
            },
            HTTP_HX_REQUEST="true",
        )

        assert page.status_code == 200
        content = page.content.decode()
        assert "Возвраты без бухгалтерского завершения" in content
        assert f"payment-refund-cases/{refund_case.id}/approve" in content
        assert response.status_code == 204
        refund_case.refresh_from_db()
        assert refund_case.status == PaymentRefundCase.Status.RESOLVED
        assert PaymentRefund.objects.for_club(club).filter(refund_case=refund_case).exists()

    def test_payment_list_requires_explicit_open_date_for_closed_payroll_refund(
        self,
        client: Client,
        settings,
        owner_user,
        club,
    ):
        from apps.billing.refund_services import approve_payment_refund_case
        from apps.clubs.timezones import club_localdate
        from apps.trainers.models import (
            TrainerEarning,
            TrainerEarningAdjustment,
            TrainerPayrollPeriodClose,
        )

        order, reservation, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        order.payment.status = Payment.Status.CONFIRMED
        order.payment.save(update_fields=["status", "updated_at"])
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=reservation.trainer,
            payment=order.payment,
            earning_source=TrainerEarning.Source.SALE,
            earning_type=TrainerEarning.EarningType.PERSONAL,
            amount=Decimal("200.00"),
            rate_percent=Decimal("20.00"),
            subscription_price=order.payment.amount,
        )
        accounting_date = club_localdate(club)
        TrainerPayrollPeriodClose.objects.create(
            club=club,
            period_start=accounting_date,
            period_end=accounting_date,
            closed_by=owner_user,
            reason="Период выплачен",
        )
        refund_case = PaymentRefundCase.objects.create(
            club=club,
            order=order,
            refund_kind=PaymentRefundCase.Kind.PARTIAL,
            status=PaymentRefundCase.Status.RECONCILIATION_REQUIRED,
            provider_refunded_at=timezone.now(),
        )
        refund = approve_payment_refund_case(
            club_id=club.id,
            case_id=refund_case.id,
            actor_user_id=owner_user.id,
            idempotency_key="htmx-closed-payroll-refund",
            amount=Decimal("100.00"),
            refund_kind=PaymentRefund.Kind.PARTIAL,
            reason="Возврат провайдера",
        )
        client.force_login(owner_user)

        page = client.get("/dashboard/billing/payments/?queue=online")
        open_date = accounting_date + timedelta(days=1)
        response = client.post(
            f"/dashboard/billing/payment-refunds/{refund.id}/complete-payroll/",
            {"effective_date": open_date.isoformat()},
            HTTP_HX_REQUEST="true",
        )

        assert refund.status == PaymentRefund.Status.PAYROLL_ACTION_REQUIRED
        assert page.status_code == 200
        assert "нужна дата корректировки зарплаты" in page.content.decode()
        assert response.status_code == 204
        refund.refresh_from_db()
        adjustment = TrainerEarningAdjustment.objects.get(
            source_refund=refund,
            source_earning=earning,
        )
        assert refund.status == PaymentRefund.Status.COMPLETED
        assert adjustment.effective_date == open_date

    def test_trainer_cannot_resolve_bank_payment_order_review(
        self,
        client: Client,
        settings,
        trainer_user,
        owner_user,
        club,
    ):
        order, _, _ = self._create_personal_manual_review_order(
            settings=settings,
            club=club,
            owner_user=owner_user,
        )
        client.force_login(trainer_user)

        response = client.post(
            f"/dashboard/billing/bank-payment-orders/{order.id}/review/",
            {
                "resolution": BankPaymentOrderReviewEvent.Resolution.REJECT,
                "reason": "Не должен иметь доступ",
            },
            HTTP_HX_REQUEST="true",
        )

        order.refresh_from_db()
        assert response.status_code == 403
        assert order.status == BankPaymentOrder.Status.MANUAL_REVIEW

    def test_payment_list_shows_seller_and_package_owner_hint(self, client: Client, owner_user, club):
        from apps.billing.tests.factories import PaymentFactory, SubscriptionFactory, TariffFactory

        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        tariff = TariffFactory(training_type=training_type)
        student = StudentFactory(club=club)
        subscription = SubscriptionFactory(club=club, student=student, tariff=tariff)
        seller = TrainerFactory(club=club, first_name="Seller", last_name="Coach")
        PaymentFactory(
            club=club,
            student=student,
            tariff=tariff,
            subscription=subscription,
            status=Payment.Status.PENDING,
            recorded_by=owner_user,
            seller_trainer=seller,
        )
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Продавец: Seller Coach" in content
        assert "Пакет закрепится за: Seller Coach" in content

    def test_payment_list_shows_permanent_group_confirmation_context(
        self,
        client: Client,
        owner_user,
        club,
    ):
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        target_trainer = TrainerFactory(club=club, first_name="Target", last_name="Coach")
        target_schedule = ScheduleFactory(
            club=club,
            trainer=target_trainer,
            training_type=training_type,
            group_name="Tue Thu Group",
        )
        tariff = TariffFactory(club=club, training_type=training_type)
        student = StudentFactory(club=club)
        start_date = date(2030, 1, 7)
        PaymentFactory(
            club=club,
            student=student,
            tariff=tariff,
            status=Payment.Status.PENDING,
            recorded_by=owner_user,
            seller_trainer=target_trainer,
            target_schedule=target_schedule,
            target_start_date=start_date,
            target_group_name_snapshot="Tue Thu Group",
            target_location_name_snapshot=target_schedule.location.name,
            target_trainer_name_snapshot="Target Coach",
            sale_trainer_id_snapshot=target_trainer.id,
        )
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/payments/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Постоянная группа: Tue Thu Group" in content
        assert target_schedule.location.name in content
        assert "старт 07.01.2030" in content
        assert "Тренер группы / продажа:" in content
        assert "Target Coach" in content
        assert "Подтверждение создаст постоянное зачисление" in content

    def test_payment_reject_action_passes_rejection_reason_to_service(self, client: Client, owner_user, club):
        client.force_login(owner_user)

        with patch("apps.htmx_admin.views.billing.billing_verify_payment") as verify_payment:
            response = client.post(
                "/dashboard/billing/payments/123/verify/",
                {"action": "reject", "rejection_reason": "Receipt mismatch"},
                HTTP_HX_REQUEST="true",
            )

        assert response.status_code == 204
        assert response["HX-Redirect"] == "/dashboard/billing/payments/"
        verify_payment.assert_called_once_with(
            payment_id=123,
            club_id=club.id,
            verified_by_id=owner_user.id,
            action="reject",
            rejection_reason="Receipt mismatch",
        )

    def test_payment_reject_action_requires_reason_without_mutation(self, client: Client, owner_user, club):
        from apps.billing.tests.factories import PaymentFactory

        payment = PaymentFactory(club=club, status=Payment.Status.PENDING, recorded_by=owner_user)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/billing/payments/{payment.id}/verify/",
            {"action": "reject", "rejection_reason": "   "},
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert "Укажите причину отклонения оплаты".encode() in response.content
        assert not response.has_header("HX-Redirect")
        payment.refresh_from_db()
        assert payment.status == Payment.Status.PENDING
        assert payment.rejection_reason == ""

    def test_debtor_tenant_isolation(self, client: Client, owner_user, club, other_club):
        from apps.attendance.tests.factories import CheckinFactory
        from apps.billing.tests.factories import DebtFactory

        checkin_mine = CheckinFactory(club=club)
        checkin_other = CheckinFactory(club=other_club)
        DebtFactory(club=club, student=checkin_mine.student, checkin=checkin_mine)
        DebtFactory(club=other_club, student=checkin_other.student, checkin=checkin_other)
        client.force_login(owner_user)
        response = client.get("/dashboard/billing/")
        content = response.content.decode()
        assert str(checkin_mine.student) in content
        assert str(checkin_other.student) not in content

    def test_debtor_export_preserves_filters_and_returns_xlsx(self, client: Client, owner_user, club):
        from io import BytesIO

        import openpyxl

        from apps.attendance.tests.factories import CheckinFactory
        from apps.billing.tests.factories import DebtFactory

        active_student = StudentFactory(
            club=club,
            first_name="Active",
            last_name="Debtor",
            status="active",
        )
        churned_student = StudentFactory(
            club=club,
            first_name="Churned",
            last_name="Debtor",
            status="churned",
        )
        active_checkin = CheckinFactory(club=club, student=active_student)
        churned_checkin = CheckinFactory(club=club, student=churned_student)
        DebtFactory(
            club=club,
            student=active_student,
            checkin=active_checkin,
            tariff_price=Decimal("7000"),
        )
        DebtFactory(
            club=club,
            student=churned_student,
            checkin=churned_checkin,
            tariff_price=Decimal("9000"),
        )

        client.force_login(owner_user)
        page_response = client.get("/dashboard/billing/?filter=large&student_status=active")
        content = page_response.content.decode()
        assert "/dashboard/billing/debtors/export/?" in content
        assert "filter=large" in content
        assert "student_status=active" in content

        response = client.get("/dashboard/billing/debtors/export/?filter=large&student_status=active")
        assert response.status_code == 200
        assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        wb = openpyxl.load_workbook(BytesIO(response.content))
        rows = list(wb.active.iter_rows(values_only=True))
        exported_students = {row[0] for row in rows[1:]}
        assert exported_students == {"Active Debtor"}

    def test_debtor_write_off_success_refreshes_list(self, client: Client, owner_user, club):
        from apps.attendance.tests.factories import CheckinFactory
        from apps.billing.tests.factories import DebtFactory

        student = StudentFactory(club=club, first_name="Writeoff", last_name="Student")
        checkin = CheckinFactory(club=club, student=student)
        debt = DebtFactory(
            club=club,
            student=student,
            checkin=checkin,
            tariff_price=Decimal("1500"),
        )

        client.force_login(owner_user)
        page_response = client.get("/dashboard/billing/")
        content = page_response.content.decode()
        assert f"/dashboard/billing/debts/{debt.id}/write-off/" in content
        assert f"/dashboard/students/{student.id}/card/" in content
        assert f'hx-get="/dashboard/students/{student.id}/"' not in content
        assert 'name="reason"' in content

        response = client.post(
            f"/dashboard/billing/debts/{debt.id}/write-off/",
            {"reason": "Admin correction"},
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        debt.refresh_from_db()
        assert debt.resolution_type == "writeoff"
        assert debt.resolved_at is not None
        event = DebtWriteOffEvent.objects.for_club(club.id).get(debt=debt)
        assert event.written_off_by_id == owner_user.id
        assert event.reason == "Admin correction"
        content = response.content.decode()
        assert "Нет должников" in content
        assert "Writeoff Student" not in content

    def test_debtor_write_off_reserved_debt_shows_safe_error(self, client: Client, owner_user, club):
        from apps.attendance.tests.factories import CheckinFactory
        from apps.billing.tests.factories import DebtFactory, PaymentFactory, TariffFactory

        tt = TrainingTypeFactory(club=club)
        tariff = TariffFactory(training_type=tt)
        student = StudentFactory(club=club, first_name="Reserved", last_name="Student")
        checkin = CheckinFactory(
            club=club,
            student=student,
            training_type=tt,
            subscription=None,
            is_debt=True,
        )
        debt = DebtFactory(
            club=club,
            student=student,
            checkin=checkin,
            tariff_price=Decimal("2500"),
        )
        payment = PaymentFactory(
            tariff=tariff,
            student=student,
            recorded_by=owner_user,
            status=Payment.Status.PENDING,
        )
        debt.settlement_payment = payment
        debt.save(update_fields=["settlement_payment"])

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/billing/debts/{debt.id}/write-off/",
            {"reason": "Bypass attempt"},
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        debt.refresh_from_db()
        assert debt.resolved_at is None
        assert DebtWriteOffEvent.objects.for_club(club.id).filter(debt=debt).count() == 0
        content = response.content.decode()
        assert "ожидающей оплате" in content
        assert "Reserved Student" not in content

    def test_unauthenticated_redirects(self, client: Client):
        response = client.get("/dashboard/billing/")
        assert response.status_code == 302


class TestSubscriptionManagement:
    """ADMIN-07: Subscription create, freeze, discounts."""

    def test_subscription_list_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/billing/subscriptions/")
        assert response.status_code == 200
        assert "Абонементы".encode() in response.content

    def test_subscription_list_shows_data(self, client: Client, owner_user, club):
        from apps.billing.tests.factories import SubscriptionFactory

        sub = SubscriptionFactory(club=club)
        client.force_login(owner_user)
        response = client.get("/dashboard/billing/subscriptions/")
        assert response.status_code == 200
        assert sub.tariff.name.encode() in response.content

    def test_subscription_list_filters_and_labels_cancelled_refund_state(
        self,
        client: Client,
        owner_user,
        club,
    ):
        cancelled = SubscriptionFactory(
            club=club,
            status=Subscription.Status.CANCELLED,
            trainings_left=5,
        )
        SubscriptionFactory(club=club, status=Subscription.Status.ACTIVE)
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/subscriptions/?status=cancelled")

        assert response.status_code == 200
        content = response.content.decode()
        assert cancelled.tariff.name in content
        assert "Отменён" in content
        assert "остаток недоступен" in content
        assert "Заморозить" not in content

    def test_subscription_create_form_includes_package_owner_field(self, client: Client, owner_user, club):
        TrainerFactory(club=club, first_name="Owner", last_name="Coach")
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/subscriptions/")

        assert response.status_code == 200
        content = response.content.decode()
        assert 'name="package_owner_trainer_id"' in content
        expected_required_binding = (
            'x-bind:required="'
            "(packageOwnerKind === 'personal' || packageOwnerKind === 'mini_group') "
            '&& !sellerTrainerId"'
        )
        assert expected_required_binding in content
        assert "Владелец пакета" in content
        assert "Owner Coach" in content
        assert 'name="payment_method"' in content
        assert 'value="cash"' in content
        assert 'value="transfer"' in content
        assert 'value="online"' in content
        assert "СБП — создать ссылку" in content
        assert "СОЗДАТЬ ССЫЛКУ СБП" in content

    def test_create_subscription_view_passes_package_owner_to_service(self, client: Client, owner_user, club):
        from apps.billing.tests.factories import TariffFactory

        tariff = TariffFactory(training_type=TrainingTypeFactory(club=club))
        student = StudentFactory(club=club)
        seller = TrainerFactory(club=club)
        package_owner = TrainerFactory(club=club)
        client.force_login(owner_user)

        with patch("apps.htmx_admin.views.billing.billing_create_subscription") as create_subscription:
            response = client.post(
                "/dashboard/billing/subscriptions/create/",
                {
                    "student_id": str(student.id),
                    "tariff_id": str(tariff.id),
                    "seller_trainer_id": str(seller.id),
                    "package_owner_trainer_id": str(package_owner.id),
                    "payment_method": Payment.Method.TRANSFER,
                },
                HTTP_HX_REQUEST="true",
            )

        assert response.status_code == 204
        create_subscription.assert_called_once_with(
            club_id=club.id,
            student_id=student.id,
            tariff_id=tariff.id,
            seller_trainer_id=seller.id,
            package_owner_trainer_id=package_owner.id,
            recorded_by_id=owner_user.id,
            payment_method=Payment.Method.TRANSFER,
        )

    def test_create_subscription_view_creates_owner_bank_payment_order(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.billing.tests.factories import TariffFactory

        tariff = TariffFactory(training_type=TrainingTypeFactory(club=club))
        student = StudentFactory(club=club)
        seller = TrainerFactory(club=club)
        package_owner = TrainerFactory(club=club)
        client.force_login(owner_user)

        with (
            patch(
                "apps.htmx_admin.views.billing.billing_create_bank_payment_order",
                return_value=SimpleNamespace(id=77),
            ) as create_bank_order,
            patch("apps.htmx_admin.views.billing.billing_create_subscription") as create_subscription,
        ):
            response = client.post(
                "/dashboard/billing/subscriptions/create/",
                {
                    "student_id": str(student.id),
                    "tariff_id": str(tariff.id),
                    "seller_trainer_id": str(seller.id),
                    "package_owner_trainer_id": str(package_owner.id),
                    "payment_method": Payment.Method.ONLINE,
                },
                HTTP_HX_REQUEST="true",
            )

        assert response.status_code == 204
        assert response.headers["HX-Redirect"] == "/dashboard/billing/payments/?bank_order=77"
        create_bank_order.assert_called_once_with(
            club_id=club.id,
            student_id=student.id,
            tariff_id=tariff.id,
            source=BankPaymentOrder.Source.OWNER,
            created_by_id=owner_user.id,
            seller_trainer_id=seller.id,
            package_owner_trainer_id=package_owner.id,
        )
        create_subscription.assert_not_called()

    @pytest.mark.parametrize("payment_method", ["", "crypto"])
    def test_create_subscription_view_rejects_missing_or_unknown_payment_method(
        self,
        payment_method,
        client: Client,
        owner_user,
        club,
    ):
        from apps.billing.tests.factories import TariffFactory

        tariff = TariffFactory(training_type=TrainingTypeFactory(club=club))
        student = StudentFactory(club=club)
        client.force_login(owner_user)

        with patch("apps.htmx_admin.views.billing.billing_create_subscription") as create_subscription:
            response = client.post(
                "/dashboard/billing/subscriptions/create/",
                {
                    "student_id": str(student.id),
                    "tariff_id": str(tariff.id),
                    "payment_method": payment_method,
                },
                HTTP_HX_REQUEST="true",
            )

        assert response.status_code == 400
        create_subscription.assert_not_called()

    def test_subscription_list_shows_package_owner(self, client: Client, owner_user, club):
        from apps.billing.tests.factories import SubscriptionFactory
        from apps.trainers.models import TrainerPackageAllocation

        package_owner = TrainerFactory(club=club, first_name="Owner", last_name="Coach")
        sub = SubscriptionFactory(club=club)
        TrainerPackageAllocation.objects.create(
            club=club,
            subscription=sub,
            student=sub.student,
            tariff=sub.tariff,
            training_type=sub.tariff.training_type,
            owner_trainer=package_owner,
            sessions_total_snapshot=sub.tariff.trainings_limit,
            sessions_remaining_snapshot=sub.trainings_left,
            amount_snapshot=sub.paid_amount or sub.tariff.price,
        )
        client.force_login(owner_user)

        response = client.get("/dashboard/billing/subscriptions/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Владелец пакета" in content
        assert "Owner Coach" in content

    def test_subscription_list_htmx_partial(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/billing/subscriptions/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert b"<aside" not in response.content

    def test_freeze_form_renders(self, client: Client, owner_user, club):
        from apps.billing.tests.factories import SubscriptionFactory

        sub = SubscriptionFactory(club=club, status="active")
        client.force_login(owner_user)
        response = client.get(f"/dashboard/billing/subscriptions/{sub.id}/freeze/")
        assert response.status_code == 200
        assert "Заморозка абонемента".encode() in response.content

    def test_freeze_form_shows_pending_request_state(self, client: Client, owner_user, trainer_user, club):
        from apps.billing.tests.factories import SubscriptionFactory, SubscriptionFreezeFactory

        sub = SubscriptionFactory(club=club, status=Subscription.Status.ACTIVE)
        SubscriptionFreezeFactory(
            subscription=sub,
            frozen_by=trainer_user,
            status=SubscriptionFreeze.FreezeStatus.PENDING,
        )
        client.force_login(owner_user)

        response = client.get(f"/dashboard/billing/subscriptions/{sub.id}/freeze/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "уже есть заявка на заморозку" in content
        assert f'hx-post="/dashboard/billing/subscriptions/{sub.id}/freeze/"' not in content

    def test_freeze_form_404_other_club(self, client: Client, owner_user, club, other_club):
        from apps.billing.tests.factories import SubscriptionFactory

        sub = SubscriptionFactory(club=other_club)
        client.force_login(owner_user)
        response = client.get(f"/dashboard/billing/subscriptions/{sub.id}/freeze/")
        assert response.status_code == 404

    def test_subscription_list_shows_pending_freeze_requests_only(
        self,
        client: Client,
        owner_user,
        trainer_user,
        club,
        other_club,
    ):
        from apps.billing.tests.factories import (
            SubscriptionFactory,
            SubscriptionFreezeFactory,
            TariffFactory,
        )

        tariff = TariffFactory(training_type=TrainingTypeFactory(club=club), name="Kids 8")
        student = StudentFactory(club=club, first_name="FreezePending", last_name="Kid")
        sub = SubscriptionFactory(tariff=tariff, student=student, status=Subscription.Status.ACTIVE)
        pending_freeze = SubscriptionFreezeFactory(
            subscription=sub,
            frozen_by=trainer_user,
            status=SubscriptionFreeze.FreezeStatus.PENDING,
            days=5,
            reason=SubscriptionFreeze.Reason.INJURY,
        )
        approved_student = StudentFactory(club=club, first_name="ApprovedOnly")
        approved_sub = SubscriptionFactory(tariff=tariff, student=approved_student, status=Subscription.Status.FROZEN)
        approved_freeze = SubscriptionFreezeFactory(
            subscription=approved_sub,
            frozen_by=trainer_user,
            status=SubscriptionFreeze.FreezeStatus.APPROVED,
        )
        foreign_sub = SubscriptionFactory(club=other_club, status=Subscription.Status.ACTIVE)
        foreign_freeze = SubscriptionFreezeFactory(
            subscription=foreign_sub,
            frozen_by=trainer_user,
            status=SubscriptionFreeze.FreezeStatus.PENDING,
        )

        client.force_login(owner_user)
        response = client.get("/dashboard/billing/subscriptions/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Заявки на заморозку" in content
        assert "FreezePending" in content
        assert "Kids 8" in content
        assert "5 дн." in content
        assert trainer_user.username in content
        assert f"/dashboard/billing/freezes/{pending_freeze.id}/approve/" in content
        assert f"/dashboard/billing/freezes/{pending_freeze.id}/reject/" in content
        assert f"/dashboard/billing/freezes/{approved_freeze.id}/approve/" not in content
        assert f"/dashboard/billing/freezes/{foreign_freeze.id}/approve/" not in content

    def test_freeze_approval_action_approves_and_redirects_to_refresh_page(
        self,
        client: Client,
        owner_user,
        trainer_user,
        club,
    ):
        from apps.billing.tests.factories import SubscriptionFactory, SubscriptionFreezeFactory

        sub = SubscriptionFactory(club=club, status=Subscription.Status.ACTIVE)
        freeze = SubscriptionFreezeFactory(
            subscription=sub,
            frozen_by=trainer_user,
            status=SubscriptionFreeze.FreezeStatus.PENDING,
        )

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/billing/freezes/{freeze.id}/approve/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 204
        assert response.headers["HX-Redirect"] == "/dashboard/billing/subscriptions/"
        freeze.refresh_from_db()
        sub.refresh_from_db()
        assert freeze.status == SubscriptionFreeze.FreezeStatus.APPROVED
        assert freeze.approved_by_id == owner_user.id
        assert freeze.decision_at is not None
        assert sub.status == Subscription.Status.FROZEN

    def test_freeze_approval_action_rejects_with_reason_and_redirects_to_refresh_page(
        self,
        client: Client,
        owner_user,
        trainer_user,
        club,
    ):
        from apps.billing.tests.factories import SubscriptionFactory, SubscriptionFreezeFactory

        sub = SubscriptionFactory(club=club, status=Subscription.Status.ACTIVE)
        freeze = SubscriptionFreezeFactory(
            subscription=sub,
            frozen_by=trainer_user,
            status=SubscriptionFreeze.FreezeStatus.PENDING,
        )

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/billing/freezes/{freeze.id}/reject/",
            {"decision_reason": "Нужен документ"},
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 204
        assert response.headers["HX-Redirect"] == "/dashboard/billing/subscriptions/"
        freeze.refresh_from_db()
        sub.refresh_from_db()
        assert freeze.status == SubscriptionFreeze.FreezeStatus.REJECTED
        assert freeze.rejected_by_id == owner_user.id
        assert freeze.decision_at is not None
        assert freeze.decision_reason == "Нужен документ"
        assert sub.status == Subscription.Status.ACTIVE

    def test_freeze_approval_action_404_for_other_club(
        self,
        client: Client,
        owner_user,
        trainer_user,
        club,
        other_club,
    ):
        from apps.billing.tests.factories import SubscriptionFactory, SubscriptionFreezeFactory

        freeze = SubscriptionFreezeFactory(
            subscription=SubscriptionFactory(club=other_club, status=Subscription.Status.ACTIVE),
            frozen_by=trainer_user,
            status=SubscriptionFreeze.FreezeStatus.PENDING,
        )

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/billing/freezes/{freeze.id}/approve/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 404
        freeze.refresh_from_db()
        assert freeze.status == SubscriptionFreeze.FreezeStatus.PENDING

    def test_freeze_approval_action_denies_non_management_role(
        self,
        client: Client,
        trainer_user,
        club,
    ):
        from apps.billing.tests.factories import SubscriptionFactory, SubscriptionFreezeFactory

        freeze = SubscriptionFreezeFactory(
            subscription=SubscriptionFactory(club=club, status=Subscription.Status.ACTIVE),
            frozen_by=trainer_user,
            status=SubscriptionFreeze.FreezeStatus.PENDING,
        )

        client.force_login(trainer_user)
        response = client.post(
            f"/dashboard/billing/freezes/{freeze.id}/approve/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 403
        freeze.refresh_from_db()
        assert freeze.status == SubscriptionFreeze.FreezeStatus.PENDING

    def test_unauthenticated_redirects(self, client: Client):
        response = client.get("/dashboard/billing/subscriptions/")
        assert response.status_code == 302


class TestTrainerViews:
    """ADMIN-08: Trainer list, salary view."""

    def test_trainer_list_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/trainers/")
        assert response.status_code == 200
        assert "ТРЕНЕРЫ".encode() in response.content

    def test_trainer_list_htmx_partial(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/trainers/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert b"<aside" not in response.content

    @patch("apps.htmx_admin.views.trainers.get_trainers_with_stats")
    def test_trainer_list_default_range_uses_club_local_today(
        self,
        mock_stats,
        client: Client,
        owner_user,
        club,
        monkeypatch,
    ):
        club.timezone = "Asia/Yekaterinburg"
        club.save(update_fields=["timezone"])
        mock_stats.return_value = []
        monkeypatch.setattr(
            "apps.clubs.timezones.timezone.now",
            lambda: datetime(2026, 6, 28, 20, 30, tzinfo=UTC),
        )

        client.force_login(owner_user)
        response = client.get("/dashboard/trainers/")

        assert response.status_code == 200
        _, kwargs = mock_stats.call_args
        assert kwargs["date_from"] == date(2026, 6, 1)
        assert kwargs["date_to"] == date(2026, 6, 29)

    def test_trainer_list_shows_trainer(self, client: Client, owner_user, club):
        from apps.trainers.tests.factories import TrainerFactory

        TrainerFactory(club=club, first_name="CoachIvan", last_name="Petrov")
        client.force_login(owner_user)
        response = client.get("/dashboard/trainers/")
        content = response.content.decode()
        assert "CoachIvan" in content
        assert "Petrov" in content

    def test_trainer_list_tenant_isolation(self, client: Client, owner_user, club, other_club):
        from apps.trainers.tests.factories import TrainerFactory

        TrainerFactory(club=club, first_name="MyTrainer")
        TrainerFactory(club=other_club, first_name="OtherTrainer")
        client.force_login(owner_user)
        response = client.get("/dashboard/trainers/")
        content = response.content.decode()
        assert "MyTrainer" in content
        assert "OtherTrainer" not in content

    def test_trainer_session_filter_uses_canonical_name_for_mapped_schedule(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.attendance.tests.factories import TrainingGroupFactory

        trainer = TrainerFactory(club=club)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        location = LocationFactory(club=club)
        group = TrainingGroupFactory(
            club=club,
            name="Canonical session group",
            training_type=training_type,
            location=location,
            responsible_trainer=trainer,
        )
        schedule = ScheduleFactory(
            club=club,
            training_group=group,
            trainer=trainer,
            training_type=training_type,
            location=location,
            start_time=time(10, 0),
        )
        Schedule.objects.for_club(club).filter(id=schedule.id).update(
            group_name="Legacy session snapshot"
        )
        schedule.refresh_from_db()
        client.force_login(owner_user)

        response = client.get("/dashboard/trainers/sessions/")

        assert response.status_code == 200
        content = response.content.decode()
        assert f">Canonical session group · {schedule.start_time:%H:%M}</option>" in content
        assert "Legacy session snapshot ·" not in content

    def test_trainer_create_success_keeps_list_tab_active(self, client: Client, owner_user, club):
        from apps.trainers.models import Trainer

        client.force_login(owner_user)
        response = client.post(
            "/dashboard/trainers/create/",
            {
                "first_name": "New",
                "last_name": "Coach",
                "phone": "",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert response["HX-Trigger"] == "closeSlideOver"
        content = response.content.decode()
        assert "New" in content
        assert "Coach" in content
        assert "СПИСОК ТРЕНЕРОВ" in content
        assert "background-color: var(--branding-dark); color: var(--branding-on-primary);" in content
        assert Trainer.objects.for_club(club).filter(first_name="New", last_name="Coach").exists()

    def test_trainer_create_rejects_out_of_range_rate_without_mutation(self, client: Client, owner_user, club):
        from apps.trainers.models import Trainer, TrainerRate

        location = LocationFactory(club=club)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP, is_active=True)

        client.force_login(owner_user)
        response = client.post(
            "/dashboard/trainers/create/",
            {
                "first_name": "Bad",
                "last_name": "Rate",
                f"location_{location.id}": "1",
                f"percent_{location.id}_{training_type.id}": "150",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert response["HX-Retarget"] == "#slide-over"
        content = response.content.decode()
        assert "Ставка должна быть 0..100%" in content
        assert not Trainer.objects.for_club(club).filter(first_name="Bad", last_name="Rate").exists()
        assert not TrainerRate.objects.for_club(club).filter(
            location=location,
            training_type=training_type,
        ).exists()

    def test_trainer_detail_missing_rate_warning_copy_matches_salary_contract(self, client: Client, owner_user, club):
        from apps.trainers.models import TrainerLocation

        trainer = TrainerFactory(club=club, first_name="Copy", last_name="Coach")
        location = LocationFactory(club=club, name="Copy Hall")
        TrainingTypeFactory(club=club, name="Group Copy", kind=TrainingType.Kind.GROUP, is_active=True)
        TrainingTypeFactory(club=club, name="Personal Copy", kind=TrainingType.Kind.PERSONAL, is_active=True)
        TrainerLocation.objects.create(club=club, trainer=trainer, location=location)

        client.force_login(owner_user)
        response = client.get(f"/dashboard/trainers/{trainer.id}/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Не заполнены ставки (2)" in content
        assert "Для этих комбинаций локация × тип тренировки автоматическое начисление зарплаты" in content
        assert "Персональные и мини-групповые чек-ины остановятся с ошибкой" in content
        assert "trainer_rate_not_set" in content
        assert "чек-ины упадут" not in content

    def test_trainer_detail_missing_rate_warning_respects_group_sale_location_fallback(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.trainers.models import TrainerLocation, TrainerRate

        trainer = TrainerFactory(club=club, first_name="Fallback", last_name="Coach")
        source_location = LocationFactory(club=club, name="Source Hall")
        fallback_location = LocationFactory(club=club, name="Fallback Hall")
        group_type = TrainingTypeFactory(
            club=club,
            name="Group Fallback",
            kind=TrainingType.Kind.GROUP,
            is_active=True,
        )
        personal_type = TrainingTypeFactory(
            club=club,
            name="Personal Exact",
            kind=TrainingType.Kind.PERSONAL,
            is_active=True,
        )
        TrainerLocation.objects.create(club=club, trainer=trainer, location=source_location)
        TrainerLocation.objects.create(club=club, trainer=trainer, location=fallback_location)
        TrainerRate.objects.create(
            club=club,
            trainer=trainer,
            location=source_location,
            training_type=group_type,
            percent=Decimal("20"),
        )
        TrainerRate.objects.create(
            club=club,
            trainer=trainer,
            location=source_location,
            training_type=personal_type,
            percent=Decimal("30"),
        )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/trainers/{trainer.id}/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Не заполнены ставки (1)" in content
        assert "Fallback Hall — Personal Exact" in content
        assert "Fallback Hall — Group Fallback" not in content

    def test_trainer_detail_counts_only_students_enrolled_in_trainer_schedules(
        self,
        client: Client,
        owner_user,
        club,
    ):
        trainer = TrainerFactory(club=club, first_name="Scoped", last_name="Coach")
        other_trainer = TrainerFactory(club=club, first_name="Other", last_name="Coach")
        schedule = ScheduleFactory(club=club, trainer=trainer)
        duplicate_schedule = ScheduleFactory(club=club, trainer=trainer, group_name="Second Group")
        other_schedule = ScheduleFactory(club=club, trainer=other_trainer)
        scoped_student = StudentFactory(club=club, status="active")
        ScheduleEnrollment.objects.create(club=club, student=scoped_student, schedule=schedule)
        ScheduleEnrollment.objects.create(club=club, student=scoped_student, schedule=duplicate_schedule)
        ScheduleEnrollment.objects.create(
            club=club,
            student=StudentFactory(club=club, status="active"),
            schedule=other_schedule,
        )
        StudentFactory(club=club, status="active")

        client.force_login(owner_user)
        response = client.get(f"/dashboard/trainers/{trainer.id}/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Учеников" in content
        assert '<span class="text-[13px] font-bold" style="color: var(--branding-text);">1</span>' in content

    @patch("apps.htmx_admin.views.trainers.get_trainer_earnings_summary")
    def test_trainer_salary_renders(self, mock_summary, client: Client, owner_user, club, monkeypatch):
        from apps.trainers.tests.factories import TrainerFactory

        club.timezone = "Asia/Yekaterinburg"
        club.save(update_fields=["timezone"])
        monkeypatch.setattr(
            "apps.clubs.timezones.timezone.now",
            lambda: datetime(2026, 6, 28, 20, 30, tzinfo=UTC),
        )
        trainer = TrainerFactory(club=club, first_name="SalaryTest")
        mock_summary.return_value = {
            "total_amount": Decimal("25000"),
            "total_sessions": 20,
            "by_type": {
                "group": {"count": 15, "total": Decimal("18000")},
                "personal": {"count": 5, "total": Decimal("7000")},
            },
        }
        client.force_login(owner_user)
        response = client.get(f"/dashboard/trainers/{trainer.id}/salary/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "SalaryTest" in content
        assert "25000" in content
        _, kwargs = mock_summary.call_args
        assert kwargs["date_from"] == date(2026, 6, 1)
        assert kwargs["date_to"] == date(2026, 6, 29)

    def test_trainer_salary_shows_package_owner_audit_context(self, client: Client, owner_user, club):
        from apps.attendance.tests.factories import CheckinFactory, ScheduleFactory
        from apps.trainers.models import TrainerEarning, TrainerEarningAdjustment
        from apps.trainers.tests.factories import TrainerFactory

        target_date = date.today().replace(day=1)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        owner_trainer = TrainerFactory(club=club, first_name="Owner", last_name="Coach")
        actual_trainer = TrainerFactory(club=club, first_name="Actual", last_name="Coach")
        schedule = ScheduleFactory(
            club=club,
            trainer=actual_trainer,
            training_type=training_type,
            group_name="Personal Slot",
        )
        checkin = CheckinFactory(
            club=club,
            schedule=schedule,
            trainer=actual_trainer,
            training_type=training_type,
            date=target_date,
        )
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=actual_trainer,
            checkin=checkin,
            earning_type=TrainingType.Kind.PERSONAL,
            amount=Decimal("3000.00"),
            rate_percent=Decimal("50.00"),
            subscription_price=Decimal("6000.00"),
        )
        TrainerEarningAdjustment.objects.create(
            club=club,
            trainer=actual_trainer,
            amount_basis_snapshot=earning.amount,
            payable_amount_delta=Decimal("0.00"),
            affects_payroll=False,
            direction=TrainerEarningAdjustment.Direction.INFO,
            kind=TrainerEarningAdjustment.Kind.PACKAGE_TRANSFER,
            effective_date=target_date,
            source_checkin=checkin,
            counterparty_trainer=owner_trainer,
            reason="package_owner_differs_from_actual_trainer",
            created_by=owner_user,
        )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/trainers/{actual_trainer.id}/salary/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "АУДИТ И КОРРЕКТИРОВКИ" in content
        assert "Personal Slot" in content
        assert "Пакет тренера: Owner Coach" in content
        assert "package_owner_differs_from_actual_trainer" in content
        assert "Передача пакета" in content
        assert "Информация" in content
        assert "Package transfer" not in content
        assert "не влияет на выплату" in content

    def test_trainer_earning_correction_form_renders(self, client: Client, owner_user, club):
        from apps.attendance.tests.factories import CheckinFactory, ScheduleFactory
        from apps.trainers.models import TrainerEarning
        from apps.trainers.tests.factories import TrainerFactory

        target_date = date.today().replace(day=1)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        source_trainer = TrainerFactory(club=club, first_name="Actual", last_name="Coach")
        TrainerFactory(club=club, first_name="Owner", last_name="Coach")
        schedule = ScheduleFactory(
            club=club,
            trainer=source_trainer,
            training_type=training_type,
            group_name="Personal Slot",
        )
        checkin = CheckinFactory(
            club=club,
            schedule=schedule,
            trainer=source_trainer,
            training_type=training_type,
            date=target_date,
        )
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=source_trainer,
            checkin=checkin,
            earning_type=TrainingType.Kind.PERSONAL,
            amount=Decimal("3000.00"),
            rate_percent=Decimal("50.00"),
            subscription_price=Decimal("6000.00"),
        )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/trainers/earnings/{earning.id}/correction/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "КОРРЕКТИРОВКА ВЫПЛАТЫ" in content
        assert "Actual Coach" in content
        assert "Owner Coach" in content
        assert "Personal Slot" in content
        assert 'name="idempotency_key"' in content

    def test_trainer_earning_correction_post_creates_manual_adjustments(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.attendance.tests.factories import CheckinFactory, ScheduleFactory
        from apps.trainers.models import TrainerEarning, TrainerEarningAdjustment
        from apps.trainers.tests.factories import TrainerFactory

        target_date = date.today().replace(day=1)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        source_trainer = TrainerFactory(club=club, first_name="Actual", last_name="Coach")
        target_trainer = TrainerFactory(club=club, first_name="Owner", last_name="Coach")
        schedule = ScheduleFactory(club=club, trainer=source_trainer, training_type=training_type)
        checkin = CheckinFactory(
            club=club,
            schedule=schedule,
            trainer=source_trainer,
            training_type=training_type,
            date=target_date,
        )
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=source_trainer,
            checkin=checkin,
            earning_type=TrainingType.Kind.PERSONAL,
            amount=Decimal("3000.00"),
            rate_percent=Decimal("50.00"),
            subscription_price=Decimal("6000.00"),
        )

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/trainers/earnings/{earning.id}/correction/",
            {
                "target_trainer_id": str(target_trainer.id),
                "reason": "Package owner should receive this payout",
                "idempotency_key": "htmx-correction-1",
            },
        )

        assert response.status_code == 200
        assert response.headers["HX-Trigger"] == "trainerUpdated"
        content = response.content.decode()
        assert "Корректировка сохранена" in content
        rows = list(
            TrainerEarningAdjustment.objects.for_club(club)
            .filter(kind=TrainerEarningAdjustment.Kind.MANUAL_ADJUSTMENT)
            .order_by("direction")
        )
        assert len(rows) == 2
        assert {row.idempotency_key for row in rows} == {"htmx-correction-1"}
        assert {row.correction_group_id for row in rows if row.correction_group_id} == {
            rows[0].correction_group_id
        }
        debit = next(row for row in rows if row.direction == TrainerEarningAdjustment.Direction.DEBIT)
        credit = next(row for row in rows if row.direction == TrainerEarningAdjustment.Direction.CREDIT)
        assert debit.trainer == source_trainer
        assert debit.payable_amount_delta == Decimal("-3000.00")
        assert credit.trainer == target_trainer
        assert credit.payable_amount_delta == Decimal("3000.00")

    def test_trainer_salary_shows_manual_credit_without_target_earning(self, client: Client, owner_user, club):
        from apps.attendance.tests.factories import CheckinFactory, ScheduleFactory
        from apps.trainers.models import TrainerEarning, TrainerEarningAdjustment
        from apps.trainers.tests.factories import TrainerFactory

        target_date = date.today().replace(day=1)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        source_trainer = TrainerFactory(club=club, first_name="Actual", last_name="Coach")
        target_trainer = TrainerFactory(club=club, first_name="Owner", last_name="Coach")
        schedule = ScheduleFactory(
            club=club,
            trainer=source_trainer,
            training_type=training_type,
            group_name="Manual Credit Source",
        )
        checkin = CheckinFactory(
            club=club,
            schedule=schedule,
            trainer=source_trainer,
            training_type=training_type,
            date=target_date,
        )
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=source_trainer,
            checkin=checkin,
            earning_type=TrainingType.Kind.PERSONAL,
            amount=Decimal("3000.00"),
            rate_percent=Decimal("50.00"),
            subscription_price=Decimal("6000.00"),
        )
        group_id = "22222222-2222-4222-8222-222222222222"
        TrainerEarningAdjustment.objects.create(
            club=club,
            trainer=source_trainer,
            amount_basis_snapshot=earning.amount,
            payable_amount_delta=Decimal("-3000.00"),
            affects_payroll=True,
            direction=TrainerEarningAdjustment.Direction.DEBIT,
            kind=TrainerEarningAdjustment.Kind.MANUAL_ADJUSTMENT,
            effective_date=target_date,
            source_checkin=checkin,
            counterparty_trainer=target_trainer,
            reason="Owner receives payout",
            created_by=owner_user,
            correction_group_id=group_id,
            idempotency_key="manual-credit-audit",
        )
        TrainerEarningAdjustment.objects.create(
            club=club,
            trainer=target_trainer,
            amount_basis_snapshot=earning.amount,
            payable_amount_delta=Decimal("3000.00"),
            affects_payroll=True,
            direction=TrainerEarningAdjustment.Direction.CREDIT,
            kind=TrainerEarningAdjustment.Kind.MANUAL_ADJUSTMENT,
            effective_date=target_date,
            source_checkin=checkin,
            counterparty_trainer=source_trainer,
            reason="Owner receives payout",
            created_by=owner_user,
            correction_group_id=group_id,
            idempotency_key="manual-credit-audit",
        )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/trainers/{target_trainer.id}/salary/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Owner receives payout" in content
        assert "Actual Coach" in content
        assert "Manual Credit Source" in content
        assert "Ручная корректировка" in content
        assert "Начисление" in content
        assert "Manual adjustment" not in content
        assert "Manual_adjustment" not in content
        assert "+3000" in content or "3000₽" in content

    def test_trainer_salary_hides_correction_button_after_manual_debit(self, client: Client, owner_user, club):
        from apps.attendance.tests.factories import CheckinFactory, ScheduleFactory
        from apps.trainers.models import TrainerEarning, TrainerEarningAdjustment
        from apps.trainers.tests.factories import TrainerFactory

        target_date = date.today().replace(day=1)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        source_trainer = TrainerFactory(club=club, first_name="Actual", last_name="Coach")
        target_trainer = TrainerFactory(club=club, first_name="Owner", last_name="Coach")
        schedule = ScheduleFactory(
            club=club,
            trainer=source_trainer,
            training_type=training_type,
            group_name="Already Corrected",
        )
        checkin = CheckinFactory(
            club=club,
            schedule=schedule,
            trainer=source_trainer,
            training_type=training_type,
            date=target_date,
        )
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=source_trainer,
            checkin=checkin,
            earning_type=TrainingType.Kind.PERSONAL,
            amount=Decimal("3000.00"),
            rate_percent=Decimal("50.00"),
            subscription_price=Decimal("6000.00"),
        )
        TrainerEarningAdjustment.objects.create(
            club=club,
            trainer=source_trainer,
            amount_basis_snapshot=earning.amount,
            payable_amount_delta=Decimal("-3000.00"),
            affects_payroll=True,
            direction=TrainerEarningAdjustment.Direction.DEBIT,
            kind=TrainerEarningAdjustment.Kind.MANUAL_ADJUSTMENT,
            effective_date=target_date,
            source_checkin=checkin,
            counterparty_trainer=target_trainer,
            reason="Already corrected",
            created_by=owner_user,
            correction_group_id="44444444-4444-4444-8444-444444444444",
            idempotency_key="already-corrected",
        )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/trainers/{source_trainer.id}/salary/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "Already corrected" in content
        assert "Списание" in content
        assert f"/dashboard/trainers/earnings/{earning.id}/correction/" not in content

    def test_trainer_salary_shows_closed_period_and_hides_correction_button(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.attendance.tests.factories import CheckinFactory, ScheduleFactory
        from apps.trainers.models import TrainerEarning
        from apps.trainers.services import close_trainer_payroll_period
        from apps.trainers.tests.factories import TrainerFactory

        target_date = date.today().replace(day=1)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        source_trainer = TrainerFactory(club=club, first_name="Closed", last_name="Coach")
        schedule = ScheduleFactory(club=club, trainer=source_trainer, training_type=training_type)
        checkin = CheckinFactory(
            club=club,
            schedule=schedule,
            trainer=source_trainer,
            training_type=training_type,
            date=target_date,
        )
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=source_trainer,
            checkin=checkin,
            earning_type=TrainingType.Kind.PERSONAL,
            amount=Decimal("3000.00"),
            rate_percent=Decimal("50.00"),
            subscription_price=Decimal("6000.00"),
        )
        close_trainer_payroll_period(
            club_id=club.id,
            period_start=target_date,
            period_end=target_date,
            reason="Payroll approved",
            actor_user_id=owner_user.id,
        )

        client.force_login(owner_user)
        response = client.get(
            f"/dashboard/trainers/{source_trainer.id}/salary/"
            f"?date_from={target_date.isoformat()}&date_to={target_date.isoformat()}"
        )

        assert response.status_code == 200
        content = response.content.decode()
        assert "Период выплат закрыт" in content
        assert "Закрыть период выплат по клубу" not in content
        assert "Payroll approved" in content
        assert f"/dashboard/trainers/earnings/{earning.id}/correction/" not in content

    def test_trainer_salary_hides_payment_correction_on_club_local_closed_day(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.trainers.models import TrainerEarning
        from apps.trainers.services import close_trainer_payroll_period

        club.timezone = "Asia/Yekaterinburg"
        club.save(update_fields=["timezone"])
        closed_local_date = date(2026, 7, 16)
        verified_at = datetime(2026, 7, 15, 20, 30, tzinfo=UTC)
        source_trainer = TrainerFactory(club=club, first_name="Boundary", last_name="Coach")
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        tariff = TariffFactory(club=club, training_type=training_type)
        subscription = SubscriptionFactory(club=club, tariff=tariff)
        close_trainer_payroll_period(
            club_id=club.id,
            period_start=closed_local_date,
            period_end=closed_local_date,
            reason="Payroll approved",
            actor_user_id=owner_user.id,
        )
        payment = PaymentFactory(
            club=club,
            student=subscription.student,
            tariff=tariff,
            subscription=subscription,
            status=Payment.Status.CONFIRMED,
            verified_at=verified_at,
        )
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=source_trainer,
            payment=payment,
            earning_source=TrainerEarning.Source.SALE,
            earning_type=TrainingType.Kind.GROUP,
            amount=Decimal("1000.00"),
            rate_percent=Decimal("20.00"),
            subscription_price=Decimal("5000.00"),
        )

        client.force_login(owner_user)
        response = client.get(
            f"/dashboard/trainers/{source_trainer.id}/salary/"
            f"?date_from={closed_local_date.isoformat()}&date_to={closed_local_date.isoformat()}"
        )

        assert response.status_code == 200
        content = response.content.decode()
        assert "Период выплат закрыт" in content
        assert f"/dashboard/trainers/earnings/{earning.id}/correction/" not in content

    def test_admin_can_view_and_close_trainer_payroll_period(
        self,
        client: Client,
        club,
    ):
        from apps.attendance.tests.factories import CheckinFactory, ScheduleFactory
        from apps.trainers.models import TrainerEarning, TrainerPayrollPeriodClose
        from apps.trainers.tests.factories import TrainerFactory

        admin_user = UserFactory()
        ClubMembership.objects.create(user=admin_user, club=club, role=ClubMembership.Role.ADMIN)
        target_date = date.today().replace(day=1)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        source_trainer = TrainerFactory(club=club, first_name="Admin", last_name="Coach")
        schedule = ScheduleFactory(club=club, trainer=source_trainer, training_type=training_type)
        checkin = CheckinFactory(
            club=club,
            schedule=schedule,
            trainer=source_trainer,
            training_type=training_type,
            date=target_date,
        )
        TrainerEarning.objects.create(
            club=club,
            trainer=source_trainer,
            checkin=checkin,
            earning_type=TrainingType.Kind.PERSONAL,
            amount=Decimal("3000.00"),
            rate_percent=Decimal("50.00"),
            subscription_price=Decimal("6000.00"),
        )

        client.force_login(admin_user)
        response = client.get(
            f"/dashboard/trainers/{source_trainer.id}/salary/"
            f"?date_from={target_date.isoformat()}&date_to={target_date.isoformat()}"
        )
        assert response.status_code == 200
        assert "Закрыть период выплат по клубу" in response.content.decode()

        response = client.post(
            f"/dashboard/trainers/{source_trainer.id}/payroll-close/",
            {
                "date_from": target_date.isoformat(),
                "date_to": target_date.isoformat(),
                "reason": "Admin approved payroll",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 204
        assert TrainerPayrollPeriodClose.objects.for_club(club).filter(
            period_start=target_date,
            period_end=target_date,
            closed_by=admin_user,
        ).exists()

    def test_trainer_is_denied_management_payroll_routes_without_mutation(
        self,
        client: Client,
        club,
    ):
        from apps.attendance.tests.factories import CheckinFactory, ScheduleFactory
        from apps.trainers.models import TrainerEarning, TrainerEarningAdjustment, TrainerPayrollPeriodClose
        from apps.trainers.tests.factories import TrainerFactory

        trainer_user = UserFactory()
        ClubMembership.objects.create(user=trainer_user, club=club, role=ClubMembership.Role.TRAINER)
        target_date = date.today().replace(day=1)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        source_trainer = TrainerFactory(
            club=club,
            first_name="Denied",
            last_name="Coach",
            user=trainer_user,
        )
        target_trainer = TrainerFactory(club=club, first_name="Target", last_name="Coach")
        schedule = ScheduleFactory(club=club, trainer=source_trainer, training_type=training_type)
        checkin = CheckinFactory(
            club=club,
            schedule=schedule,
            trainer=source_trainer,
            training_type=training_type,
            date=target_date,
        )
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=source_trainer,
            checkin=checkin,
            earning_type=TrainingType.Kind.PERSONAL,
            amount=Decimal("3000.00"),
            rate_percent=Decimal("50.00"),
            subscription_price=Decimal("6000.00"),
        )

        client.force_login(trainer_user)
        salary_response = client.get(f"/dashboard/trainers/{source_trainer.id}/salary/")
        correction_get_response = client.get(f"/dashboard/trainers/earnings/{earning.id}/correction/")
        correction_post_response = client.post(
            f"/dashboard/trainers/earnings/{earning.id}/correction/",
            {
                "target_trainer_id": target_trainer.id,
                "reason": "Denied trainer mutation",
                "idempotency_key": "denied-trainer-correction",
            },
        )
        close_response = client.post(
            f"/dashboard/trainers/{source_trainer.id}/payroll-close/",
            {
                "date_from": target_date.isoformat(),
                "date_to": target_date.isoformat(),
                "reason": "Denied trainer close",
            },
        )

        assert salary_response.status_code == 403
        assert correction_get_response.status_code == 403
        assert correction_post_response.status_code == 403
        assert close_response.status_code == 403
        assert not TrainerEarningAdjustment.objects.for_club(club).filter(
            source_checkin=checkin,
            kind=TrainerEarningAdjustment.Kind.MANUAL_ADJUSTMENT,
        ).exists()
        assert not TrainerPayrollPeriodClose.objects.for_club(club).exists()

    def test_trainer_earning_correction_post_rejects_closed_period(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.attendance.tests.factories import CheckinFactory, ScheduleFactory
        from apps.trainers.models import TrainerEarning
        from apps.trainers.services import close_trainer_payroll_period
        from apps.trainers.tests.factories import TrainerFactory

        target_date = date.today().replace(day=1)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.PERSONAL)
        source_trainer = TrainerFactory(club=club, first_name="Actual", last_name="Coach")
        target_trainer = TrainerFactory(club=club, first_name="Owner", last_name="Coach")
        schedule = ScheduleFactory(club=club, trainer=source_trainer, training_type=training_type)
        checkin = CheckinFactory(
            club=club,
            schedule=schedule,
            trainer=source_trainer,
            training_type=training_type,
            date=target_date,
        )
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=source_trainer,
            checkin=checkin,
            earning_type=TrainingType.Kind.PERSONAL,
            amount=Decimal("3000.00"),
            rate_percent=Decimal("50.00"),
            subscription_price=Decimal("6000.00"),
        )
        close_trainer_payroll_period(
            club_id=club.id,
            period_start=target_date,
            period_end=target_date,
            reason="Payroll approved",
            actor_user_id=owner_user.id,
        )

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/trainers/earnings/{earning.id}/correction/",
            {
                "target_trainer_id": str(target_trainer.id),
                "reason": "Move payout",
                "idempotency_key": "closed-htmx-correction",
            },
        )

        assert response.status_code == 400
        content = response.content.decode()
        assert "Период выплат закрыт" in content

    def test_trainer_salary_shows_payment_manual_adjustment_audit(self, client: Client, owner_user, club):
        from apps.billing.tests.factories import PaymentFactory
        from apps.trainers.models import TrainerEarning, TrainerEarningAdjustment
        from apps.trainers.tests.factories import TrainerFactory

        local_salary_date = date.today().replace(day=1)
        training_type = TrainingTypeFactory(club=club, kind=TrainingType.Kind.GROUP)
        tariff = TariffFactory(training_type=training_type, price=Decimal("5000.00"))
        subscription = SubscriptionFactory(club=club, tariff=tariff)
        source_trainer = TrainerFactory(club=club, first_name="Sale", last_name="Coach")
        target_trainer = TrainerFactory(club=club, first_name="Target", last_name="Coach")
        payment = PaymentFactory(
            club=club,
            student=subscription.student,
            tariff=tariff,
            subscription=subscription,
            status=Payment.Status.CONFIRMED,
            verified_at=datetime.combine(
                local_salary_date - timedelta(days=1),
                time(22, 30),
                tzinfo=UTC,
            ),
        )
        earning = TrainerEarning.objects.create(
            club=club,
            trainer=source_trainer,
            payment=payment,
            earning_source=TrainerEarning.Source.SALE,
            earning_type=TrainingType.Kind.GROUP,
            amount=Decimal("1000.00"),
            rate_percent=Decimal("20.00"),
            subscription_price=Decimal("5000.00"),
        )
        TrainerEarningAdjustment.objects.create(
            club=club,
            trainer=source_trainer,
            amount_basis_snapshot=earning.amount,
            payable_amount_delta=Decimal("-1000.00"),
            affects_payroll=True,
            direction=TrainerEarningAdjustment.Direction.DEBIT,
            kind=TrainerEarningAdjustment.Kind.MANUAL_ADJUSTMENT,
            effective_date=local_salary_date,
            source_payment=payment,
            source_subscription=subscription,
            counterparty_trainer=target_trainer,
            reason="Sale attribution correction",
            created_by=owner_user,
            correction_group_id="33333333-3333-4333-8333-333333333333",
            idempotency_key="payment-adjustment-audit",
        )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/trainers/{source_trainer.id}/salary/")

        assert response.status_code == 200
        content = response.content.decode()
        assert "оплата #" in content
        assert "Sale attribution correction" in content
        assert "Target Coach" in content
        assert local_salary_date.strftime("%d.%m.%Y") in content
        assert "Ручная корректировка" in content
        assert "Списание" in content

    def test_trainer_salary_404_other_club(self, client: Client, owner_user, club, other_club):
        from apps.trainers.tests.factories import TrainerFactory

        trainer = TrainerFactory(club=other_club)
        client.force_login(owner_user)
        response = client.get(f"/dashboard/trainers/{trainer.id}/salary/")
        assert response.status_code == 404

    def test_unauthenticated_redirects(self, client: Client):
        response = client.get("/dashboard/trainers/")
        assert response.status_code == 302


class TestPnLReport:
    """ADMIN-09: P&L report with metrics."""

    @patch("apps.htmx_admin.views.reports.get_pnl_report")
    @patch("apps.htmx_admin.views.reports.get_business_metrics")
    def test_pnl_renders(self, mock_metrics, mock_pnl, client: Client, owner_user, club):
        mock_pnl.return_value = {
            "income": Decimal("100000"),
            "salary_expenses": Decimal("30000"),
            "manual_expenses": Decimal("10000"),
            "margin": Decimal("60000"),
        }
        mock_metrics.return_value = {
            "arpm": Decimal("5000"),
            "churn_rate": Decimal("0.10"),
            "retention_rate": Decimal("0.90"),
            "ltv": Decimal("50000"),
        }
        client.force_login(owner_user)
        response = client.get("/dashboard/reports/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "100 000" in content or "100000" in content or "100,000" in content

    @patch("apps.htmx_admin.views.reports.get_pnl_report")
    @patch("apps.htmx_admin.views.reports.get_business_metrics")
    def test_pnl_with_date_range(self, mock_metrics, mock_pnl, client: Client, owner_user, club):
        mock_pnl.return_value = {
            "income": Decimal("0"),
            "salary_expenses": Decimal("0"),
            "manual_expenses": Decimal("0"),
            "margin": Decimal("0"),
        }
        mock_metrics.return_value = {"arpm": None, "churn_rate": None, "retention_rate": None, "ltv": None}
        client.force_login(owner_user)
        response = client.get("/dashboard/reports/?date_from=2026-01-01&date_to=2026-01-31")
        assert response.status_code == 200

    @patch("apps.htmx_admin.views.reports.get_pnl_report")
    @patch("apps.htmx_admin.views.reports.get_business_metrics")
    def test_pnl_htmx_partial(self, mock_metrics, mock_pnl, client: Client, owner_user, club):
        mock_pnl.return_value = {
            "income": Decimal("0"),
            "salary_expenses": Decimal("0"),
            "manual_expenses": Decimal("0"),
            "margin": Decimal("0"),
        }
        mock_metrics.return_value = {"arpm": None, "churn_rate": None, "retention_rate": None, "ltv": None}
        client.force_login(owner_user)
        response = client.get("/dashboard/reports/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert b"<aside" not in response.content

    @patch("apps.htmx_admin.views.reports.get_pnl_report")
    @patch("apps.htmx_admin.views.reports.get_business_metrics")
    def test_pnl_shows_metrics(self, mock_metrics, mock_pnl, client: Client, owner_user, club):
        mock_pnl.return_value = {
            "income": Decimal("200000"),
            "salary_expenses": Decimal("50000"),
            "manual_expenses": Decimal("20000"),
            "margin": Decimal("130000"),
        }
        mock_metrics.return_value = {
            "arpm": Decimal("3500"),
            "churn_rate": Decimal("0.05"),
            "retention_rate": Decimal("0.95"),
            "ltv": Decimal("70000"),
        }
        client.force_login(owner_user)
        response = client.get("/dashboard/reports/")
        content = response.content.decode()
        assert "ФИНАНСЫ" in content
        assert "МАРЖА" in content
        assert "130 000" in content  # margin value from mock
        assert "200 000" in content  # income from mock
        assert "КЛЮЧЕВЫЕ МЕТРИКИ" in content

    def test_expense_form_preserves_report_date_range(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get(
            "/dashboard/reports/expense/create/?date_from=2026-02-01&date_to=2026-02-28"
        )

        assert response.status_code == 200
        content = response.content.decode()
        assert 'name="date_from" value="2026-02-01"' in content
        assert 'name="date_to" value="2026-02-28"' in content

    @patch("apps.htmx_admin.views.reports.create_expense")
    def test_expense_create_preserves_report_date_range_and_uses_service(
        self,
        mock_create_expense,
        client: Client,
        owner_user,
        club,
    ):
        client.force_login(owner_user)

        response = client.post(
            "/dashboard/reports/expense/create/",
            {
                "name": "February rent",
                "amount": "1234.50",
                "date": "2026-02-10",
                "category": "Аренда",
                "date_from": "2026-02-01",
                "date_to": "2026-02-28",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 204
        assert response["HX-Redirect"] == "/dashboard/reports/?date_from=2026-02-01&date_to=2026-02-28"
        mock_create_expense.assert_called_once_with(
            club_id=club.id,
            name="February rent",
            amount=Decimal("1234.50"),
            date=date(2026, 2, 10),
            category="Аренда",
            is_recurring=False,
        )

    @patch("apps.htmx_admin.views.reports.delete_expense")
    def test_expense_delete_preserves_report_date_range_and_uses_service(
        self,
        mock_delete_expense,
        client: Client,
        owner_user,
        club,
    ):
        expense = ExpenseFactory(club=club)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/reports/expense/{expense.id}/delete/?date_from=2026-02-01&date_to=2026-02-28",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 204
        assert response["HX-Redirect"] == "/dashboard/reports/?date_from=2026-02-01&date_to=2026-02-28"
        mock_delete_expense.assert_called_once_with(expense_id=expense.id, club_id=club.id)

    def test_pnl_export_preserves_date_range_and_returns_expense_rows(self, client: Client, owner_user, club):
        from openpyxl import load_workbook

        ExpenseFactory(
            club=club,
            name="February rent",
            amount=Decimal("1234.50"),
            date=date(2026, 2, 10),
            is_recurring=False,
        )
        ExpenseFactory(
            club=club,
            name="Outside range",
            amount=Decimal("9999.00"),
            date=date(2026, 3, 1),
            is_recurring=False,
        )

        client.force_login(owner_user)
        response = client.get("/dashboard/reports/export/?date_from=2026-02-01&date_to=2026-02-28")

        assert response.status_code == 200
        assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert response["Content-Disposition"] == 'attachment; filename="pnl_20260201_20260228.xlsx"'

        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        sheet = workbook["P&L"]
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        assert ("Финансовый отчёт: 01.02.2026 — 28.02.2026", None, None, None) in rows
        assert ("Прочие расходы", 1234.5, None, None) in rows
        assert ("February rent", 1234.5, "Нет", "10.02.2026") in rows
        assert not any(row[0] == "Outside range" for row in rows)

    def test_unauthenticated_redirects(self, client: Client):
        response = client.get("/dashboard/reports/")
        assert response.status_code == 302


class TestOnboardingWizard:
    """ADMIN-10: Onboarding wizard flow."""

    def test_wizard_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/onboarding/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "Step" in content
        assert "Добро пожаловать" in content
        assert "wizard-content" in content

    def test_wizard_step_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        # Start wizard first to create draft
        client.get("/dashboard/onboarding/")
        response = client.get("/dashboard/onboarding/step/1/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "Грейды" in content

    def test_step_2_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        response = client.get("/dashboard/onboarding/step/2/")
        assert response.status_code == 200
        assert "Шаг 2: Тренеры" in response.content.decode()

    def test_step_3_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        response = client.get("/dashboard/onboarding/step/3/")
        assert response.status_code == 200
        assert "Шаг 3: Расписание" in response.content.decode()

    def test_step_4_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        response = client.get("/dashboard/onboarding/step/4/")
        assert response.status_code == 200
        assert "Шаг 4: Ученики" in response.content.decode()

    def test_step_5_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        response = client.get("/dashboard/onboarding/step/5/")
        assert response.status_code == 200
        assert "Шаг 5: Тарифы" in response.content.decode()

    def test_skip_advances_step(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        response = client.post("/dashboard/onboarding/skip/1/")
        assert response.status_code == 200
        assert b"form" in response.content.lower()  # next step's form rendered

    def test_skip_get_is_read_only_and_returns_405(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        draft = OnboardingDraft.objects.for_club(club).get(is_completed=False)

        response = client.get("/dashboard/onboarding/skip/1/")

        assert response.status_code == 405
        draft.refresh_from_db()
        assert draft.current_step == 1
        assert draft.data["_skipped_steps"] == []

    def test_schedule_options_disambiguate_duplicate_names_with_stable_ids(
        self,
        client: Client,
        owner_user,
        club,
    ):
        trainer_a = TrainerFactory(club=club, first_name="Ivan", last_name="Petrov")
        trainer_b = TrainerFactory(club=club, first_name="Ivan", last_name="Petrov")
        location_a = LocationFactory(club=club, name="Main")
        location_b = LocationFactory(club=club, name="Main")
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")

        response = client.get("/dashboard/onboarding/step/3/")

        content = response.content.decode()
        assert f'value="existing:{trainer_a.id}"' in content
        assert f'value="existing:{trainer_b.id}"' in content
        assert f"Ivan Petrov · в клубе #{trainer_a.id}" in content
        assert f"Ivan Petrov · в клубе #{trainer_b.id}" in content
        assert f'value="{location_a.id}"' in content
        assert f'value="{location_b.id}"' in content
        assert f"Main · #{location_a.id}" in content
        assert f"Main · #{location_b.id}" in content

    def test_skip_last_step_finishes_with_post_and_redirects_to_dashboard(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        client.post("/dashboard/onboarding/skip/3/")
        response = client.post("/dashboard/onboarding/skip/5/")
        assert response.status_code == 302
        assert response.url == "/dashboard/"

    def test_skip_last_step_hx_finishes_and_redirects_to_dashboard(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        client.post("/dashboard/onboarding/skip/3/")
        response = client.post("/dashboard/onboarding/skip/5/", HTTP_HX_REQUEST="true")
        assert response.status_code == 204
        assert response["HX-Redirect"] == "/dashboard/"

    def test_finish_get_is_read_only_and_returns_405(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        draft = OnboardingDraft.objects.for_club(club).get(is_completed=False)
        response = client.get("/dashboard/onboarding/finish/")
        assert response.status_code == 405
        draft.refresh_from_db()
        assert draft.is_completed is False

    @patch("django_q.tasks.async_task")
    def test_finish_post_redirects_to_dashboard(self, mock_async, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        draft = OnboardingDraft.objects.for_club(club).get(is_completed=False)
        client.post("/dashboard/onboarding/skip/3/")
        response = client.post(
            "/dashboard/onboarding/finish/",
            data={"draft_id": draft.id},
            HTTP_HX_REQUEST="true",
        )
        assert response.status_code == 204
        assert response["HX-Redirect"] == "/dashboard/"

    def test_finish_post_without_csrf_is_rejected(self, owner_user, club):
        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(owner_user)
        draft = OnboardingDraft.objects.create(
            club=club,
            data={"_schema_version": 2, "_skipped_steps": [3]},
        )

        response = csrf_client.post("/dashboard/onboarding/finish/", data={"draft_id": draft.id})

        assert response.status_code == 403
        draft.refresh_from_db()
        assert draft.is_completed is False

    def test_finish_error_returns_to_populated_schedule_with_api_error_code(
        self,
        client: Client,
        owner_user,
        club,
    ):
        client.force_login(owner_user)
        draft = OnboardingDraft.objects.create(
            club=club,
            current_step=5,
            data={
                "_schema_version": 2,
                "_skipped_steps": [],
                "3": {
                    "schedules": [
                        {
                            "day": 0,
                            "start": "10:00",
                            "end": "11:00",
                            "group": "Adults",
                            "trainer_id": None,
                            "trainer_ref": None,
                            "location_id": None,
                            "legacy_trainer_name": "Ivan",
                        }
                    ]
                },
            },
        )

        response = client.post(
            "/dashboard/onboarding/finish/",
            data={"draft_id": draft.id},
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        content = response.content.decode()
        assert "Шаг 3: Расписание" in content
        assert 'data-error-code="onboarding_schedule_trainer_unresolved"' in content
        assert "Adults" in content
        assert "Ivan" in content
        draft.refresh_from_db()
        assert draft.is_completed is False

    def test_wizard_resumes_from_draft(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        # Start wizard
        client.get("/dashboard/onboarding/")
        # Skip step 1 -> now on step 2
        client.post("/dashboard/onboarding/skip/1/")
        # Revisit wizard -- should resume from step 2
        response = client.get("/dashboard/onboarding/")
        assert response.status_code == 200
        assert b"form" in response.content.lower()  # step content rendered

    def test_step_post_normalizes_checkboxes_and_indexed_rows(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        client.get("/dashboard/onboarding/")
        location = LocationFactory(club=club)

        response = client.post(
            "/dashboard/onboarding/step/1/",
            data={"disciplines": ["bjj", "boxing"]},
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        draft = OnboardingDraft.objects.for_club(club).get(is_completed=False)
        assert draft.data["1"] == {"disciplines": ["bjj", "boxing"], "use_templates": False}

        response = client.post(
            "/dashboard/onboarding/step/2/",
            data={
                "trainers[0].client_ref": "00000000-0000-0000-0000-000000000001",
                "trainers[0].first_name": "",
                "trainers[0].last_name": "",
                "trainers[0].phone": "",
                "trainers[1].client_ref": "00000000-0000-0000-0000-000000000002",
                "trainers[1].first_name": "Ivan",
                "trainers[1].last_name": "Petrov",
                "trainers[1].phone": "+79007654321",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        draft.refresh_from_db()
        assert draft.data["2"]["trainers"] == [
            {
                "client_ref": "00000000-0000-0000-0000-000000000002",
                "first_name": "Ivan",
                "last_name": "Petrov",
                "phone": "+79007654321",
            }
        ]

        response = client.post(
            "/dashboard/onboarding/step/3/",
            data={
                "schedules[0].day": "2",
                "schedules[0].start": "19:00",
                "schedules[0].end": "20:30",
                "schedules[0].group": "Adults",
                "schedules[0].trainer_choice": "draft:00000000-0000-0000-0000-000000000002",
                "schedules[0].location_id": str(location.id),
                "schedules[0].legacy_trainer_name": "",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        draft.refresh_from_db()
        assert draft.data["3"]["schedules"] == [
            {
                "day": 2,
                "start": "19:00",
                "end": "20:30",
                "group": "Adults",
                "trainer_id": None,
                "trainer_ref": "00000000-0000-0000-0000-000000000002",
                "location_id": location.id,
                "legacy_trainer_name": "",
            }
        ]

        response = client.post(
            "/dashboard/onboarding/step/4/",
            data={
                "students[0].first_name": "Nina",
                "students[0].last_name": "Ivanova",
                "students[0].phone": "+79001234567",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        draft.refresh_from_db()
        assert draft.data["4"]["students"] == [
            {"first_name": "Nina", "last_name": "Ivanova", "phone": "+79001234567"}
        ]

        response = client.post(
            "/dashboard/onboarding/step/5/",
            data={
                "tariffs[0].name": "Monthly",
                "tariffs[0].price": "5000",
                "tariffs[0].training_limit": "",
                "tariffs[0].duration_days": "30",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 204
        assert response["HX-Redirect"] == "/dashboard/"
        draft.refresh_from_db()
        assert draft.data["5"]["tariffs"] == [
            {"name": "Monthly", "price": 5000, "training_limit": None, "duration_days": 30}
        ]
        assert draft.is_completed is True
        schedule = Schedule.objects.for_club(club).get(group_name="Adults")
        assert schedule.trainer.first_name == "Ivan"
        assert schedule.location_id == location.id

    def test_unauthenticated_redirects(self, client: Client):
        response = client.get("/dashboard/onboarding/")
        assert response.status_code == 302


class TestRetentionTasks:
    """ADMIN-11: Retention tasks view."""

    def test_retention_tasks_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/retention/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "RETENTION TASKS" in content

    def test_retention_tasks_htmx_partial(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/retention/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert b"<aside" not in response.content

    def test_retention_tasks_shows_data(self, client: Client, owner_user, club):
        from apps.retention.tests.factories import RetentionTaskFactory

        task = RetentionTaskFactory(club=club)
        client.force_login(owner_user)
        response = client.get("/dashboard/retention/")
        content = response.content.decode()
        assert task.student.first_name in content

    def test_retention_tasks_status_filter(self, client: Client, owner_user, club):
        from django.utils import timezone

        from apps.retention.tests.factories import RetentionTaskFactory

        RetentionTaskFactory(club=club, resolved_at=None)
        RetentionTaskFactory(club=club, resolved_at=timezone.now())
        client.force_login(owner_user)
        response = client.get("/dashboard/retention/?status=open")
        assert response.status_code == 200

    def test_retention_tasks_tenant_isolation(self, client: Client, owner_user, club, other_club):
        from apps.retention.tests.factories import RetentionTaskFactory

        task_mine = RetentionTaskFactory(club=club)
        task_other = RetentionTaskFactory(club=other_club)
        client.force_login(owner_user)
        response = client.get("/dashboard/retention/")
        content = response.content.decode()
        assert task_mine.student.first_name in content
        assert task_other.student.first_name not in content

    def test_unauthenticated_redirects(self, client: Client):
        response = client.get("/dashboard/retention/")
        assert response.status_code == 302


class TestPushNotifications:
    """ADMIN-12: Push notifications management."""

    def test_push_form_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/notifications/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "PUSH-УВЕДОМЛЕНИЯ" in content
        assert "segment_type" in content
        assert "text" in content

    def test_push_form_htmx_partial(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/notifications/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert b"<aside" not in response.content

    @patch("apps.htmx_admin.views.notifications.send_mass_notification")
    def test_push_send(self, mock_send, client: Client, owner_user, club):
        from apps.notifications.models import MassNotification

        mock_notif = MassNotification(
            id=1,
            club=club,
            text="Test",
            segment_type="club",
            segment_filter={},
            recipient_count=5,
            sent_by=owner_user,
        )
        mock_send.return_value = mock_notif
        client.force_login(owner_user)
        response = client.post(
            "/dashboard/notifications/",
            {
                "text": "Test notification",
                "segment_type": "club",
                "segment_value": "",
            },
        )
        assert response.status_code == 200
        assert b"5 recipients" in response.content
        mock_send.assert_called_once()

    def test_push_preview_uses_current_enrollment_and_ignores_training_reminder_opt_out(
        self,
        client: Client,
        owner_user,
        club,
    ):
        schedule = ScheduleFactory(club=club)
        deliverable_user = UserFactory()
        opted_out_user = UserFactory()
        deliverable_student = StudentFactory(club=club, user=deliverable_user, status=Student.Status.ACTIVE)
        opted_out_student = StudentFactory(club=club, user=opted_out_user, status=Student.Status.ACTIVE)
        ClubMembership.objects.create(
            club=club,
            user=deliverable_user,
            role=ClubMembership.Role.STUDENT,
        )
        ClubMembership.objects.create(
            club=club,
            user=opted_out_user,
            role=ClubMembership.Role.STUDENT,
        )
        ScheduleEnrollment.objects.create(
            club=club,
            schedule=schedule,
            student=deliverable_student,
            status=ScheduleEnrollment.Status.ACTIVE,
        )
        ScheduleEnrollment.objects.create(
            club=club,
            schedule=schedule,
            student=opted_out_student,
            status=ScheduleEnrollment.Status.ACTIVE,
        )
        PushSubscriptionFactory(user=deliverable_user)
        PushSubscriptionFactory(user=opted_out_user)
        NotificationPreference.objects.create(
            user=opted_out_user,
            disabled_categories=["training_reminders"],
        )

        client.force_login(owner_user)
        response = client.post(
            "/dashboard/notifications/preview/",
            {
                "segment_type": "group",
                "segment_value": str(schedule.id),
            },
        )

        assert response.status_code == 200
        assert b"<strong>2</strong>" in response.content

    def test_push_preview_renders_invalid_segment_error(
        self,
        client: Client,
        owner_user,
        club,
    ):
        client.force_login(owner_user)
        response = client.post(
            "/dashboard/notifications/preview/",
            {
                "segment_type": "group",
                "segment_value": "not-an-id",
            },
        )

        assert response.status_code == 200
        assert "Укажите корректный ID группы" in response.content.decode()

    @patch("apps.htmx_admin.views.notifications.send_mass_notification")
    def test_push_send_invalid_segment_does_not_call_service_or_create_row(
        self,
        mock_send,
        client: Client,
        owner_user,
        club,
    ):
        client.force_login(owner_user)
        response = client.post(
            "/dashboard/notifications/",
            {
                "text": "Must not send",
                "segment_type": "location",
                "segment_value": "not-an-id",
            },
        )

        assert response.status_code == 200
        assert "Укажите корректный ID локации" in response.content.decode()
        mock_send.assert_not_called()
        assert not MassNotification.objects.for_club(club).exists()

    def test_push_empty_text_shows_error(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.post(
            "/dashboard/notifications/",
            {
                "text": "",
                "segment_type": "club",
            },
        )
        assert response.status_code == 200
        assert b"Message text is required" in response.content

    def test_unauthenticated_redirects(self, client: Client):
        response = client.get("/dashboard/notifications/")
        assert response.status_code == 302


class TestClubSettings:
    """ADMIN-13: Club settings (branding + general)."""

    def test_settings_page_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/settings/general/")
        assert response.status_code == 200
        content = response.content.decode()
        assert "Внешний вид клуба" in content
        assert "primary_color" in content
        assert "PNG, JPG или WEBP. Максимум 2 МБ." in content
        assert "PNG, JPG, SVG или WEBP" not in content

    def test_settings_save_branding(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.post(
            "/dashboard/settings/general/",
            {
                "club_name_display": "New Club Name",
                "primary_color": "#FF0000",
                "accent_color": "#00FF00",
                "freeze_max_days": "30",
                "max_push_per_week": "3",
            },
        )
        # Save returns 204 with HX-Refresh header to trigger full page reload.
        assert response.status_code in (200, 204)
        from apps.clubs.models import ClubSettings

        settings = ClubSettings.objects.get(club=club)
        assert settings.club_name_display == "New Club Name"
        assert settings.primary_color == "#FF0000"

    def test_settings_save_shows_success(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.post(
            "/dashboard/settings/general/",
            {
                "club_name_display": "Test",
                "primary_color": "#000000",
                "accent_color": "#FF6B00",
                "freeze_max_days": "30",
                "max_push_per_week": "3",
            },
        )
        # Save returns 204 with HX-Refresh header; client refreshes to show success.
        assert response.status_code in (200, 204)
        assert response.headers.get("HX-Refresh") == "true" or "Настройки сохранены".encode() in response.content

    def test_settings_cache_invalidation(self, client: Client, owner_user, club):
        from django.core.cache import cache

        cache.set(f"club_branding:{club.id}", "stale_cached_value", timeout=300)
        client.force_login(owner_user)
        client.post(
            "/dashboard/settings/general/",
            {
                "club_name_display": "Updated",
                "primary_color": "#111111",
                "accent_color": "#222222",
                "freeze_max_days": "30",
                "max_push_per_week": "3",
            },
        )
        # Cache was invalidated -- old stale value is gone.
        # Context processor may re-cache during render, so check value is fresh, not stale.
        cached = cache.get(f"club_branding:{club.id}")
        assert cached != "stale_cached_value"

    def test_settings_tenant_isolation(self, client: Client, owner_user, club, other_club):
        from apps.clubs.models import ClubSettings

        ClubSettings.objects.get_or_create(club=other_club, defaults={"club_name_display": "Other"})
        client.force_login(owner_user)
        response = client.get("/dashboard/settings/")
        content = response.content.decode()
        assert "Other" not in content

    def test_settings_htmx_partial(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/settings/general/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert b"<aside" not in response.content

    def test_unauthenticated_redirects(self, client: Client):
        response = client.get("/dashboard/settings/")
        assert response.status_code == 302


class TestNotificationSettings:
    def test_notifications_settings_renders_and_backfills_missing_templates(self, client: Client, owner_user, club):
        existing = NotificationTemplateFactory(
            club=club,
            trigger_type=NotificationTemplate.TriggerType.SUB_EXPIRY_7D,
            title_template="Keep existing title",
        )

        client.force_login(owner_user)
        response = client.get("/dashboard/settings/notifications/")

        assert response.status_code == 200
        body = response.content.decode()
        assert "Шаблоны уведомлений" in body
        assert "для учеников, родителей и тренерских задач" in body
        assert "Абонементы" in body
        assert "Вовлечение" in body
        _assert_no_nested_forms(body)
        assert NotificationTemplate.objects.for_club(club).count() == len(NotificationTemplate.TriggerType.values)
        assert NotificationTemplate.objects.for_club(club).filter(
            trigger_type=NotificationTemplate.TriggerType.TRAINING_REMINDER,
        ).exists()
        assert NotificationTemplate.objects.for_club(club).filter(
            trigger_type=NotificationTemplate.TriggerType.TRAINING_REMINDER_24H,
        ).exists()
        existing.refresh_from_db()
        assert existing.title_template == "Keep existing title"

    def test_owner_can_save_notification_timing_settings(self, client: Client, owner_user, club):
        from apps.clubs.models import ClubSettings

        client.force_login(owner_user)
        response = client.post(
            "/dashboard/settings/notifications/",
            {
                "max_push_per_week": "6",
                "feedback_delay_hours": "5",
                "quiet_hours_start": "22:30",
                "quiet_hours_end": "08:15",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert "Настройки сохранены" in response.content.decode()
        settings = ClubSettings.objects.get(club=club)
        assert settings.max_push_per_week == 6
        assert settings.feedback_delay_hours == 5
        assert settings.quiet_hours_start == time(22, 30)
        assert settings.quiet_hours_end == time(8, 15)

    def test_notification_timing_settings_reject_invalid_values_without_mutation(
        self,
        client: Client,
        owner_user,
        club,
    ):
        from apps.clubs.models import ClubSettings

        settings, _ = ClubSettings.objects.get_or_create(club=club)
        settings.max_push_per_week = 3
        settings.feedback_delay_hours = 2
        settings.quiet_hours_start = time(21, 0)
        settings.quiet_hours_end = time(9, 0)
        settings.save()

        client.force_login(owner_user)
        response = client.post(
            "/dashboard/settings/notifications/",
            {
                "max_push_per_week": "51",
                "feedback_delay_hours": "12",
                "quiet_hours_start": "22:30",
                "quiet_hours_end": "08:15",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        content = response.content.decode()
        assert "Максимум уведомлений" in content
        assert "Настройки сохранены" not in content
        settings.refresh_from_db()
        assert settings.max_push_per_week == 3
        assert settings.feedback_delay_hours == 2
        assert settings.quiet_hours_start == time(21, 0)
        assert settings.quiet_hours_end == time(9, 0)

    def test_owner_can_update_notification_template_form(self, client: Client, owner_user, club):
        template = NotificationTemplateFactory(
            club=club,
            trigger_type=NotificationTemplate.TriggerType.TRAINING_REMINDER,
            is_enabled=True,
        )

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/settings/notifications/templates/{template.id}/form/",
            {
                "title_template": "  Custom title  ",
                "body_template": "  Custom body for {name}  ",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 204
        assert response.headers["HX-Redirect"] == "/dashboard/settings/notifications/"
        template.refresh_from_db()
        assert template.title_template == "Custom title"
        assert template.body_template == "Custom body for {name}"
        assert template.is_enabled is False

    def test_owner_can_update_expiry_template_days_before(self, client: Client, owner_user, club):
        template = NotificationTemplateFactory(
            club=club,
            trigger_type=NotificationTemplate.TriggerType.SUB_EXPIRY_7D,
            title_template="Expiry title",
            body_template="Expiry body {days}",
            days_before=7,
            is_enabled=True,
        )

        client.force_login(owner_user)
        form_response = client.get(
            f"/dashboard/settings/notifications/templates/{template.id}/form/",
            HTTP_HX_REQUEST="true",
        )

        assert form_response.status_code == 200
        form_body = form_response.content.decode()
        assert 'name="days_before"' in form_body
        assert 'value="7"' in form_body

        response = client.post(
            f"/dashboard/settings/notifications/templates/{template.id}/form/",
            {
                "title_template": "Updated expiry title",
                "body_template": "Updated expiry body {days}",
                "is_enabled": "on",
                "days_before": "10",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 204
        template.refresh_from_db()
        assert template.title_template == "Updated expiry title"
        assert template.body_template == "Updated expiry body {days}"
        assert template.days_before == 10
        assert template.is_enabled is True

    def test_non_expiry_template_form_hides_days_before(self, client: Client, owner_user, club):
        template = NotificationTemplateFactory(
            club=club,
            trigger_type=NotificationTemplate.TriggerType.TRAINING_REMINDER,
            is_enabled=True,
        )

        client.force_login(owner_user)
        response = client.get(
            f"/dashboard/settings/notifications/templates/{template.id}/form/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert 'name="days_before"' not in response.content.decode()

    def test_follow_up_template_form_uses_trainer_task_safe_helper_copy(self, client: Client, owner_user, club):
        template = NotificationTemplateFactory(
            club=club,
            trigger_type=NotificationTemplate.TriggerType.FOLLOW_UP,
            is_enabled=True,
        )

        client.force_login(owner_user)
        response = client.get(
            f"/dashboard/settings/notifications/templates/{template.id}/form/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        body = response.content.decode()
        assert "Напоминание по задаче" in body
        assert "получатель зависит от типа шаблона" in body
        assert "данные ученика при отправке" not in body

    def test_notification_template_form_rejects_blank_text_without_mutation(self, client: Client, owner_user, club):
        template = NotificationTemplateFactory(
            club=club,
            trigger_type=NotificationTemplate.TriggerType.MISSED_TRAINING,
            title_template="Original title",
            body_template="Original body",
            is_enabled=True,
        )

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/settings/notifications/templates/{template.id}/form/",
            {
                "title_template": "   ",
                "body_template": "Updated body",
                "is_enabled": "on",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert "Заголовок и текст обязательны" in response.content.decode()
        template.refresh_from_db()
        assert template.title_template == "Original title"
        assert template.body_template == "Original body"
        assert template.is_enabled is True

    def test_owner_can_toggle_notification_template_with_tenant_scope(
        self,
        client: Client,
        owner_user,
        club,
        other_club,
    ):
        from apps.clubs.models import ClubSettings

        settings, _ = ClubSettings.objects.get_or_create(club=club)
        settings.max_push_per_week = 3
        settings.feedback_delay_hours = 2
        settings.save()
        template = NotificationTemplateFactory(
            club=club,
            trigger_type=NotificationTemplate.TriggerType.PARENT_CHECKIN,
            is_enabled=True,
        )
        other_template = NotificationTemplateFactory(
            club=other_club,
            trigger_type=NotificationTemplate.TriggerType.PARENT_CHECKIN,
            is_enabled=True,
        )

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/settings/notifications/templates/{template.id}/toggle/",
            {
                "max_push_per_week": "50",
                "feedback_delay_hours": "48",
            },
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert "Настройки сохранены" not in response.content.decode()
        template.refresh_from_db()
        settings.refresh_from_db()
        assert template.is_enabled is False
        assert settings.max_push_per_week == 3
        assert settings.feedback_delay_hours == 2

        forbidden_response = client.post(
            f"/dashboard/settings/notifications/templates/{other_template.id}/toggle/",
            HTTP_HX_REQUEST="true",
        )

        assert forbidden_response.status_code == 404
        other_template.refresh_from_db()
        assert other_template.is_enabled is True

    def test_trainer_cannot_use_notification_settings(self, client: Client, trainer_user, club):
        template = NotificationTemplateFactory(club=club)

        client.force_login(trainer_user)

        settings_response = client.get("/dashboard/settings/notifications/")
        form_response = client.get(
            f"/dashboard/settings/notifications/templates/{template.id}/form/",
            HTTP_HX_REQUEST="true",
        )
        toggle_response = client.post(
            f"/dashboard/settings/notifications/templates/{template.id}/toggle/",
            HTTP_HX_REQUEST="true",
        )

        assert settings_response.status_code == 403
        assert form_response.status_code == 403
        assert toggle_response.status_code == 403

    def test_admin_push_disable_posts_server_unsubscribe(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/settings/notifications/")

        assert response.status_code == 200
        body = response.content.decode()
        assert "sub.unsubscribe()" in body
        assert "/api/notifications/unsubscribe/" in body
        assert "endpoint: sub.endpoint" in body


class TestSettingsDocuments:
    def test_document_type_form_renders_in_slide_over(self, client: Client, owner_user):
        client.force_login(owner_user)
        response = client.get("/dashboard/settings/documents/types/form/", HTTP_HX_REQUEST="true")
        assert response.status_code == 200
        assert 'name="name"' in response.content.decode()

    def test_settings_documents_tab_renders(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/settings/documents/")
        assert response.status_code == 200
        body = response.content.decode()
        assert "Документы клуба" in body

    def test_settings_tabs_include_documents(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.get("/dashboard/settings/catalog/")
        body = response.content.decode()
        assert "/dashboard/settings/documents/" in body
        assert "Документы" in body

    def test_trainer_cannot_open_documents_settings(self, client: Client, trainer_user, club):
        client.force_login(trainer_user)
        response = client.get("/dashboard/settings/documents/")
        assert response.status_code == 403

    def test_owner_can_create_document_type_from_settings(self, client: Client, owner_user, club):
        client.force_login(owner_user)
        response = client.post(
            "/dashboard/settings/documents/types/form/",
            {
                "name": "Медсправка",
                "description": "Для допуска к тренировкам",
                "scope": "children",
                "is_required": "on",
            },
            HTTP_HX_REQUEST="true",
        )
        assert response.status_code in (200, 204)
        created = DocumentTypeFactory._meta.model.objects.get(club=club, name="Медсправка")
        assert created.is_required is True
        assert response.headers.get("HX-Retarget") == "#content"
        assert "closeSlideOver" in (response.headers.get("HX-Trigger") or "")

    def test_owner_can_toggle_document_type_from_settings(self, client: Client, owner_user, club):
        document_type = DocumentTypeFactory(club=club, name="Паспорт", is_active=True)
        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/settings/documents/types/{document_type.id}/toggle/",
            HTTP_HX_REQUEST="true",
        )
        assert response.status_code == 200
        document_type.refresh_from_db()
        assert document_type.is_active is False
        assert "Паспорт" in response.content.decode()

    def test_toggle_duplicate_name_shows_inline_error(self, client: Client, owner_user, club):
        DocumentTypeFactory(club=club, name="Паспорт", is_active=True)
        archived_same_name = DocumentTypeFactory(club=club, name="Паспорт", is_active=False)
        client.force_login(owner_user)

        response = client.post(
            f"/dashboard/settings/documents/types/{archived_same_name.id}/toggle/",
            HTTP_HX_REQUEST="true",
        )

        assert response.status_code == 200
        assert "существует" in response.content.decode().lower() or "duplicate" in response.content.decode().lower()


class TestStudentDocumentsAdminCard:
    def test_student_card_documents_empty_state(self, client: Client, owner_user, club):
        student = StudentFactory(club=club)
        client.force_login(owner_user)
        response = client.get(f"/dashboard/students/{student.id}/card/")
        assert response.status_code == 200
        assert "Для этого ученика не настроены документы" in response.content.decode()

    def test_student_card_renders_documents_block_with_archived_history(self, client: Client, owner_user, club):
        student = StudentFactory(club=club, is_child=False)
        active_type = DocumentTypeFactory(club=club, name="Договор", scope="all", is_required=True)
        archived_type = DocumentTypeFactory(club=club, name="Старый документ", scope="adults", is_active=False)
        StudentDocumentFactory(
            student=student,
            document_type=archived_type,
            is_provided=True,
            file=SimpleUploadedFile("old.pdf", b"%PDF-1.4 archived", content_type="application/pdf"),
        )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/students/{student.id}/card/")
        body = response.content.decode()

        assert response.status_code == 200
        assert "ДОКУМЕНТЫ" in body
        assert active_type.name in body
        assert archived_type.name in body
        assert "АРХИВ" in body or "Неактив" in body

    def test_owner_can_mark_student_document_provided(self, client: Client, owner_user, club):
        student = StudentFactory(club=club)
        document_type = DocumentTypeFactory(club=club, name="Согласие")
        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/students/{student.id}/documents/mark/",
            {"document_type_id": str(document_type.id), "is_provided": "1"},
            HTTP_HX_REQUEST="true",
        )
        assert response.status_code == 200
        body = response.content.decode()
        assert 'id="student-documents-block"' in body
        assert "ПРЕДОСТАВЛЕН" in body or "ЗАГРУЖЕН" in body

    def test_owner_can_upload_student_document(self, client: Client, owner_user, club):
        student = StudentFactory(club=club)
        document_type = DocumentTypeFactory(club=club, name="Паспорт")
        upload = SimpleUploadedFile("passport.pdf", b"%PDF-1.4 passport", content_type="application/pdf")

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/students/{student.id}/documents/upload/",
            {"document_type_id": str(document_type.id), "file": upload},
            HTTP_HX_REQUEST="true",
        )
        assert response.status_code == 200
        body = response.content.decode()
        assert 'id="student-documents-block"' in body
        assert "ЗАГРУЖЕН" in body
        assert "Открыть файл" in body

    def test_invalid_upload_shows_error(self, client: Client, owner_user, club):
        student = StudentFactory(club=club)
        document_type = DocumentTypeFactory(club=club, name="Паспорт")
        upload = SimpleUploadedFile("bad.exe", b"bad", content_type="application/x-msdownload")

        client.force_login(owner_user)
        response = client.post(
            f"/dashboard/students/{student.id}/documents/upload/",
            {"document_type_id": str(document_type.id), "file": upload},
            HTTP_HX_REQUEST="true",
        )
        assert response.status_code == 422
        assert "Unsupported file type" in response.content.decode()

    def test_owner_can_open_uploaded_document(self, client: Client, owner_user, club):
        student = StudentFactory(club=club)
        document_type = DocumentTypeFactory(club=club, name="Паспорт")
        student_document = StudentDocumentFactory(
            student=student,
            document_type=document_type,
            is_provided=True,
            file=SimpleUploadedFile("passport.pdf", b"%PDF-1.4 open", content_type="application/pdf"),
        )

        client.force_login(owner_user)
        response = client.get(
            f"/dashboard/students/{student.id}/documents/{student_document.id}/open/",
        )
        assert response.status_code == 200
        assert response.get("Content-Type") == "application/pdf"
        assert response.get("X-Content-Type-Options") == "nosniff"
        assert response.get("Content-Disposition", "").startswith("attachment;")

    def test_open_document_does_not_serve_historical_html_inline(self, client: Client, owner_user, club):
        student = StudentFactory(club=club)
        document_type = DocumentTypeFactory(club=club, name="Паспорт")
        student_document = StudentDocumentFactory(
            student=student,
            document_type=document_type,
            is_provided=True,
            file=SimpleUploadedFile(
                "legacy.html",
                b"<html><script>alert('xss')</script></html>",
                content_type="text/html",
            ),
        )

        client.force_login(owner_user)
        response = client.get(
            f"/dashboard/students/{student.id}/documents/{student_document.id}/open/",
        )

        assert response.status_code == 200
        assert response.get("Content-Type") != "text/html"
        assert response.get("X-Content-Type-Options") == "nosniff"
        assert "inline" not in response.get("Content-Disposition", "").lower()

    def test_open_document_returns_404_for_other_student(self, client: Client, owner_user, club):
        student = StudentFactory(club=club)
        other_student = StudentFactory(club=club)
        document_type = DocumentTypeFactory(club=club, name="Паспорт")
        student_document = StudentDocumentFactory(
            student=other_student,
            document_type=document_type,
            is_provided=True,
            file=SimpleUploadedFile("passport.pdf", b"%PDF-1.4 open", content_type="application/pdf"),
        )

        client.force_login(owner_user)
        response = client.get(f"/dashboard/students/{student.id}/documents/{student_document.id}/open/")

        assert response.status_code == 404

    def test_trainer_cannot_use_student_document_routes(self, client: Client, trainer_user, club):
        student = StudentFactory(club=club)
        document_type = DocumentTypeFactory(club=club, name="Паспорт")
        client.force_login(trainer_user)
        response = client.post(
            f"/dashboard/students/{student.id}/documents/mark/",
            {"document_type_id": str(document_type.id), "is_provided": "1"},
            HTTP_HX_REQUEST="true",
        )
        assert response.status_code == 403

    def test_trainer_cannot_upload_or_open_student_documents(self, client: Client, trainer_user, club):
        student = StudentFactory(club=club)
        document_type = DocumentTypeFactory(club=club, name="Паспорт")
        student_document = StudentDocumentFactory(
            student=student,
            document_type=document_type,
            is_provided=True,
            file=SimpleUploadedFile("passport.pdf", b"%PDF-1.4 open", content_type="application/pdf"),
        )
        upload = SimpleUploadedFile("passport.pdf", b"%PDF-1.4 upload", content_type="application/pdf")

        client.force_login(trainer_user)
        upload_response = client.post(
            f"/dashboard/students/{student.id}/documents/upload/",
            {"document_type_id": str(document_type.id), "file": upload},
            HTTP_HX_REQUEST="true",
        )
        open_response = client.get(f"/dashboard/students/{student.id}/documents/{student_document.id}/open/")

        assert upload_response.status_code == 403
        assert open_response.status_code == 403

    def test_invalid_student_document_requests_return_404(self, client: Client, owner_user):
        client.force_login(owner_user)

        mark_response = client.post(
            "/dashboard/students/999999/documents/mark/",
            {"document_type_id": "bad", "is_provided": "1"},
            HTTP_HX_REQUEST="true",
        )
        upload_response = client.post(
            "/dashboard/students/999999/documents/upload/",
            {"document_type_id": "1"},
            HTTP_HX_REQUEST="true",
        )

        assert mark_response.status_code == 404
        assert upload_response.status_code == 404
