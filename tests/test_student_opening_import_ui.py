from dataclasses import replace
from io import BytesIO
from unittest.mock import patch

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from apps.billing.models import Payment
from apps.billing.tests.test_opening_issuer import command as opening_command
from apps.billing.tests.test_opening_issuer import group_command as opening_group_command
from apps.clubs.models import ClubMembership
from apps.clubs.tests.factories import ClubFactory
from apps.students.imports.parsing import ENTITLEMENTS_SHEET, WorkbookRow
from apps.students.imports.runner import run_batch_chunk
from apps.students.imports.schemas import COLUMNS
from apps.students.imports.storage import private_path
from apps.students.imports.workbooks import workbook_bytes
from apps.students.models import OpeningImportBatch, OpeningImportItemReceipt, Student

pytestmark = pytest.mark.django_db
command = opening_command
group_command = opening_group_command
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def book_case(club, command, settings, tmp_path, client):
    actor, terms = command
    settings.STUDENT_IMPORT_PRIVATE_ROOT = tmp_path / "private"
    client.force_login(actor)

    def content(extra=(), source_terms=None):
        rows = []
        for index, source in enumerate((source_terms or terms, *extra), start=2):
            payload = source.as_payload()
            values = {column: payload[field] for field, column in COLUMNS.items()}
            values.update(
                {
                    "Исходная цена": payload["paid_amount"],
                    "Валюта": "RUB",
                    "Исторический долг": 0,
                    "Заморожен": False,
                    "Тип пакета": "Групповой" if source.training_group_id else "Персональный",
                }
            )
            rows.append(WorkbookRow(ENTITLEMENTS_SHEET, index, values))
        return workbook_bytes(club_id=club.id, actor_user_id=actor.id, namespace=terms.source_namespace, rows=rows)

    return actor, terms, content


def upload(client, content):
    return client.post(
        "/dashboard/students/opening-import/",
        {"workbook": SimpleUploadedFile("source.xlsx", content, content_type=MIME)},
    )


def test_preview_apply_result_download_and_renamed_replay(client, club, book_case, settings):
    _, _, content = book_case
    response = upload(client, content())
    assert response.status_code == 302
    batch = OpeningImportBatch.objects.get()
    page = client.get(response.url)
    assert page.status_code == 200 and b"data-import-ready" in page.content
    assert "Применить 1 готовых записей".encode() in page.content
    assert not Student.objects.for_club(club).exists() and not Payment.objects.for_club(club).exists()
    assert "no-store" in page["Cache-Control"]
    preview = batch.previews.get()
    data = {"revision": preview.revision, "selection_token": str(preview.selection_token)}
    with patch("django_q.tasks.async_task") as queue:
        response = client.post(f"/dashboard/students/opening-import/{batch.id}/apply/", data, HTTP_HX_REQUEST="true")
    assert response.status_code == 200
    queue.assert_called_once_with("apps.students.imports.runner.run_batch_worker", club_id=club.id, batch_id=batch.id)
    assert not Payment.objects.for_club(club).exists()
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["status"] == "completed"
    page = client.get(f"/dashboard/students/opening-import/{batch.id}/")
    assert "Открыть ученика".encode() in page.content and "Применено полностью".encode() in page.content
    assert OpeningImportItemReceipt.objects.get().channel == "htmx"
    result = client.get(f"/dashboard/students/opening-import/{batch.id}/export/")
    assert result.status_code == 200 and "no-store" in result["Cache-Control"]
    exported = b"".join(result.streaming_content)
    # Consuming the test-client streaming iterator already closes this response.
    workbook = load_workbook(BytesIO(exported), read_only=True)
    assert workbook["_Источник"]["B1"].value == batch.source_namespace
    assert "Как заполнить" in workbook.sheetnames
    workbook.close()
    settings.STUDENT_OPENING_IMPORT_ENABLED = False
    response = upload(client, exported)
    second = OpeningImportBatch.objects.latest("id")
    replay = second.previews.get()
    with patch("django_q.tasks.async_task"):
        client.post(
            f"/dashboard/students/opening-import/{second.id}/apply/",
            {"revision": replay.revision, "selection_token": str(replay.selection_token)},
        )
    assert run_batch_chunk(club_id=club.id, batch_id=second.id)["counts"] == {"replayed": 1}
    assert Payment.objects.for_club(club).count() == 1
    assert private_path(name=second.prepared_file).stat().st_mode & 0o777 == 0o600


def test_invalid_file_scope_and_revoked_actor(client, club, book_case, owner_user):
    actor, _, content = book_case
    response = client.post(
        "/dashboard/students/opening-import/",
        {"workbook": SimpleUploadedFile("fake.xlsx", b"not zip", content_type=MIME)},
    )
    assert response.status_code == 200 and not OpeningImportBatch.objects.exists()
    upload(client, content())
    batch = OpeningImportBatch.objects.get()
    foreign = OpeningImportBatch.objects.create(
        club=ClubFactory(),
        created_by=owner_user,
        expires_at=batch.expires_at,
    )
    assert client.get(f"/dashboard/students/opening-import/{foreign.id}/export/").status_code == 404
    assert client.get("/dashboard/students/opening-import/").status_code == 200
    assert f"Партия №{foreign.id}".encode() not in client.get("/dashboard/students/opening-import/").content
    ClubMembership.objects.filter(club=club, user=actor).update(role="trainer")
    assert client.get(f"/dashboard/students/opening-import/{batch.id}/").status_code == 403
    assert client.get(f"/dashboard/students/opening-import/{batch.id}/export/").status_code == 403


def test_partial_batch_item_edit_invalidates_old_preview_preserves_keys(client, club, book_case):
    _, terms, content = book_case
    other = replace(terms, entitlement_source_key="second", payment_source_key="second-pay", original_left=6)
    upload(client, content((other,)))
    batch = OpeningImportBatch.objects.get()
    preview = batch.previews.get()
    assert batch.items.filter(status="needs_review").count() == 1
    item = batch.items.get(status="needs_review")
    page = client.get(f"/dashboard/students/opening-import/{batch.id}/items/{item.id}/")
    assert page.status_code == 200
    fields = {field["label"]: field["value"] for field in page.context["fields"]}
    assert "source_subscription_key" not in fields
    fields["Осталось"] = "7"
    fields["Основание отдельной оплаты"] = "Отдельный пакет по источнику"
    response = client.post(f"/dashboard/students/opening-import/{batch.id}/items/{item.id}/", fields)
    assert response.status_code == 302
    item.refresh_from_db()
    batch.refresh_from_db()
    assert batch.revision == 2 and item.source_data["source_subscription_key"] == "second"
    with patch("django_q.tasks.async_task") as queue:
        response = client.post(
            f"/dashboard/students/opening-import/{batch.id}/apply/",
            {"revision": preview.revision, "selection_token": str(preview.selection_token)},
            HTTP_HX_REQUEST="true",
        )
    assert "устарела".encode() in response.content and not queue.called
    assert not Payment.objects.for_club(club).exists()


def test_manual_form_uses_same_draft_and_readable_values(client, club, book_case):
    _, terms, _ = book_case
    payload = terms.as_payload()
    fields = {
        column: payload[field]
        for field, column in COLUMNS.items()
        if field not in {"student_source_key", "entitlement_source_key", "payment_source_key"}
    }
    fields = {k: ("Да" if v else "Нет") if isinstance(v, bool) else "" if v is None else v for k, v in fields.items()}
    fields.update(
        {
            "Исходная цена": payload["paid_amount"],
            "Валюта": "RUB",
            "Исторический долг": "0",
            "Заморожен": "Нет",
            "Тип пакета": "Персональный",
            "Правило начислений": "За занятие",
            "Способ оплаты": "Неизвестен",
        }
    )
    response = client.post("/dashboard/students/opening-import/single/", fields)
    assert response.status_code == 302
    batch = OpeningImportBatch.objects.get()
    assert batch.items.get().status == "ready", batch.items.get().errors
    assert not Payment.objects.for_club(club).exists()
    response = client.post(
        f"/dashboard/students/opening-import/{batch.id}/apply/", {"revision": "1", "selection_token": "broken"}
    )
    assert response.status_code == 200 and "устарела".encode() in response.content


def test_browser_fixture_uses_valid_schema_and_keeps_preview_side_effect_free(settings, tmp_path):
    from apps.common.management.commands.prepare_student_opening_import_e2e import Command
    from apps.students.imports.services import prepare_batch, validate_batch
    from apps.trainers.models import TrainerEarning, TrainerSettlementEntry

    settings.STUDENT_IMPORT_PRIVATE_ROOT = tmp_path / "private"
    settings.STUDENT_OPENING_IMPORT_ENABLED = True
    settings.TRAINER_SETTLEMENTS_ENABLED = True
    settings.TRAINING_GROUP_NEW_WRITES_ENABLED = True
    fixture = Command().create_fixture()
    batch = prepare_batch(
        club_id=fixture["club_id"], actor_user_id=fixture["owner"]["id"], path=fixture["workbook_path"]
    )
    preview = validate_batch(club_id=fixture["club_id"], actor_user_id=fixture["owner"]["id"], batch_id=batch.id)
    assert len(preview.selection) == 4, list(batch.items.values("errors"))
    assert batch.items.filter(status="needs_review").count() == 1
    assert not Student.objects.for_club(fixture["club_id"]).exists()
    assert not Payment.objects.for_club(fixture["club_id"]).exists()
    assert not TrainerEarning.objects.for_club(fixture["club_id"]).exists()
    assert not TrainerSettlementEntry.objects.for_club(fixture["club_id"]).exists()


@pytest.mark.parametrize("change,clear", [(False, False), (True, False), (True, True)])
def test_preview_assignment_matches_apply(client, club, book_case, change, clear):
    from apps.students.tests.factories import StudentFactory
    from apps.trainers.tests.factories import TrainerFactory

    _, terms, content = book_case
    original = TrainerFactory(club=club)
    requested = TrainerFactory(club=club)
    student = StudentFactory(
        club=club, first_name=terms.first_name, last_name=terms.last_name,
        phone=terms.phone, assigned_trainer=original, status=Student.Status.ACTIVE,
    )
    terms = replace(terms, student_id=student.id, assigned_trainer_id=None if clear else requested.id,
                    change_assigned_trainer=change)
    response = upload(client, content(source_terms=terms))
    batch = OpeningImportBatch.objects.get()
    assert batch.items.get().status == "ready", batch.items.get().errors
    page = client.get(response.url)
    item = page.context["items_page"][0]
    assert item.assignment_action == ("preserve" if not change else "clear" if clear else "set")
    expected = original if not change else None if clear else requested
    assert item.assigned_name == (str(expected) if expected else "")
    assert ("текущее назначение сохраняется" if not change else "текущее назначение будет снято" if clear
            else f"назначить {requested}") in page.content.decode()
    preview = batch.previews.get()
    with patch("django_q.tasks.async_task"):
        client.post(f"/dashboard/students/opening-import/{batch.id}/apply/",
                    {"revision": preview.revision, "selection_token": str(preview.selection_token)})
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["status"] == "completed"
    student.refresh_from_db()
    assert student.assigned_trainer_id == (expected.id if expected else None)
    # Even after live drift the accepted review continues to show its saved assignment.
    Student.objects.for_club(club).filter(id=student.id).update(assigned_trainer=requested)
    assert client.get(response.url).context["items_page"][0].assigned_name == item.assigned_name


@pytest.mark.parametrize("history_only", [False, True])
def test_preview_status_workspace_and_historical_policy_match_apply(client, club, book_case, history_only):
    from apps.billing.models import Tariff
    from apps.students.tests.factories import StudentFactory

    _, terms, content = book_case
    student = StudentFactory(club=club, first_name=terms.first_name, last_name=terms.last_name,
                             phone=terms.phone, status=Student.Status.CHURNED)
    terms = replace(terms, student_id=student.id, confirm_student_transition=True,
                    original_used=12 if history_only else 5, original_left=0 if history_only else 7,
                    payout_policy=Tariff.PayoutPolicy.NONE)
    response = upload(client, content(source_terms=terms))
    batch = OpeningImportBatch.objects.get()
    assert batch.items.get().status == "ready", batch.items.get().errors
    page = client.get(response.url)
    item = page.context["items_page"][0]
    expected = Student.Status.CHURNED if history_only else Student.Status.ACTIVE
    assert item.preview_summary["old_status"] == Student.Status.CHURNED
    assert item.preview_summary["new_status"] == expected
    assert f"{item.old_status_label} → {item.new_status_label}" in page.content.decode()
    assert f"Историческое правило начислений: {item.payout_policy_label}" in page.content.decode()
    if history_only:
        assert "Карточка будет в разделе «Ученики»" in page.content.decode()
    preview = batch.previews.get()
    with patch("django_q.tasks.async_task"):
        client.post(f"/dashboard/students/opening-import/{batch.id}/apply/",
                    {"revision": preview.revision, "selection_token": str(preview.selection_token)})
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["status"] == "completed"
    student.refresh_from_db()
    assert student.status == expected and student.lead_status is None and student.became_student_at


@pytest.mark.parametrize("independent", [False, True])
def test_preview_membership_consequence_matches_apply(client, club, book_case, group_command, independent):
    from datetime import timedelta

    from apps.attendance.models import TrainingGroupMembership
    from apps.attendance.tests.factories import TrainingGroupMembershipFactory
    from apps.students.tests.factories import StudentFactory

    _, _, content = book_case
    _, terms, group, _, day = group_command
    if independent:
        student = StudentFactory(club=club, phone=terms.phone, first_name=terms.first_name,
                                 last_name=terms.last_name, status="active")
        TrainingGroupMembershipFactory(club=club, student=student, training_group=group,
                                       starts_on=day - timedelta(days=14))
        terms = replace(terms, student_id=student.id)
    response = upload(client, content(source_terms=terms))
    batch = OpeningImportBatch.objects.get()
    assert batch.items.get().status == "ready", batch.items.get().errors
    page = client.get(response.url)
    item = page.context["items_page"][0]
    assert item.membership_effect == ("independent" if independent else "new_payment_owned")
    assert ("Возврат этой оплаты не удаляет независимое членство" if independent
            else "Эта оплата станет владельцем членства") in page.content.decode()
    preview = batch.previews.get()
    with patch("django_q.tasks.async_task"):
        client.post(f"/dashboard/students/opening-import/{batch.id}/apply/",
                    {"revision": preview.revision, "selection_token": str(preview.selection_token)})
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["status"] == "completed"
    membership = TrainingGroupMembership.objects.for_club(club).get()
    assert membership.authority == ("independent" if independent else "payment_owned")


def test_preview_later_row_preserves_assignment_after_previous_row(client, club, book_case):
    from apps.students.tests.factories import StudentFactory
    from apps.trainers.tests.factories import TrainerFactory

    _, terms, content = book_case
    original = TrainerFactory(club=club)
    requested = TrainerFactory(club=club)
    student = StudentFactory(club=club, first_name=terms.first_name, last_name=terms.last_name,
                             phone=terms.phone, assigned_trainer=original, status=Student.Status.ACTIVE)
    first = replace(terms, student_id=student.id, assigned_trainer_id=requested.id,
                    change_assigned_trainer=True, distinct_payment_reference="First separate purchase")
    second = replace(first, entitlement_source_key="second-package", payment_source_key="second-payment",
                     change_assigned_trainer=False, distinct_payment_reference="Second separate purchase")
    response = upload(client, content((second,), source_terms=first))
    batch = OpeningImportBatch.objects.get()
    assert batch.items.filter(status="ready").count() == 2, list(batch.items.values("errors"))
    page = client.get(response.url)
    second_item = list(page.context["items_page"])[1]
    row_html = page.content.decode().split(f'data-import-item="{second_item.id}"')[1].split("</article>")[0]
    assert "с учётом изменений предыдущих строк партии" in row_html
    assert f"сохраняется — {original}" not in row_html
    preview = batch.previews.get()
    assert all(entry["scope"]["student"]["assigned_trainer_id"] == original.id for entry in preview.selection)
    with patch("django_q.tasks.async_task"):
        client.post(f"/dashboard/students/opening-import/{batch.id}/apply/",
                    {"revision": preview.revision, "selection_token": str(preview.selection_token)})
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["status"] == "completed"
    student.refresh_from_db()
    assert student.assigned_trainer_id == requested.id
    assert Payment.objects.for_club(club).count() == 2


@pytest.mark.parametrize("second_has_trainer", [True, False])
def test_preview_new_same_student_rows_show_conditional_assignment(client, club, book_case, second_has_trainer):
    from apps.trainers.tests.factories import TrainerFactory

    _, terms, content = book_case
    original = TrainerFactory(club=club)
    requested = TrainerFactory(club=club)
    first = replace(terms, assigned_trainer_id=original.id, change_assigned_trainer=False,
                    distinct_payment_reference="First separate purchase")
    second = replace(first, entitlement_source_key="second-package", payment_source_key="second-payment",
                     assigned_trainer_id=requested.id if second_has_trainer else None,
                     distinct_payment_reference="Second separate purchase")
    response = upload(client, content((second,), source_terms=first))
    batch = OpeningImportBatch.objects.get()
    assert batch.items.filter(status="ready").count() == 2, list(batch.items.values("errors"))
    page = client.get(response.url)
    second_item = list(page.context["items_page"])[1]
    assert second_item.assignment_action == "create_or_preserve"
    row_html = page.content.decode().split(f'data-import-item="{second_item.id}"')[1].split("</article>")[0]
    assert "если карточка уже создана предыдущей строкой партии — сохранить её назначение" in row_html
    assert (f"при создании — назначить {requested}" if second_has_trainer
            else "при создании — без ответственного") in row_html
    preview = batch.previews.get()
    assert all(entry["scope"]["student"] is None for entry in preview.selection)
    with patch("django_q.tasks.async_task"):
        client.post(f"/dashboard/students/opening-import/{batch.id}/apply/",
                    {"revision": preview.revision, "selection_token": str(preview.selection_token)})
    assert run_batch_chunk(club_id=club.id, batch_id=batch.id)["status"] == "completed"
    student = Student.objects.for_club(club).get()
    assert student.assigned_trainer_id == original.id
    assert Payment.objects.for_club(club).count() == 2
